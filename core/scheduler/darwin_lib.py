"""Darwinbox Exit Clearance mailer - runs inside the console as the scheduled process "DARWINBOX EXIT CLEARANCE".

Moved here from D:\\DARWIN\\darwinbox_exit_mailer.py (same logic, nothing removed). What one run does
  1. logs in to Darwinbox (saved session is reused; if it is dead it logs in again by itself)
  2. reads the Task Box (separation tasks) and skips every employee that was already e-mailed
  3. for the NEW employees: fetches company / department / e-mail / mobile / PAN and sends ONE e-mail
     (network / SMTP problems are retried every 5 s until the mail is sent or the run is cancelled)
  4. only after the mail is really sent are the employees written to the state file - so nothing is lost and
     NOTHING IS EVER SENT TWICE
  5. when something is waiting: reads the replies of the 4 people over POP3 (read-only), updates Employee_Resign.xlsx,
     and - once Manohar, Jagesh and Anuj have all replied for an employee - checks the amounts against today's RISK_REPORT
     (COMMON FILES) and submits the Darwinbox "Act" form (only when "Submit the Act form" is ON)
  6. 24 h without a reply -> a gentle reminder on the same mail thread; 24 h without the three key replies -> a picture
     by WhatsApp (and e-mail if configured), once per mail

Where things live (everything is inside the project, so git push / pull moves it to another machine):
  data/darwinbox_exit/darwinbox_sent_records.json   THE MEMORY (who was e-mailed, replies, who was submitted). Never delete.
  data/darwinbox_exit/.darwinbox_session.json       saved login cookies (treat like a password)
  data/darwinbox_exit/darwinbox_device.json         browser fingerprint captured once
  data/darwinbox_exit/Employee_Resign.xlsx          rebuilt from the memory after every change
  data/darwinbox_exit/Darwin_Act_Preview.*          what the Act step built / submitted
Passwords / tokens are NOT in this file: Darwinbox login, mailbox and WhatsApp token are set with the Credentials
button of the scheduler process (stored in .env.local, never in git).

Safety: if the memory file is missing the run REFUSES to start (a missing file would mean "everyone is new" and
re-send every old mail). Set DARWIN_ALLOW_EMPTY_STATE=1 only for a brand-new start.

Manual use (outside the scheduler):  python -m core.scheduler.darwin_lib --dry-run | --replies-only | --act-only |
  --mark-done M0214691 | --reminder-preview | --alert-preview | --no-act | --no-reminder | --no-alert | --fresh
"""
from __future__ import annotations

import ast
import base64
import csv
import html
import json
import os
import poplib
import re
import smtplib
import socket
import ssl
import sys
import time
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from datetime import date, datetime, timedelta
from email import message_from_bytes, policy
from email.message import EmailMessage
from email.utils import formatdate, make_msgid, parseaddr, parsedate_to_datetime
from pathlib import Path
from urllib.parse import parse_qs, parse_qsl, unquote, urljoin, urlparse

import json5
import requests
from bs4 import BeautifulSoup
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import padding


# ==========================================================================
# CONFIG - everything secret is filled in at run time by configure() from the console's Credentials
# (scheduler process "DARWINBOX EXIT CLEARANCE"); only non-secret defaults live here.
# ==========================================================================
# --- Darwinbox ---------------------------------------------------------------------------------
USERNAME = ""                      # set by configure()  (Credentials > Darwinbox)
PASSWORD = ""

# --- E-mail (SMTP) -----------------------------------------------
SMTP_HOST = "smtp3.netcore.co.in"
SMTP_PORT = 587
SMTP_SECURITY = "none"             # "starttls" | "ssl" | "none"  (as in the original script)
SMTP_USER = ""                     # set by configure()
SMTP_PASSWORD = ""
FROM_ADDR = ""

TO = ["manohar@mastertrust.co.in", "jagesh.aggarwal@mastertrust.co.in", "anuj.joshi@mastertrust.co.in"]
CC = ["hrd@mastertrust.co.in", "rms_del@mastertrust.co.in", "suprita.b@mastertrust.co.in", "va@mastertrust.co.in",
      "palkaachopra@mastertrust.co.in"]
SUBJECT = "Employee Clearance"

RETRY_DELAY = 5                    # seconds between retries (SMTP and network errors)
MASK_PAN = False                   # True -> PAN shown as XXXXXX491G in the e-mail

# --- Replies (POP3). Read-only: the script never deletes or changes anything in the mailbox -----
POP_HOST = "pop9.netcore.co.in"
POP_PORT = 995                     # SSL POP3
POP_USER = ""                      # set by configure()
POP_PASSWORD = ""
REPLY_SENDERS = {                  # who must reply to every mail  (address -> name shown in the log)
    "hrd@mastertrust.co.in": "HR Team",
    "jagesh.aggarwal@mastertrust.co.in": "Jagesh",
    "anuj.joshi@mastertrust.co.in": "Anuj",
    "manohar@mastertrust.co.in": "Manohar",
}
REPLY_WATCH_DAYS = 14              # keep waiting for replies to a mail for this many days
COLUMN_ALIASES = {                 # same meaning, different spelling -> ONE column in the Excel
    "Shortfall amount": ["S/F AMOUNT", "S/F", "SF", "SF amount", "SF amt", "S/F amt", "Shortfall",
                         "Shortfall amount", "Shortfall amt"],
}
POP_ATTEMPTS = 5                   # connection attempts per run (then it simply tries again next run)
POP_STOP_AFTER_OLD = 25            # scanning newest->oldest stops after this many mails older than our oldest sent mail

# --- Act step (after the replies are in: check shortfalls, fill the Darwinbox "Act" form) ------------
ACT_ENABLED = True                 # False = a normal run never touches the Act step
ACT_SUBMIT = True                  # set by configure() from the process setting "Submit the Act form".
                                   # False = TEST MODE (form built and saved to Darwin_Act_Preview.*, NOTHING submitted)
ACT_REQUIRE = ("Manohar", "Jagesh", "Anuj")      # an employee is handled only after ALL of these have replied
ACT_RISK_COLUMN = "ShortFall"      # column read from RISK_REPORT_<ddmmyyyy>.csv (today's file, COMMON FILES)
                                   # (a NEGATIVE ShortFall counts by its size, a POSITIVE one counts as 0)
ACT_TOLERANCE_PCT = 1.0            # Jagesh's amount must be within this % of the risk report, else the step STOPS
ACT_MANUAL = {}                    # {"M0214647": {"Manohar": "NA", "Jagesh": "NA", "Anuj": "..."}} for replies NOT in a table
ACT_SKIP_EMP_IDS = []              # Emp IDs the Act step must leave alone
ACT_ALERT_TO = []                  # who is e-mailed when the Act step stops on a mismatch (configure() sets the sender)
ACT_FIXED_CHOICES = {              # start of the question text -> answer (same for everybody)
    "closure of master care": "No",
    "removal of mobile number & email from savior": "No",
    "removal of mobile number & email from master dealer": "Yes",
}
ACT_TOTAL_QUESTIONS = ("outstanding in personal account", "outstanding in family, relative account",
                       "outstanding in client account")      # all three get the combined shortfall
ACT_REMARKS_QUESTION = "final remarks/comments"

# --- Reply-overdue alert: if a mail has had no reply from Manohar / Jagesh / Anuj for ALERT_AFTER_HOURS ----------
ALERT_ENABLED = True
ALERT_AFTER_HOURS = 24
ALERT_FOR = "Amit Pundir"          # whose work this is - named on the picture and in the e-mail
ALERT_EMAIL_TO = []                # e-mail the picture to these (empty = no e-mail)
WA_ENABLED = True
WA_TO = ["+917889026088", "+919871410617"]     # set by configure() from the process setting "Overdue-alert WhatsApp numbers"
WA_AUTH_TOKEN = ""                 # set by configure()  (Credentials > WhatsApp)
WA_UPLOAD_URL = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/media"
WA_SEND_URL = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/messages"
WA_TEMPLATE = "daily_report"       # approved template: header = image, body = 2 text values
WA_LANGUAGE = "en"
WA_TEXT1 = "EMPLOYEE CLEARANCE - REPLY OVERDUE"      # body value 1  (value 2 is today's date, dd-mm-yyyy)
WA_ATTEMPTS = 3

# --- Gentle reminder: a reply on the thread of the original mail, to each person who has not replied -------------
REMINDER_ENABLED = True
REMINDER_AFTER_HOURS = 24
REMINDER_CC = list(CC)

# --- Fixed values (captured from your own browser - they are NOT secrets) ----------------------
BASE_URL = "https://mastertrust.darwinbox.in"
DEVICE_FINGERPRINT = "1iks17d"     # ThumbmarkJS value the login page computes in the browser
DEVICE_USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                     "(KHTML, like Gecko) Chrome/154.0.0.0 Safari/537.36")
EXPECTED_EXCLUDE = ["canvas", "locales", "audio", "plugins"]   # fingerprint recipe at capture time
LANDING_PAGE = "dashboard"
REMEMBER_ME = True

# --- Details fetched per employee --------------------------------------------------------------
DETAIL_FIELDS = ["bank_pan_num"]
DEFAULT_FORM_ID = "a6878e79293677"
CALL_DELAY = 0.4                   # seconds between per-employee requests
PAGE_SIZE_HINT = 10                # the Task Box loads 10 rows per page
LOCK_STALE_HOURS = 6               # a lock file older than this is ignored (crashed run)

# --- run-time hooks (set by the scheduler process) ------------------------------------------------
ARGS: list = []                    # flags such as --dry-run / --no-act (the scheduler passes none)
LOG_HOOK = None                    # callable(str) - every log line is also sent here (the scheduler's run log)
CANCEL = None                      # callable() -> True when the run was cancelled: retry loops stop
STATS: dict = {}                   # what this run did (see _reset_stats)

# ==========================================================================
_REPO = Path(__file__).resolve().parents[2]
HERE = Path(os.getenv("DARWIN_DATA_DIR") or (_REPO / "data" / "darwinbox_exit"))
HERE.mkdir(parents=True, exist_ok=True)
COOKIE_FILE = HERE / ".darwinbox_session.json"
STATE_FILE = HERE / "darwinbox_sent_records.json"
LOG_FILE = HERE / "darwinbox_exit_mailer.log"
LOCK_FILE = HERE / "darwinbox_exit_mailer.lock"
PREVIEW_FILE = HERE / "email_preview.html"
EXCEL_FILE = HERE / "Employee_Resign.xlsx"
ACT_PREVIEW_XLSX = HERE / "Darwin_Act_Preview.xlsx"
ACT_PREVIEW_JSON = HERE / "Darwin_Act_Preview.json"
ACT_BLOCKED_FILE = HERE / "Darwin_Act_BLOCKED.txt"

LOGIN_URL = f"{BASE_URL}/user/login?source=logout"
HOME_URL = f"{BASE_URL}/ms/db/home"
TASKS_URL = f"{BASE_URL}/tasksApi/GetTasks"
COMPLETE_URL = f"{BASE_URL}/separation/completeaction"
DETAILS_URL = f"{BASE_URL}/formsapi/getdetailsvalues"
FORM_DETAIL_URL = f"{BASE_URL}/formsapi/getFormDetailById"
COMPLETE_ACTION_URL = f"{BASE_URL}/separation/CompleteAction"
CSRF_COOKIE = "pbqBeYWPUn"
TIMEOUT = 30
MAX_REDIRECTS = 10

COLUMNS = ["S.No", "Name", "Emp ID", "Designation", "Separation date", "Company",
           "Deptt / Location", "E-mail id", "Mobile No.", "Pan No."]

BODY_LINES = [   # (greeting, request) - grammar corrected
    ("Dear HR Team,", "Please update the contact details as per the data given below."),
    ("Dear Jagesh,", "Please provide the shortfall from BIS."),
    ("Dear Anuj,", "Please provide the details of the dealer ID, if any."),
    ("Dear Manohar,", "Please provide the trading codes, if any."),
]

SIGNATURE = ["Best regards,", "Sandeep Kumar", "011-42111079"]


# ==========================================================================
# 🎨 Logging (terminal + log file)
# ==========================================================================
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if os.name == "nt":
    os.system("")
_COLOR = sys.stdout.isatty() and not os.environ.get("NO_COLOR")
_C = {"g": "32", "y": "33", "r": "31", "c": "36", "b": "1", "d": "2"}


def log(emoji: str, msg: str, color: str = "") -> None:
    stamp = f"{datetime.now():%Y-%m-%d %H:%M:%S}"
    start, end = (f"\033[{_C[color]}m", "\033[0m") if (_COLOR and color) else ("", "")
    print(f"{stamp}  {emoji}  {start}{msg}{end}", flush=True)
    try:
        with LOG_FILE.open("a", encoding="utf-8") as fh:
            fh.write(f"{stamp}  {emoji}  {msg}\n")
    except OSError:
        pass
    if LOG_HOOK:
        try:
            LOG_HOOK(f"{emoji} {msg}")
        except Exception:
            pass


def info(m): log("🔹", m, "c")
def ok(m): log("✅", m, "g")
def warn(m): log("⚠️ ", m, "y")
def fail(m): log("❌", m, "r")


class AppError(Exception):
    pass


def _reset_stats() -> None:
    STATS.clear()
    STATS.update(emailed=0, replies=0, submitted=0, reminders=0, alerts=0, idle=False, blocked=False)


_reset_stats()


def _check_cancel() -> None:
    if CANCEL and CANCEL():
        raise AppError("The run was cancelled.")


def mask(text: str) -> str:
    return text[:2] + "•" * (len(text) - 4) + text[-2:] if len(text) > 4 else "•" * len(text)


# ==========================================================================
# 🔑 Login (RSA exactly like the page's JavaScript)
# ==========================================================================
EXCLUDE_RE = re.compile(r"setOption\(\s*['\"]exclude['\"]\s*,\s*(\[[^\]]*\])\s*\)")


@dataclass
class LoginPage:
    url: str
    public_key_pem: str
    query_fields: list
    redirect_options: list
    remember_value: str | None
    verify_code: tuple | None
    turnstile: bool
    exclude: list | None
    notes: list = field(default_factory=list)


def extract_public_key_pem(page_html: str) -> str:
    m = re.search(r"-----BEGIN PUBLIC KEY-----(.*?)-----END PUBLIC KEY-----", page_html, re.S)
    if not m:
        raise AppError("RSA public key not found in the login page (page changed?).")
    body = m.group(1).replace("\\n", "").replace("\\r", "")
    body = re.sub(r"[\"'`]\s*\+\s*[\"'`]", "", body)
    body = re.sub(r"[^A-Za-z0-9+/=]", "", body)
    lines = [body[i:i + 64] for i in range(0, len(body), 64)]
    return "-----BEGIN PUBLIC KEY-----\n" + "\n".join(lines) + "\n-----END PUBLIC KEY-----\n"


def rsa_encrypt_b64(key, text: str) -> str:
    data = text.encode("utf-8")
    if len(data) > key.key_size // 8 - 11:
        raise AppError("Value too long for the RSA key.")
    return base64.b64encode(key.encrypt(data, padding.PKCS1v15())).decode()


def parse_login_page(page_html: str, url: str) -> LoginPage:
    soup = BeautifulSoup(page_html, "html.parser")
    for el_id in ("UserLogin_username", "UserLogin_password"):
        if not soup.find(id=el_id):
            raise AppError(f"Login page layout changed: #{el_id} not found.")
    sel = soup.find("select", id="UserLogin_redirectpage")
    options = [(o.get("value", ""), o.get_text(" ", strip=True)) for o in sel.find_all("option")] if sel else []
    remember = soup.find("input", id="Field", attrs={"type": "checkbox"})
    verify = soup.find(id="UserLogin_verifyCode")
    turn = soup.find(id="cf-turnstile-token") or soup.select_one(".cf-turnstile, [data-sitekey]")
    exclude = None
    for tag in soup.find_all("script"):
        code = tag.get_text()
        if "inputDetails" in code:
            m = EXCLUDE_RE.search(code)
            if m:
                try:
                    exclude = [str(x) for x in ast.literal_eval(m.group(1))]
                except (ValueError, SyntaxError):
                    pass
            break
    return LoginPage(
        url=url, public_key_pem=extract_public_key_pem(page_html),
        query_fields=parse_qsl(urlparse(url).query, keep_blank_values=True),
        redirect_options=options,
        remember_value=remember.get("value", "on") if remember else None,
        verify_code=(verify.get("type", "text"), verify.get("value", "")) if verify else None,
        turnstile=bool(turn), exclude=exclude)


def build_payload(page: LoginPage, enc_user: str, enc_pass: str) -> dict:
    if page.turnstile:
        raise AppError("Cloudflare Turnstile is enabled on the login page - a plain HTTP client cannot pass it.")
    if page.verify_code and page.verify_code[0] != "hidden":
        raise AppError("The login page now asks for a verification code - this script cannot fill it.")
    payload = dict(page.query_fields)
    payload["UserLogin[username]"] = enc_user
    payload["UserLogin[password]"] = enc_pass
    if page.redirect_options:
        values = [v for v, _ in page.redirect_options]
        if LANDING_PAGE not in values:
            raise AppError(f"LANDING_PAGE={LANDING_PAGE!r} is not offered by the page: {values}")
        payload["UserLogin[redirectpage]"] = LANDING_PAGE
    if REMEMBER_ME and page.remember_value is not None:
        payload["Field"] = page.remember_value
    if page.verify_code and page.verify_code[1]:
        payload["UserLogin[verifyCode]"] = page.verify_code[1]
    payload["UserLogin[fingerprint]"] = DEVICE_FINGERPRINT
    return payload


def is_redirect(resp: requests.Response) -> bool:
    return resp.status_code in (301, 302, 303, 307, 308)


def bounced_to_login(resp: requests.Response) -> bool:
    return resp.status_code == 401 or (is_redirect(resp) and "/user/log" in resp.headers.get("Location", ""))


class Darwinbox:
    def __init__(self) -> None:
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": DEVICE_USER_AGENT,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
            "Upgrade-Insecure-Requests": "1",
        })

    # ---- every HTTP call goes through here: network errors / 5xx are retried forever --------------
    def _raw(self, method: str, url: str, **kw) -> requests.Response:
        kw.setdefault("timeout", TIMEOUT)
        kw.setdefault("allow_redirects", False)
        while True:
            _check_cancel()
            try:
                resp = self.session.request(method, url, **kw)
                if resp.status_code in (502, 503, 504):
                    raise requests.ConnectionError(f"HTTP {resp.status_code}")
                return resp
            except (requests.ConnectionError, requests.Timeout, requests.exceptions.ChunkedEncodingError) as exc:
                warn(f"Darwinbox network problem ({type(exc).__name__}) - retrying in {RETRY_DELAY}s")
                time.sleep(RETRY_DELAY)

    def request(self, method: str, url: str, **kw) -> requests.Response:
        """Like _raw, but if the session turns out to be dead it logs in again and retries once."""
        resp = self._raw(method, url, **kw)
        if bounced_to_login(resp):
            warn("Session rejected by Darwinbox - logging in again")
            self.session.cookies.clear()
            self.login()
            resp = self._raw(method, url, **kw)
            if bounced_to_login(resp):
                raise AppError("Still redirected to login after a fresh login.")
        return resp

    def _follow(self, resp: requests.Response) -> requests.Response:
        for _ in range(MAX_REDIRECTS):
            if not is_redirect(resp):
                return resp
            nxt = urljoin(resp.url, resp.headers.get("Location", ""))
            if not nxt or nxt == resp.url:
                return resp
            resp = self._raw("GET", nxt)
        raise AppError("Too many redirects.")

    # ---- cookies ----------------------------------------------------------------------------------
    def save_cookies(self) -> None:
        data = [{"name": c.name, "value": c.value, "domain": c.domain, "path": c.path} for c in self.session.cookies]
        COOKIE_FILE.write_text(json.dumps(data), encoding="utf-8")

    def load_cookies(self) -> bool:
        if not COOKIE_FILE.exists():
            return False
        try:
            for c in json.loads(COOKIE_FILE.read_text(encoding="utf-8")):
                self.session.cookies.set(c["name"], c["value"], domain=c["domain"], path=c["path"])
            return True
        except (ValueError, KeyError, TypeError):
            return False

    # ---- login ------------------------------------------------------------------------------------
    def login(self) -> None:
        if not USERNAME or not PASSWORD:
            raise AppError("Fill USERNAME and PASSWORD at the top of the script.")
        info("Logging in to Darwinbox")
        resp = self._raw("GET", LOGIN_URL, allow_redirects=True)
        resp.raise_for_status()
        page = parse_login_page(resp.text, resp.url)
        if page.exclude is not None and page.exclude != EXPECTED_EXCLUDE:
            warn(f"The login page's fingerprint recipe changed ({page.exclude}); if login starts failing the "
                 "fingerprint must be captured again.")
        key = serialization.load_pem_public_key(page.public_key_pem.encode())
        payload = build_payload(page, rsa_encrypt_b64(key, USERNAME), rsa_encrypt_b64(key, PASSWORD))
        origin = f"{urlparse(page.url).scheme}://{urlparse(page.url).netloc}"
        resp = self._raw("POST", page.url, data=list(payload.items()),
                         headers={"Origin": origin, "Referer": page.url})
        resp = self._follow(resp)
        path = urlparse(resp.url).path.lower()
        if "otpauth" in path:
            raise AppError("Darwinbox asked for an OTP - this unattended script cannot answer it.")
        if "/user/login" in path:
            msgs = [el.get_text(" ", strip=True)[:150]
                    for el in BeautifulSoup(resp.text, "html.parser").find_all(class_=re.compile("alert|error|danger", re.I))]
            raise AppError(f"Login rejected by Darwinbox. Page says: {msgs[:3] or 'nothing'}")
        self.save_cookies()
        ok(f"Logged in as {mask(USERNAME)} (landed on {urlparse(resp.url).path})")

    # ---- Task Box page ----------------------------------------------------------------------------
    def get_tasks_html(self, fresh: bool = False) -> str:
        if not fresh and self.load_cookies():
            info("Using the saved session")
        else:
            self.session.cookies.clear()
            self.login()
        resp = self.request("GET", TASKS_URL, headers={"Referer": HOME_URL})   # re-logs in itself if dead
        if is_redirect(resp):
            resp = self._follow(resp)
        resp.raise_for_status()
        self.save_cookies()
        return resp.text


# ==========================================================================
# 📖 Task Box rows
# ==========================================================================
def _balanced_literal(code: str, start: int) -> str:
    depth, quote, i = 0, None, start
    while i < len(code):
        ch = code[i]
        if quote:
            if ch == "\\":
                i += 1
            elif ch == quote:
                quote = None
        elif ch in ("'", '"'):
            quote = ch
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return code[start:i + 1]
        i += 1
    raise AppError("Could not find the end of the Task Box table definition (page changed?).")


def extract_rows(page_html: str) -> list[dict]:
    soup = BeautifulSoup(page_html, "html.parser")
    for tag in soup.find_all("script", src=False):
        code = tag.get_text()
        if "initialTableData" in code and "core_taskbox_tasks" in code:
            start = code.index("{", code.index("tableOptions"))
            return list(json5.loads(_balanced_literal(code, start))["options"].get("initialTableData", []))
    raise AppError("Task Box table not found in the page (not logged in, or the page layout changed).")


def _text(fragment: str) -> str:
    return BeautifulSoup(fragment or "", "html.parser").get_text(" ", strip=True)


def parse_row(row: dict) -> dict:
    emp = BeautifulSoup(row.get("Employee", ""), "html.parser")
    hover, link, code_el = emp.find("dbx-hover-data"), emp.find("a", href=True), emp.select_one(".label-two")
    user_id = hover.get("data-id") if hover else ""
    if not user_id and link:
        m = re.search(r"/id/(\d+)", link["href"])
        user_id = m.group(1) if m else ""
    dol = None
    m = re.search(r"\b(\d{2})-(\d{2})-(\d{4})\b", _text(row.get("LWD", "")))
    if m:
        try:
            dol = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass
    # meta.action_id = "<actionid>_<sep_flow>_<task_id>"  (the page's own JavaScript splits it the same way)
    meta = row.get("meta") or {}
    action_id, sep_flow, task_id = (str(meta.get("action_id", "")).split("_") + ["", "", ""])[:3]
    name = hover.get_text(" ", strip=True) if hover else (link.get_text(" ", strip=True) if link else "")
    emp_code = code_el.get_text(" ", strip=True) if code_el else ""
    return {
        "Name": name, "Emp ID": emp_code, "Designation": _text(row.get("Designation", "")),
        "Separation date": dol, "Company": "", "Deptt / Location": short_location(_text(row.get("Location", ""))),
        "E-mail id": "", "Mobile No.": "", "Pan No.": "",
        "_key": emp_code or user_id or task_id, "_user_id": user_id, "_task_id": task_id or row.get("task_id", ""),
        "_action_id": action_id, "_sep_flow": sep_flow, "_category": meta.get("category", "separation_work_task"),
        "_loc": _text(row.get("Location", "")), "_ok": False, "_note": "",
    }


def short_location(loc: str) -> str:
    """'Delhi, Delhi, India, (Delhi)' -> 'Delhi (Delhi)'"""
    parts = [p.strip() for p in (loc or "").split(",") if p.strip()]
    if len(parts) > 1 and parts[-1].startswith("("):
        return f"{parts[0]} {parts[-1]}"
    return parts[0] if parts else ""


# ==========================================================================
# 👤 Per-employee details (what the "Act" button does, then the details API)
# ==========================================================================
def csrf_token(client: Darwinbox) -> str:
    value = ""
    for c in client.session.cookies:
        if c.name == CSRF_COOKIE:
            value = c.value
    if not value:
        raise AppError(f"Cookie '{CSRF_COOKIE}' not found.")
    return value


def _ajax_headers(client: Darwinbox, json_body: bool = False) -> dict:
    h = {"Accept": "application/json, text/plain, */*", "Origin": BASE_URL, "Referer": TASKS_URL,
         "csrf-token": csrf_token(client), "X-Requested-With": "XMLHttpRequest"}
    if json_body:
        h["Content-Type"] = "application/json"
    return h


def _json(resp: requests.Response, label: str) -> dict:
    if is_redirect(resp):
        raise AppError(f"{label}: redirected to {resp.headers.get('Location')} (session problem?)")
    try:
        data = resp.json()
        if not isinstance(data, dict):
            raise ValueError
        return data
    except ValueError:
        raise AppError(f"{label}: HTTP {resp.status_code}, reply is not JSON (starts: {resp.text[:100]!r})")


def open_task_form(client: Darwinbox, rec: dict) -> dict:
    data = {"actionid": rec["_action_id"], "sep_flow": rec["_sep_flow"], "mode": "open", "task_id": rec["_task_id"],
            "is_submit_allowed": "1", "fromTaskBox": "true", "category": rec["_category"],
            CSRF_COOKIE: csrf_token(client)}
    return _json(client.request("POST", COMPLETE_URL, data=data, headers=_ajax_headers(client)), "open task form")


def form_params(reply: dict) -> dict:
    """form_id / form_process_id / workflow id come from the <db-form ...> tag in the reply's HTML."""
    found: dict[str, str] = {}
    soup = BeautifulSoup(str(reply.get("update") or ""), "html.parser")
    form = soup.find("db-form")
    if form:
        for attr, key in (("form-id", "form_id"), ("form-process-id", "form_process_id"),
                          ("custom-flow-request-id", "custom_workflow_id"), ("type", "type"),
                          ("dynamic-config", "dynamic_config"), ("config", "config")):
            if form.get(attr):
                found[key] = str(form[attr])
    for name, key in (("FormFramework[id]", "form_id"), ("SeprationEmployeeWorkFLow[id]", "custom_workflow_id")):
        inp = soup.find("input", attrs={"name": name})
        if inp and inp.get("value"):
            found.setdefault(key, inp["value"])
    for k, v in parse_qs(str(reply.get("form_query_params") or "").lstrip("?"), keep_blank_values=True).items():
        key = re.sub(r"[^a-z]", "", k.lower())
        if key == "formid":
            found.setdefault("form_id", v[0])
        elif key in ("formprocessid", "processid"):
            found.setdefault("form_process_id", v[0])
        elif key in ("customworkflowid", "workflowid", "wfid"):
            found.setdefault("custom_workflow_id", v[0])
    return found


def get_details(client: Darwinbox, rec: dict, params: dict) -> dict:
    if not params.get("form_process_id"):
        raise AppError("form_process_id not found in the task form reply")
    body = {
        "add_variables": 1, "fields": DETAIL_FIELDS, "array_fields": [], "separation_fields": [],
        "type": params.get("type") or "separation_workflow", "preset_dropdown": [],
        "preset_dropdown_dynamic_list": [], "engagement_themes": [], "pillar": [], "scale": [],
        "user_id": rec["_user_id"], "survey_id": None, "form_type": 1, "employeement_details": [],
        "token": None, "reviewParameter": [], "form_id": params.get("form_id") or DEFAULT_FORM_ID,
        "reviewAssessment": [], "encryptedUserId": None, "form_process_id": params["form_process_id"],
        "module": "default", "candidate_id": "", "config": params.get("config", ""),
        "dynamic_config": params.get("dynamic_config") or '{"type":"separation_workflow"}',
        "custom_workflow_id": params.get("custom_workflow_id") or rec["_sep_flow"],
        "is_form_disabled": False, "user_language": None, "is_preview": False,
        "hasDisableEditingField": False, "field_of_lookup_field_objects": {}, CSRF_COOKIE: csrf_token(client),
    }
    return _json(client.request("POST", DETAILS_URL, json=body, headers=_ajax_headers(client, True)),
                 "employee details")


def apply_details(rec: dict, reply: dict) -> None:
    v = reply.get("add_variables_value")
    if not v:
        raise AppError(f"details call refused: {reply.get('message') or 'no data returned'}")
    name = " ".join(x for x in (v.get("firstname"), v.get("lastname")) if x).strip()
    rec["Name"] = name or rec["Name"]
    rec["Emp ID"] = v.get("employee_no") or rec["Emp ID"]
    rec["Designation"] = v.get("designation") or rec["Designation"]
    rec["Company"] = v.get("company_name") or ""
    dept = (v.get("department") or "").strip()
    loc = short_location(rec["_loc"] or (v.get("location") or ""))
    rec["Deptt / Location"] = " / ".join(x for x in (dept, loc) if x) or rec["Deptt / Location"]
    rec["E-mail id"] = (v.get("email") or "").strip()
    mobile = (v.get("primary_mobile") or v.get("personal_number") or "").strip()
    code = (v.get("primary_mobile_country_code") or "").strip().lstrip("+")
    rec["Mobile No."] = f"+{code} {mobile}" if mobile and code and not mobile.startswith("+") else mobile
    pan = str(((reply.get("fields_value") or {}).get("bank_pan_num"))
              or ((reply.get("field_data_values") or {}).get("bank_pan_num")) or "").strip().upper()
    rec["Pan No."] = ("X" * (len(pan) - 4) + pan[-4:]) if (MASK_PAN and len(pan) > 4) else pan
    rec["_ok"] = True


def enrich(client: Darwinbox, records: list[dict]) -> None:
    for i, rec in enumerate(records, start=1):
        label = f"{rec['Name']} ({rec['Emp ID']})"
        try:
            apply_details(rec, get_details(client, rec, form_params(open_task_form(client, rec))))
            ok(f"[{i}/{len(records)}] details fetched: {label}")
        except (AppError, requests.RequestException) as exc:
            rec["_note"] = str(exc)
            warn(f"[{i}/{len(records)}] {label}: {exc} - will be retried on the next run")
        if i < len(records):
            time.sleep(CALL_DELAY)


# ==========================================================================
# 🗂️  Already-sent memory
# ==========================================================================
def load_state() -> dict:
    state: dict = {"sent": {}, "batches": [], "pop_seen2": [], "act_alert": ""}
    if not STATE_FILE.exists() and os.getenv("DARWIN_ALLOW_EMPTY_STATE", "").strip().lower() not in ("1", "true", "yes"):
        raise AppError(f"{STATE_FILE.name} is missing in {HERE}. Without it every employee would look NEW and all the old "
                       "mails would be sent again, so the run is refused. Copy the file back (git pull / backup), or set "
                       "DARWIN_ALLOW_EMPTY_STATE=1 for a brand-new start.")
    if STATE_FILE.exists():
        try:
            data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        except ValueError:
            raise AppError(f"{STATE_FILE.name} is damaged. Fix or delete it (deleting means everyone is sent again).")
        if isinstance(data, dict):
            state["sent"] = data.get("sent") or {}
            state["batches"] = data.get("batches") or []
            state["pop_seen2"] = data.get("pop_seen2") or []
            state["act_alert"] = data.get("act_alert") or ""
            if data.get("reply_parser") != REPLY_PARSER:      # parsing improved: read the mailbox again once
                state["pop_seen2"] = []
    if state["sent"] and not state["batches"]:
        # written by an older version: group the sent employees by the minute they were e-mailed
        groups: dict[str, list] = {}
        for key, v in state["sent"].items():
            groups.setdefault(v.get("sent_at", ""), []).append(key)
        state["batches"] = [{"message_id": None, "sent_at": at, "keys": keys, "replies": {}}
                            for at, keys in sorted(groups.items()) if at]
    return state


def save_state(state: dict) -> None:
    tmp = STATE_FILE.with_suffix(".tmp")
    state["reply_parser"] = REPLY_PARSER
    tmp.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, STATE_FILE)


def record_of(rec: dict) -> dict:
    return {c: (rec[c].isoformat() if isinstance(rec[c], date) else rec[c]) for c in COLUMNS if c != "S.No"}


# ==========================================================================
# 📗 Employee_Resign.xlsx  (saved next to the script, never attached to the mail)
# ==========================================================================
REPLY_PARSER = 3                   # bump when reply parsing changes: mailbox is then re-read once


def norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


STANDARD_NORM = {norm(c) for c in COLUMNS} | {"empcode", "employeecode", "employeeid", "contactno", "pan", "co",
                                              "dol", "slno", "srno", "mobile", "phone", "phoneno"}
MOBILE_NORM = {"mobileno", "contactno", "mobile", "phone", "phoneno"}
EMP_NORM = {"empid", "empcode", "employeeid", "employeecode"}


ALIAS_NORM = {norm(sp): name for name, spellings in COLUMN_ALIASES.items() for sp in spellings}


def column_label(head: str) -> str:
    """Header text people typed -> the Excel header: alias if known, else 'First word capital, rest small'."""
    alias = ALIAS_NORM.get(norm(head))
    if alias:
        return alias
    words = head.split()
    if not words:
        return head
    first = words[0] if (words[0].isupper() and len(words[0]) <= 4) else words[0].capitalize()
    return " ".join([first] + [w.lower() for w in words[1:]])


def reply_extras(state: dict) -> tuple[list, dict, dict]:
    """(extra column names in first-seen order, {emp: {column: value}}, {emp: mobile}) from all replies."""
    order: list[str] = []
    by_emp: dict[str, dict] = {}
    mobile: dict[str, str] = {}
    replies = sorted((r for b in state["batches"] for r in b["replies"].values()), key=lambda r: r["received_at"])
    for r in replies:
        for emp, cols in r["rows"].items():
            for head, val in cols.items():
                if not val:
                    continue
                n = norm(head)
                if n in MOBILE_NORM:
                    mobile[emp] = val
                elif n not in STANDARD_NORM:
                    label = column_label(head)
                    if label not in order:
                        order.append(label)
                    by_emp.setdefault(emp, {})[label] = val
    return order, by_emp, mobile


def write_excel(state: dict) -> None:
    """Rebuild the workbook from everything remembered. Rebuilding (instead of appending) means a run that
    could not write the file (e.g. it was open in Excel) is repaired by the next run."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        warn("openpyxl is not installed - Employee_Resign.xlsx not written (pip install openpyxl)")
        return
    records = [v["record"] for v in state["sent"].values() if v.get("record")]
    extras, by_emp, mobile = reply_extras(state)
    headers = COLUMNS + extras
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Resign"
    ws.append(headers)
    for n, r in enumerate(records, start=1):
        emp = str(r.get("Emp ID") or "").upper()
        row = [n]
        for c in COLUMNS[1:]:
            v = r.get(c) or None
            if c == "Mobile No." and not v:
                v = mobile.get(emp)
            if c == "Separation date" and v:
                try:
                    v = date.fromisoformat(v)
                except ValueError:
                    pass
            row.append(v)
        row += [by_emp.get(emp, {}).get(x) for x in extras]
        ws.append(row)
    thin = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Calibri", size=11, bold=(cell.row == 1))
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="B4C6E7")
            elif headers[cell.column - 1] in ("Emp ID", "Mobile No.", "Pan No."):
                cell.number_format = "@"
            elif headers[cell.column - 1] == "Separation date":
                cell.number_format = "DD-MMM-YY"
    for idx in range(1, len(headers) + 1):
        longest = max(len(str(c.value)) for c in ws[get_column_letter(idx)] if c.value is not None)
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 4, 60)
    ws.freeze_panes = "A2"
    tmp = EXCEL_FILE.with_name(EXCEL_FILE.stem + ".tmp.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, EXCEL_FILE)
        ok(f"{EXCEL_FILE.name} updated ({len(records)} employee(s), {len(extras)} reply column(s))")
    except OSError as exc:
        warn(f"Could not write {EXCEL_FILE.name} ({exc}). Close it if it is open in Excel - "
             "the next run rebuilds it.")
        tmp.unlink(missing_ok=True)


# ==========================================================================
# 📥 Replies (POP3)
# ==========================================================================
_QUOTE_HTML = re.compile(
    r'<div[^>]+id=["\']?(?:divRplyFwdMsg|appendonsend)|<blockquote|class=["\']?gmail_quote|'
    r'-{2,}\s*Original Message|<b>\s*(?:<span[^>]*>)?\s*From:|border-top:\s*solid\s*#(?:B5C4DF|E1E1E1)', re.I)
_QUOTE_TEXT = re.compile(r"^\s*(?:-{2,}\s*Original Message|From:\s.+|On .{5,120}wrote:|>.*)$", re.I | re.M)


def cut_quoted_html(page: str) -> str:
    """Keep only what the replier wrote: cut at the start of the quoted original."""
    m = _QUOTE_HTML.search(page)
    if not m:
        return page
    lt = page.rfind("<", 0, m.start() + 1)
    return page[: lt if lt != -1 else m.start()]


def cut_quoted_text(text: str) -> str:
    m = _QUOTE_TEXT.search(text)
    return text[: m.start()] if m else text


def message_bodies(msg) -> tuple[str, str]:
    out = []
    for kind in ("html", "plain"):
        try:
            part = msg.get_body((kind,))
            out.append(part.get_content() if part else "")
        except Exception:
            out.append("")
    return out[0], out[1]


def parse_reply_tables(page: str) -> dict:
    """{EMP ID: {column header: value}} from EVERY table in the mail that has an Emp ID column - the replier's
    own text, a forwarded trail, a quoted original. If a cell appears in several tables the first one wins."""
    result: dict[str, dict] = {}
    clean = lambda c: re.sub(r"\s+", " ", c.get_text(" ", strip=True))
    for table in BeautifulSoup(page, "html.parser").find_all("table"):
        trs = table.find_all("tr")
        if len(trs) < 2:
            continue
        header = [clean(c) for c in trs[0].find_all(["th", "td"])]
        idx = next((i for i, h in enumerate(header) if norm(h) in EMP_NORM), None)
        if idx is None:
            continue
        for tr in trs[1:]:
            cells = [clean(c) for c in tr.find_all(["td", "th"])]
            if len(cells) <= idx or not cells[idx].strip():
                continue
            row = result.setdefault(cells[idx].strip().upper(), {})
            for h, v in zip(header, cells):
                if h:
                    row.setdefault(h, v)
    return result


def parse_reply(msg, sent: dict) -> dict:
    """What did the person ADD to our table? Compared with what we sent:
    - any column that is not one of ours (e.g. 'S/F AMOUNT', 'Terminal Id Status', 'trading a/c')
    - a Mobile No. that differs from the one we sent (we sent it blank when Darwinbox had none)
    Tables are read wherever they are: in a reply, in a forwarded trail, or quoted below."""
    page, plain = message_bodies(msg)
    tables = parse_reply_tables(page) if page else {}
    changes: dict[str, dict] = {}
    for emp, cols in tables.items():
        record = (sent.get(emp) or {}).get("record") or {}
        for head, val in cols.items():
            if not val:
                continue
            n = norm(head)
            if n in MOBILE_NORM:
                if val != (record.get("Mobile No.") or ""):
                    changes.setdefault(emp, {})[head] = val
            elif n not in STANDARD_NORM:
                changes.setdefault(emp, {})[head] = val
    if plain:
        text = cut_quoted_text(plain)
    else:
        text = BeautifulSoup(cut_quoted_html(page), "html.parser").get_text(" ", strip=True)
    text = re.sub(r"\s+", " ", text).strip()
    return {"status": "table" if changes else "text", "rows": changes, "ids": sorted(tables), "text": text[:500]}


def msg_time(msg) -> datetime:
    try:
        dt = parsedate_to_datetime(str(msg.get("Date")))
        return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt
    except (TypeError, ValueError):
        return datetime.now()


def _sent_at(batch: dict) -> datetime:
    return datetime.strptime(batch["sent_at"], "%Y-%m-%d %H:%M:%S")


def assign_batch(state: dict, msg, ids: set, received: datetime):
    """Which of our mails is this a reply to? 1) its Message-ID in In-Reply-To/References,
    2) the same employees, 3) the latest mail sent before the reply (only for an exact 'RE: <subject>')."""
    refs = " ".join(str(msg.get(h, "")) for h in ("In-Reply-To", "References"))
    batches = state["batches"]
    for b in reversed(batches):
        if b.get("message_id") and b["message_id"] in refs:
            return b
    best, best_n = None, 0
    for b in batches:
        if _sent_at(b) <= received + timedelta(minutes=1):
            n = len(ids & set(b["keys"]))
            if n and n >= best_n:
                best, best_n = b, n
    if best:
        return best
    subject = re.sub(r"^(?:(?:re|fw|fwd)\s*:\s*)+", "", str(msg.get("Subject", "")).lower()).strip()
    if subject == SUBJECT.lower():
        earlier = [b for b in batches if _sent_at(b) <= received + timedelta(minutes=1)]
        return earlier[-1] if earlier else None
    return None


def register_reply(state: dict, msg, sender: str, received: datetime, uid: str) -> bool:
    parsed = parse_reply(msg, state["sent"])
    ids = set(parsed["ids"]) or set(re.findall(r"\b[A-Z]\d{6,8}\b", parsed["text"].upper()))
    batch = assign_batch(state, msg, ids, received)
    role = REPLY_SENDERS[sender]
    if batch is None:
        info(f"Mail from {role} ({str(msg.get('Subject', ''))[:50]!r}) does not belong to any mail we sent - ignored")
        return False
    old = batch["replies"].get(sender)
    if old and old["status"] == "table" and parsed["status"] == "text":
        info(f"{role} sent a further message without new data - keeping the earlier reply")
        return False
    batch["replies"][sender] = {"role": role, "received_at": received.strftime("%Y-%m-%d %H:%M:%S"), "uid": uid,
                                "status": parsed["status"], "rows": parsed["rows"], "text": parsed["text"]}
    STATS["replies"] += 1
    if parsed["status"] == "table":
        cols = sorted({column_label(h) if norm(h) not in MOBILE_NORM else "Mobile No." for r in parsed["rows"].values() for h in r})
        ok(f"Reply from {role}: new data for {len(parsed['rows'])} employee(s) in column(s) {', '.join(cols)}"
           + (" (replaces the earlier reply)" if old else ""))
    else:
        warn(f"{role} replied but NO new column/value was found in any table - please read it yourself: "
             f"{parsed['text'][:150]!r}")
    return True


def pop_connect():
    ctx = ssl.create_default_context()
    ctx.set_ciphers("DEFAULT@SECLEVEL=1")
    pop = poplib.POP3_SSL(POP_HOST, POP_PORT, timeout=TIMEOUT, context=ctx)
    pop.user(POP_USER or SMTP_USER)
    pop.pass_(POP_PASSWORD or SMTP_PASSWORD)
    return pop


def batch_open(state: dict, b: dict) -> bool:
    """True while at least one employee of this mail is not submitted yet - only such mails still wait for replies."""
    return any(k in state["sent"] and not state["sent"][k].get("acted") for k in b["keys"])


def collect_replies(state: dict) -> bool:
    """Read the mailbox over POP3 and store the replies. Returns True if the state changed."""
    if not POP_HOST:
        info("POP_HOST is empty - replies are not collected")
        return False
    now = datetime.now()
    active = [b for b in state["batches"] if now - _sent_at(b) <= timedelta(days=REPLY_WATCH_DAYS) and batch_open(state, b)]
    if not active:
        return False
    since = min(_sent_at(b) for b in active) - timedelta(minutes=5)

    pop = None
    for attempt in range(1, POP_ATTEMPTS + 1):
        try:
            pop = pop_connect()
            break
        except poplib.error_proto as exc:                      # wrong login etc. - retrying would not help
            warn(f"POP3 login/protocol problem: {exc} - replies not collected this run")
            return False
        except (OSError, ssl.SSLError) as exc:
            warn(f"POP3 connection attempt {attempt}/{POP_ATTEMPTS} failed ({type(exc).__name__}: {exc})")
            time.sleep(RETRY_DELAY)
    if pop is None:
        warn("POP3 not reachable - will try again on the next run")
        return False

    seen = set(state["pop_seen2"])
    uidls: dict[str, int] = {}
    try:
        for line in pop.uidl()[1]:
            num, uid = line.decode().split(None, 1)
            uidls[uid.strip()] = int(num)
        todo = sorted(((n, uid) for uid, n in uidls.items() if uid not in seen), reverse=True)   # newest first
        info(f"Mailbox: {len(uidls)} message(s), {len(todo)} not checked yet "
             f"(scanning back to {since:%d-%b %H:%M})")
        matches, old_streak, scanned = [], 0, 0
        for pos, (n, uid) in enumerate(todo):
            head = message_from_bytes(b"\r\n".join(pop.top(n, 0)[1]), policy=policy.default)
            scanned += 1
            if scanned % 250 == 0:
                info(f"  ... {scanned} messages scanned")
            received = msg_time(head)
            if received < since:
                old_streak += 1
                seen.add(uid)
                if old_streak >= POP_STOP_AFTER_OLD:      # everything further back is older than our mails
                    seen.update(u for _, u in todo[pos + 1:])
                    break
                continue
            old_streak = 0
            sender = parseaddr(str(head.get("From", "")))[1].lower()
            if sender in REPLY_SENDERS and "employee clearance" in str(head.get("Subject", "")).lower():
                matches.append((n, uid, sender, received))       # remembered as seen only after it is stored
            else:
                seen.add(uid)
        info(f"{scanned} scanned, {len(matches)} look like replies from the {len(REPLY_SENDERS)} people")
        for n, uid, sender, received in sorted(matches):          # oldest first: a later reply replaces an earlier one
            full = message_from_bytes(b"\r\n".join(pop.retr(n)[1]), policy=policy.default)
            register_reply(state, full, sender, received, uid)
            seen.add(uid)
    except (OSError, ssl.SSLError, poplib.error_proto) as exc:
        warn(f"POP3 problem while reading ({type(exc).__name__}: {exc}) - continuing next run")
    finally:
        try:
            pop.quit()
        except Exception:
            pass
        state["pop_seen2"] = sorted(u for u in seen if u in uidls)
    return True


def log_reply_status(state: dict) -> None:
    now = datetime.now()
    for b in state["batches"]:
        if now - _sent_at(b) > timedelta(days=REPLY_WATCH_DAYS) or not batch_open(state, b):
            continue
        missing = [name for addr, name in REPLY_SENDERS.items() if addr not in b["replies"]]
        label = f"Mail of {b['sent_at']} ({len(b['keys'])} employee(s))"
        if not missing:
            if not b.get("done_logged"):
                ok(f"{label}: all {len(REPLY_SENDERS)} replies received")
                b["done_logged"] = True
        else:
            warn(f"{label}: {len(REPLY_SENDERS) - len(missing)}/{len(REPLY_SENDERS)} replies - "
                 f"still waiting for {', '.join(missing)}")


def finish_replies(state: dict, force_excel: bool = False) -> None:
    changed = False
    try:
        changed = collect_replies(state)
    except Exception as exc:                      # reading replies must never break the run
        warn(f"Reply collection failed: {type(exc).__name__}: {exc}")
    log_reply_status(state)
    if state["sent"] or state["batches"]:
        save_state(state)
    if changed or force_excel or (state["sent"] and not EXCEL_FILE.exists()):
        write_excel(state)


# ==========================================================================
# 📧 E-mail
# ==========================================================================
def _fmt(rec: dict, col: str, n: int) -> str:
    if col == "S.No":
        return str(n)
    val = rec[col]
    if isinstance(val, date):
        return f"{val.day}-{val:%b-%y}"
    return str(val or "")


def build_message(records: list[dict]) -> EmailMessage:
    esc = html.escape
    cell = "border:1px solid #000;padding:4px 10px;text-align:center;white-space:nowrap;"
    head = "".join(f'<th style="{cell}background:#B4C6E7;font-weight:bold;">{esc(c)}</th>' for c in COLUMNS)
    body_rows = "".join(
        "<tr>" + "".join(f'<td style="{cell}">{esc(_fmt(r, c, n))}</td>' for c in COLUMNS) + "</tr>"
        for n, r in enumerate(records, start=1))
    paras = "".join(f'<p style="margin:0 0 12px 0;">{esc(g)}<br>{esc(t)}</p>' for g, t in BODY_LINES)
    html_body = (
        '<div style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#000;">'
        f'{paras}'
        f'<table style="border-collapse:collapse;font-size:10.5pt;"><tr>{head}</tr>{body_rows}</table>'
        f'<p style="margin:16px 0 0 0;">{"<br>".join(esc(x) for x in SIGNATURE)}</p></div>')

    widths = [max(len(c), *(len(_fmt(r, c, n)) for n, r in enumerate(records, 1))) for c in COLUMNS]
    line = lambda cells: " | ".join(v.ljust(w) for v, w in zip(cells, widths))
    table_txt = "\n".join([line(COLUMNS), "-+-".join("-" * w for w in widths)] +
                          [line([_fmt(r, c, n) for c in COLUMNS]) for n, r in enumerate(records, 1)])
    text_body = "\n\n".join(f"{g}\n{t}" for g, t in BODY_LINES) + "\n\n" + table_txt + "\n\n" + "\n".join(SIGNATURE) + "\n"

    msg = EmailMessage()
    msg["Subject"] = SUBJECT
    msg["From"] = FROM_ADDR
    msg["To"] = ", ".join(TO)
    if CC:
        msg["Cc"] = ", ".join(CC)
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=(FROM_ADDR.split("@")[-1] or "localhost"))
    msg.set_content(text_body)
    msg.add_alternative(html_body, subtype="html")
    return msg


def send_with_retry(msg: EmailMessage, recipients: list | None = None) -> None:
    """Send the mail. Any network / server problem -> wait RETRY_DELAY seconds and try again, forever.
    The only exception is a wrong SMTP username/password: retrying that every 5 seconds could lock the
    mailbox, so the run stops and nothing is marked as sent (the next run tries again)."""
    custom = recipients
    recipients = custom or (TO + CC)
    attempt = 0
    while True:
        _check_cancel()
        attempt += 1
        try:
            if SMTP_SECURITY == "ssl":
                server = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=TIMEOUT, context=ssl.create_default_context())
            else:
                server = smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=TIMEOUT)
            with server:
                server.ehlo()
                if SMTP_SECURITY == "starttls":
                    server.starttls(context=ssl.create_default_context())
                    server.ehlo()
                if SMTP_USER:
                    server.login(SMTP_USER, SMTP_PASSWORD)
                refused = server.send_message(msg, from_addr=FROM_ADDR, to_addrs=recipients)
            if refused:
                warn(f"Some recipients were refused by the server: {list(refused)}")
            ok(f"E-mail sent to {', '.join(custom or TO)}" + (f" (cc {', '.join(CC)})" if CC and not custom else "")
               + (f"  [attempt {attempt}]" if attempt > 1 else ""))
            return
        except smtplib.SMTPAuthenticationError as exc:
            raise AppError(f"SMTP login was rejected ({exc.smtp_code}). Check SMTP_USER / SMTP_PASSWORD - "
                           "not retrying, to avoid locking the mailbox.")
        except (smtplib.SMTPException, OSError, ssl.SSLError, socket.timeout) as exc:
            warn(f"E-mail attempt {attempt} failed ({type(exc).__name__}: {exc}) - retrying in {RETRY_DELAY}s")
            time.sleep(RETRY_DELAY)


# ==========================================================================
# 🏁 Act step: check the shortfalls -> fill in the "Act" form on Darwinbox
# ==========================================================================
# Who is handled : every e-mailed employee whose replies from ACT_REQUIRE have all arrived and who is still in the Task Box.
# What is checked: Manohar's trading a/c codes and Jagesh's "code - amount" are looked up in RISK_REPORT_<ddmmyyyy>.csv
#                  (today's file, in this folder). Jagesh's amounts must match the report (ACT_TOLERANCE_PCT).
#                  Anything that does not match / cannot be read STOPS the whole step and warns - nothing is submitted.
# What is filled : the 3 Yes/No answers are fixed, the 3 "Outstanding" boxes get the combined shortfall,
#                  Final Remarks gets a short summary of all the replies.
ZERO_TOKENS = {"na", "nil", "nill", "none", "no", "zero", "0", "0.0", "0.00", "-", "--", "not", "applicable"}
CODE_PART = r"(?=[A-Za-z0-9]*[A-Za-z])(?=[A-Za-z0-9]*\d)[A-Za-z0-9]{5,12}"
CODE_RE = re.compile(r"(?<![A-Za-z0-9])" + CODE_PART + r"(?![A-Za-z0-9])")
PAIR_RE = re.compile(r"(?<![A-Za-z0-9])(?P<code>" + CODE_PART + r")\s*[-–—:=]+\s*\(?\s*"
                     r"(?P<amt>-?\s*[\d,]*\.?\d+|zero|nil)\s*\)?", re.I)
ROLE_HINTS = {"Manohar": ("trad", "account", "acct", "code"),
              "Jagesh": ("sf", "shortfall"),
              "Anuj": ("terminal", "dealer", "status")}


def shortfall_of(value: Decimal) -> Decimal:
    """Risk report rule: a NEGATIVE ShortFall is a shortfall (its size counts); a POSITIVE one counts as 0."""
    return -value if value < 0 else Decimal(0)


def fmt_amt(d: Decimal) -> str:
    s = f"{d.quantize(Decimal('0.01')):f}"
    return s[:-3] if s.endswith(".00") else s


def to_dec(text) -> Decimal:
    t = str(text).strip().lower().replace(",", "").replace(" ", "")
    return Decimal(0) if t in ("zero", "nil", "nill") else Decimal(t)


def is_nothing(text: str) -> bool:
    """'NA', 'N/A', 'Nil', 'ZERO', '0', 'NA / NA' ... (blank is NOT nothing - blank means not answered)."""
    t = re.sub(r"n\s*/\s*a", "na", (text or "").lower())
    tokens = [x for x in re.split(r"[\s,;/&]+", t) if x]
    return bool(tokens) and all(x in ZERO_TOKENS for x in tokens)


def pick_cell(cols: dict, role: str):
    """The cell of a reply row that belongs to this person. None = no such column, '' = blank cell."""
    hints = ROLE_HINTS.get(role, ())
    plain = {h: v for h, v in cols.items() if norm(h) not in MOBILE_NORM and norm(h) not in EMP_NORM}
    hit = {h: v for h, v in plain.items() if any(x in norm(h) for x in hints)}
    if not hit:
        other = {h: v for h, v in plain.items() if norm(h) not in STANDARD_NORM}
        hit = other if len(other) == 1 else {}
    if not hit:
        return None
    return " ".join(str(v).strip() for v in hit.values() if v and str(v).strip())


def parse_manohar(text: str) -> tuple[list, str]:
    t = (text or "").strip()
    if is_nothing(t):
        return [], ""
    codes = list(dict.fromkeys(m.group(0).upper() for m in CODE_RE.finditer(t)))
    return (codes, "") if codes else ([], f"cannot read any trading a/c code in Manohar's reply: {t!r}")


def parse_jagesh(text: str) -> tuple[list, bool, str]:
    """-> ([(account, amount he stated)], said 'zero/NA' only, problem)"""
    t = (text or "").strip()
    try:
        pairs = [(m.group("code").upper(), to_dec(m.group("amt"))) for m in PAIR_RE.finditer(t)]
    except InvalidOperation:
        return [], False, f"cannot read the amount in Jagesh's reply: {t!r}"
    if pairs:
        return pairs, False, ""
    if is_nothing(t):
        return [], True, ""
    return [], False, f"cannot read Jagesh's S/F amount: {t!r}"


def risk_report_path() -> Path:
    """Today's RISK_REPORT: the console's COMMON FILES folder first, then this folder. An old one is never used."""
    name = f"RISK_REPORT_{date.today():%d%m%Y}.csv"
    common = (os.getenv("SA_COMMON_DIR") or "").strip()
    for base in ([Path(common)] if common else []) + [HERE]:
        if (base / name).exists():
            return base / name
    return (Path(common) if common else HERE) / name


def read_risk(path: Path, codes: set) -> dict:
    """Streams the (huge) csv once and keeps only the wanted accounts: {code: [ {shortfall, comp, name}, ... ]}"""
    found: dict[str, list] = {c: [] for c in codes}
    if not codes:
        return found
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.reader(fh)
        head = [h.strip().lower() for h in next(reader, [])]
        col = ACT_RISK_COLUMN.lower()
        if "accountid" not in head or col not in head:
            raise AppError(f"{path.name}: columns ACCOUNTID and {ACT_RISK_COLUMN} are needed, found: {head[:12]}...")
        ia, isf = head.index("accountid"), head.index(col)
        icomp = head.index(col + "comp") if (col + "comp") in head else -1
        inm = head.index("name") if "name" in head else -1
        need = max(ia, isf, icomp, inm)
        for row in reader:
            if len(row) > need and row[ia].strip().upper() in found:
                found[row[ia].strip().upper()].append({
                    "shortfall": row[isf].strip(), "comp": row[icomp].strip() if icomp >= 0 else "",
                    "name": row[inm].strip() if inm >= 0 else ""})
    return found


def act_candidates(state: dict, skip: set) -> tuple[list, list]:
    """(ready, waiting): ready = every reply in ACT_REQUIRE has arrived with a value for the employee."""
    ready, waiting = [], []
    for key, v in state["sent"].items():
        if v.get("acted") or key.upper() in skip:
            continue
        manual = {k.upper(): m for k, m in ACT_MANUAL.items()}.get(key.upper())
        if manual is not None:
            cells = {role: str(manual.get(role, "")).strip() for role in ACT_REQUIRE}
            why = [f"ACT_MANUAL has no value for {role}" for role, c in cells.items() if not c]
            warn(f"{v.get('name')} ({key}): using the answers typed in ACT_MANUAL, not the table cells")
            (waiting if why else ready).append({"key": key, "v": v, "cells": cells, "why": why})
            continue
        batches = [b for b in state["batches"] if key in b.get("keys", [])]
        batch = batches[-1] if batches else None
        cells, why = {}, []
        for role in ACT_REQUIRE:
            row = None
            for r in (batch or {}).get("replies", {}).values():
                if r.get("role") == role:
                    row = {k.upper(): c for k, c in (r.get("rows") or {}).items()}.get(key.upper())
            if row is None:
                why.append(f"{role} has not replied for this employee yet")
                continue
            cell = pick_cell(row, role)
            if cell is None:
                why.append(f"{role}'s reply has no usable column")
            elif not cell:
                why.append(f"{role} left the cell blank")
            else:
                cells[role] = cell
        (waiting if why else ready).append({"key": key, "v": v, "cells": cells, "why": why})
    return ready, waiting


def lookup(risk: dict, code: str, problems: list):
    rows = risk.get(code) or []
    if not rows:
        problems.append(f"a/c {code} is not in the risk report")
        return None
    if len(rows) > 1:
        problems.append(f"a/c {code} appears {len(rows)} times in the risk report")
        return None
    try:
        return Decimal(rows[0]["shortfall"].replace(",", "") or "0"), rows[0]
    except InvalidOperation:
        problems.append(f"a/c {code}: risk report {ACT_RISK_COLUMN} {rows[0]['shortfall']!r} is not a number")
        return None


def make_plan(cand: dict, risk: dict) -> dict:
    cells = cand["cells"]
    plan = {"emp": cand["key"], "name": cand["v"].get("name", ""), "status": "READY", "problems": [], "notes": [],
            "m_text": cells["Manohar"], "j_text": cells["Jagesh"], "a_text": cells["Anuj"],
            "m_accts": [], "j_accts": [], "j_zero": False}
    problems = plan["problems"]
    codes_m, p1 = parse_manohar(cells["Manohar"])
    pairs, plan["j_zero"], p2 = parse_jagesh(cells["Jagesh"])
    problems += [p for p in (p1, p2) if p]
    accounts: dict[str, Decimal] = {}
    for code in codes_m:
        got = lookup(risk, code, problems)
        if got:
            val, row = got
            amt = shortfall_of(val)
            accounts[code] = amt
            plan["m_accts"].append((code, amt, row["shortfall"], row["comp"]))
            if val > 0:
                plan["notes"].append(f"{code}: risk report {ACT_RISK_COLUMN} is +{row['shortfall']} (positive) -> counted as 0")
    tol = Decimal(str(ACT_TOLERANCE_PCT)) / 100
    for code, said in pairs:
        got = lookup(risk, code, problems)
        if not got:
            continue
        val, row = got
        amt = shortfall_of(val)
        diff, allowed = abs(abs(said) - amt), max(Decimal("0.01"), amt * tol)
        if diff > allowed:
            problems.append(f"Jagesh says a/c {code} = {fmt_amt(said)} but the risk report {ACT_RISK_COLUMN} is "
                            f"{fmt_amt(val)} (shortfall {fmt_amt(amt)}, difference {fmt_amt(diff)})")
        plan["j_accts"].append((code, said, val, amt))
        if code in accounts:
            plan["notes"].append(f"a/c {code} is in both replies - counted once")
        else:
            accounts[code] = amt
    plan["total"] = sum(accounts.values(), Decimal(0))
    plan["total_str"] = fmt_amt(plan["total"])
    plan["remarks"] = build_remarks(plan)
    if problems:
        plan["status"] = "BLOCKED"
    return plan


def build_remarks(plan: dict, limit: int = 200) -> str:
    m = ", ".join(f"{c} ({fmt_amt(a)})" for c, a, _, _ in plan["m_accts"]) or "NA"
    if plan["j_accts"]:
        j = ", ".join(f"{c} {fmt_amt(a)}" for c, _, _, a in plan["j_accts"])
    else:
        j = "Zero" if re.sub(r"[^a-z0-9]", "", plan["j_text"].lower()) in ("zero", "0", "000", "nil") else "NA"
    a = " ".join(plan["a_text"].split())
    tail = f", Total shortfall {plan['total_str']}" if plan["total"] != 0 else ""
    text = ""
    for sep in (" as received from ", " per "):
        text = f"Trading A/c - {m}{sep}Manohar, S/F amount {j}{sep}Jagesh, {a}{sep}Anuj{tail}"
        if len(text) <= limit:
            return text
    return text[:limit]


# ---- the form itself -----------------------------------------------------------------------------------
def label_of(q: dict) -> str:
    cfg = q.get("config") or {}
    return " ".join(str(cfg.get("question") or cfg.get("questionText") or q.get("name") or "").split())


def get_form_definition(client, rec: dict, params: dict) -> list:
    body = {"form_id": params.get("form_id") or DEFAULT_FORM_ID,
            "custom_workflow_id": params.get("custom_workflow_id") or rec["_sep_flow"],
            "type": params.get("type") or "separation_workflow",
            "dynamic_config": params.get("dynamic_config") or '{"type":"separation_workflow"}',
            "token": None, "isOldForm": False, "user_id": rec["_user_id"], "user_language": None,
            CSRF_COOKIE: csrf_token(client)}
    data = _json(client.request("POST", FORM_DETAIL_URL, json=body, headers=_ajax_headers(client, True)),
                 "form definition")
    questions = (data.get("form_details") or {}).get("questions") or data.get("questions")
    if not questions:
        raise AppError(f"form definition has no questions ({data.get('message')})")
    return sorted(questions, key=lambda q: q.get("order", 0))


def build_answers(questions: list, plan: dict, details: dict) -> dict:
    meta = details.get("meta_data") or {}
    out: dict = {}
    n_choice = n_total = n_remarks = 0
    for q in questions:
        cfg, qtype = q.get("config") or {}, q.get("question_type") or (q.get("config") or {}).get("questionType")
        label = label_of(q)
        low, order, mand = label.lower(), q.get("order", 0), cfg.get("isMandatory", True)
        if qtype == "choices":
            want = next((ans for pre, ans in ACT_FIXED_CHOICES.items() if low.startswith(pre)), None)
            if want is None:
                raise AppError(f"unexpected question on the form: {label!r} - not submitting")
            choice = next((c for c in cfg.get("choices") or [] if str(c.get("value", "")).strip().lower() == want.lower()), None)
            if choice is None:
                raise AppError(f"question {label!r} has no '{want}' choice")
            item = {"answer": [choice["id"]], "numericValue": None, "comment": None, "otherValue": None,
                    "type": "choices", "uploadedFiles": [], "hidden": False, "order": order}
            n_choice += 1
        elif qtype == "text-box":
            if low in ACT_TOTAL_QUESTIONS:
                value, n_total = plan["total_str"], n_total + 1
            elif low == ACT_REMARKS_QUESTION:
                value, n_remarks = plan["remarks"], n_remarks + 1
            else:
                raise AppError(f"unexpected question on the form: {label!r} - not submitting")
            limit = cfg.get("maximumCharacters")
            if limit and len(value) > int(limit):
                raise AppError(f"{label!r}: {len(value)} characters, the form allows {limit}")
            item = {"answer": value, "comment": None, "type": "text-box", "defaultValue": None,
                    "hidden": False, "order": order}
        elif qtype == "system-attributes":
            field = cfg.get("field") or "bank_pan_num"
            val = (details.get("fields_value") or {}).get(field) or (details.get("field_data_values") or {}).get(field) or ""
            m = meta.get(field) or {}
            out[q["custom_id"]] = {
                "answer": str(val), "comment": None, "label": cfg.get("question") or m.get("label"),
                "fieldType": cfg.get("fieldType") or m.get("type") or "textfield", "validations": m.get("validations"),
                "regex": None, "section": cfg.get("array") or m.get("section"), "mandatory": cfg.get("isMandatory", False),
                "type": "system-attributes", "fieldName": field, "isEditable": cfg.get("isEditAllowed", False),
                "updateInProfile": cfg.get("updateInProfile", False), "is_separation_field": cfg.get("isSeparationField", False),
                "defaultValue": None, "isCustom": m.get("isCustom", False), "hidden": False, "order": order,
                "smartTag": None, "isMandatory": mand}
            continue
        else:
            raise AppError(f"unexpected question type {qtype!r} ({label!r}) - not submitting")
        if "smartTag" in cfg:
            item["smartTag"] = None
        item["isMandatory"] = mand
        out[q["custom_id"]] = item
    if n_choice != len(ACT_FIXED_CHOICES) or n_total != len(ACT_TOTAL_QUESTIONS) or n_remarks != 1:
        raise AppError("the form no longer has the expected questions "
                       f"(Yes/No {n_choice}, outstanding {n_total}, remarks {n_remarks}) - not submitting")
    out["allSmartTags"] = {"answer": [], "hidden": True, "type": "smart_tags"}
    return out


def build_submission(client, rec: dict, plan: dict) -> dict:
    """Does what the Act button does and builds - but does not send - the final form post."""
    reply = open_task_form(client, rec)
    if str(reply.get("status", "")).lower() != "success":
        raise AppError(f"Act form refused: {reply.get('message') or reply.get('status')}")
    params = form_params(reply)
    soup = BeautifulSoup(str(reply.get("update") or ""), "html.parser")
    form = soup.find("form", id="edit_com_bus_form") or soup
    hidden = []
    for inp in form.find_all("input"):
        name, typ = inp.get("name"), (inp.get("type") or "text").lower()
        if name and typ not in ("submit", "button", "file", "image", "reset", "checkbox", "radio"):
            hidden.append((name, inp.get("value") or ""))
    names = [n for n, _ in hidden]
    for need in ("phase_id", "SeprationEmployeeWorkFLow[id]", "FormFramework[id]", "NewFormInput"):
        if need not in names:
            raise AppError(f"the Act form has no '{need}' field - layout changed, not submitting")
    answers = build_answers(get_form_definition(client, rec, params), plan, get_details(client, rec, params))
    answers_json = json.dumps(answers, separators=(",", ":"), ensure_ascii=False)
    fields = []
    for name, val in hidden:
        fields.append((name, answers_json if name == "NewFormInput" else ("true" if name == "NewFormInputValid" else val)))
    if CSRF_COOKIE not in names:
        fields.insert(0, (CSRF_COOKIE, unquote(csrf_token(client))))
    return {"fields": fields, "answers": answers}


def send_submission(client, sub: dict) -> tuple[bool, str]:
    """ONE attempt, never retried blindly: a repeated POST could submit twice."""
    files = [(n, (None, v)) for n, v in sub["fields"]]
    files.append(("AddAttachmentWorkFlow[name]", ("", b"", "application/octet-stream")))
    headers = {"Accept": "application/json, text/javascript, */*; q=0.01", "Origin": BASE_URL,
               "Referer": TASKS_URL, "X-Requested-With": "XMLHttpRequest"}
    resp = client.session.request("POST", COMPLETE_ACTION_URL, files=files, headers=headers,
                                  timeout=TIMEOUT, allow_redirects=False)
    if bounced_to_login(resp):
        raise AppError("Darwinbox rejected the session - nothing was submitted")
    data = _json(resp, "submit")
    return str(data.get("status", "")).lower() == "success", str(data.get("message") or data)


# ---- preview files + alert ---------------------------------------------------------------------------
def _risk_text(items) -> str:
    return "; ".join(items) if items else "-"


def write_act_preview(rows: list) -> None:
    """Darwin_Act_Preview.xlsx (readable) + Darwin_Act_Preview.json (the exact fields that are / would be posted)."""
    payloads = []
    for r in rows:
        sub = r.get("sub")
        if sub:
            payloads.append({"emp_id": r["emp"], "name": r["name"], "status": r["status"], "url": COMPLETE_ACTION_URL,
                             "fields": [{"name": n, "value": ("<csrf token>" if n == CSRF_COOKIE else v)} for n, v in sub["fields"]],
                             "NewFormInput": sub["answers"]})
    try:
        tmp = ACT_PREVIEW_JSON.with_suffix(".tmp")
        tmp.write_text(json.dumps(payloads, indent=2, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, ACT_PREVIEW_JSON)
    except OSError as exc:
        warn(f"Could not write {ACT_PREVIEW_JSON.name} ({exc})")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        warn("openpyxl is not installed - Darwin_Act_Preview.xlsx not written")
        return
    headers = ["S.No", "Emp ID", "Name", "Status", "Manohar - trading a/c", "Risk report -> counted (Manohar a/c)",
               "Jagesh - S/F amount", "Risk report -> counted (Jagesh a/c)", "Combined shortfall", "Closure of Master Care",
               "Removal Savior/RMS", "Removal Master Dealer", "Final remarks", "Notes / problems"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Act Preview"
    ws.append(headers)
    yes_no = list(ACT_FIXED_CHOICES.values())
    for n, r in enumerate(rows, start=1):
        plan = r.get("plan")
        if plan:
            ws.append([n, r["emp"], r["name"], r["status"], plan["m_text"],
                       _risk_text([f"{c} = {raw} -> {fmt_amt(a)}" for c, a, raw, _ in plan["m_accts"]]),
                       plan["j_text"], _risk_text([f"{c} = {fmt_amt(v)} -> {fmt_amt(a)}" for c, _, v, a in plan["j_accts"]]),
                       plan["total_str"], *yes_no, plan["remarks"],
                       " | ".join(plan["problems"] + plan["notes"] + ([r["error"]] if r.get("error") else []))])
        else:
            ws.append([n, r["emp"], r["name"], r["status"], *[None] * 8, r.get("error") or ""])
    thin = Side(style="thin", color="000000")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=(cell.column >= 13))
            cell.font = Font(name="Calibri", size=11, bold=(cell.row == 1))
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="B4C6E7")
            elif headers[cell.column - 1] in ("Emp ID", "Combined shortfall"):
                cell.number_format = "@"
    for idx in range(1, len(headers) + 1):
        longest = max(len(str(c.value)) for c in ws[get_column_letter(idx)] if c.value is not None)
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 4, 70)
    ws.freeze_panes = "A2"
    tmp = ACT_PREVIEW_XLSX.with_name(ACT_PREVIEW_XLSX.stem + ".tmp.xlsx")
    try:
        wb.save(tmp)
        os.replace(tmp, ACT_PREVIEW_XLSX)
    except OSError as exc:
        warn(f"Could not write {ACT_PREVIEW_XLSX.name} ({exc}). Close it if it is open in Excel.")
        tmp.unlink(missing_ok=True)


def act_blocked(state: dict, problems: list, submit: bool) -> None:
    text = ("The Darwinbox Act step was STOPPED - nothing was submitted.\n\n" + "\n".join(f"- {p}" for p in problems)
            + "\n\nFix the reply / risk report (or add the Emp ID to ACT_SKIP_EMP_IDS) and run again.\n")
    ACT_BLOCKED_FILE.write_text(text, encoding="utf-8")
    STATS["blocked"] = True
    fail("ACT STEP STOPPED - " + str(len(problems)) + " problem(s), see " + ACT_BLOCKED_FILE.name)
    for p in problems:
        fail("   " + p)
    if submit and ACT_ALERT_TO and state.get("act_alert") != text:
        msg = EmailMessage()
        msg["Subject"] = "Darwinbox Act step stopped - please check"
        msg["From"], msg["To"] = FROM_ADDR, ", ".join(ACT_ALERT_TO)
        msg["Date"] = formatdate(localtime=True)
        msg["Message-ID"] = make_msgid(domain=(FROM_ADDR.split("@")[-1] or "localhost"))
        msg.set_content(text)
        send_with_retry(msg, recipients=ACT_ALERT_TO)
        state["act_alert"] = text
        save_state(state)


def act_step(client, state: dict, fresh: bool = False) -> int:
    submit = ACT_SUBMIT or "--submit" in ARGS
    log("🏁", "Act step - " + ("LIVE: forms WILL be submitted" if submit else "TEST MODE: nothing is submitted"), "b")
    skip = {s.upper() for s in ACT_SKIP_EMP_IDS}
    ready, waiting = act_candidates(state, skip)
    rows: list[dict] = [{"emp": w["key"], "name": w["v"].get("name", ""), "status": "WAITING",
                         "error": "; ".join(w["why"])} for w in waiting]
    for w in waiting:
        info(f"Waiting: {w['v'].get('name')} ({w['key']}) - {'; '.join(w['why'])}")
    if not ready:
        ok("Act step: nobody has all replies yet" if waiting else "Act step: nothing to do")
        if rows:
            write_act_preview(rows)
        return 0
    risk_file = risk_report_path()
    need_risk = any(parse_manohar(c["cells"]["Manohar"])[0] or parse_jagesh(c["cells"]["Jagesh"])[0] for c in ready)
    if need_risk and not risk_file.exists():          # no account to check (all NA / zero) = no risk report needed
        warn(f"{risk_file.name} not found in {HERE} - Act step skipped (an old risk report is never used)")
        return 1
    task_rows = extract_rows(client.get_tasks_html(fresh=fresh))
    in_box = {r["_key"].upper(): r for r in (parse_row(x) for x in task_rows) if r["_key"]}
    cands = []
    for c in ready:
        if c["key"].upper() in in_box:
            cands.append(c)
        else:
            rows.append({"emp": c["key"], "name": c["v"].get("name", ""), "status": "NOT IN TASK BOX",
                         "error": "all replies are in, but the task is not in the Task Box (already done?)"})
            info(f"{c['v'].get('name')} ({c['key']}): not in the Task Box - skipped")
    if not cands:
        write_act_preview(rows)
        ok("Act step: no ready employee is in the Task Box")
        return 0
    wanted = set()
    for c in cands:
        wanted.update(parse_manohar(c["cells"]["Manohar"])[0])
        wanted.update(code for code, _ in parse_jagesh(c["cells"]["Jagesh"])[0])
    if wanted:
        info(f"Reading {risk_file.name} for {len(wanted)} account(s)")
    else:
        info("No trading a/c to check - the risk report is not needed")
    risk = read_risk(risk_file, wanted) if wanted else {}
    plans = [make_plan(c, risk) for c in cands]
    problems = [f"{p['name']} ({p['emp']}): {x}" for p in plans for x in p["problems"]]
    for p in plans:
        (fail if p["status"] == "BLOCKED" else ok)(
            f"{p['name']} ({p['emp']}): total {p['total_str']}  |  {p['remarks']}")
        for note in p["notes"]:
            warn(f"   {note}")
    entries = [{"emp": p["emp"], "name": p["name"], "status": p["status"], "plan": p} for p in plans]
    if problems and submit:                               # stop BEFORE anything is opened or sent
        write_act_preview(rows + entries)
        act_blocked(state, problems, submit)
        return 3
    rc = 0
    for i, e in enumerate(entries, start=1):              # build the exact post (read-only calls) for the preview
        if e["status"] != "READY":
            continue
        try:
            e["sub"] = build_submission(client, in_box[e["emp"].upper()], e["plan"])
        except (AppError, requests.RequestException) as exc:
            e["status"], e["error"], rc = "ERROR", str(exc), 1
            fail(f"{e['name']} ({e['emp']}): could not build the form: {exc}")
        time.sleep(CALL_DELAY)
    if problems:                                          # test mode only: show, do not pretend all is fine
        write_act_preview(rows + entries)
        act_blocked(state, problems, False)
        return 3
    ACT_BLOCKED_FILE.unlink(missing_ok=True)
    if not submit:
        write_act_preview(rows + entries)
        ok(f"TEST MODE: nothing submitted. Check {ACT_PREVIEW_XLSX.name} and {ACT_PREVIEW_JSON.name}; "
           "set ACT_SUBMIT = True to go live.")
        return rc
    for e in entries:
        if e["status"] != "READY":
            continue
        try:
            done, message = send_submission(client, e["sub"])
        except requests.RequestException as exc:
            e["status"], e["error"], rc = "UNKNOWN", f"{type(exc).__name__} - outcome unknown", 1
            warn(f"{e['name']} ({e['emp']}): no answer ({type(exc).__name__}). Outcome unknown - the Task Box is "
                 "checked again on the next run, so it is never submitted twice.")
            continue
        except AppError as exc:
            e["status"], e["error"], rc = "ERROR", str(exc), 1
            fail(f"{e['name']} ({e['emp']}): {exc}")
            continue
        if done:
            e["status"] = "SUBMITTED"
            state["sent"][e["emp"]]["acted"] = {"at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                                "total": e["plan"]["total_str"], "remarks": e["plan"]["remarks"]}
            save_state(state)                              # remembered at once - a crash cannot cause a second submit
            STATS["submitted"] += 1
            ok(f"SUBMITTED {e['name']} ({e['emp']}) - total {e['plan']['total_str']} - Darwinbox: {message}")
        else:
            e["status"], e["error"], rc = "REFUSED", message, 1
            fail(f"{e['name']} ({e['emp']}): Darwinbox did not accept it: {message}")
        time.sleep(CALL_DELAY)
    write_act_preview(rows + entries)
    return rc


def act_step_safe(client, state: dict, fresh: bool = False) -> int:
    try:
        return act_step(client, state, fresh)
    except AppError as exc:
        fail(f"Act step: {exc}")
    except requests.RequestException as exc:
        fail(f"Act step - network error: {type(exc).__name__}: {exc}")
    except Exception as exc:                              # the Act step must never break the mail / reply work
        fail(f"Act step crashed: {type(exc).__name__}: {exc}")
    return 1


# ==========================================================================
# ⏰ Reply-overdue alert  (picture -> e-mail + WhatsApp)
# ==========================================================================
# A mail we sent is "overdue" when ALERT_AFTER_HOURS have passed, somebody in ALERT_ROLES has still not replied and
# at least one of its employees is not submitted yet. Then ONE picture is drawn (who has / has not replied, which
# employees are waiting) and sent by e-mail (attached) and by WhatsApp. Sent once per mail - and once per person, so a
# failed number is simply tried again on the next run without anybody getting it twice. State: batch["alert"].
ALERT_ROLES = ACT_REQUIRE


def _alert_font(size: int, bold: bool = False):
    from PIL import ImageFont
    names = (["arialbd.ttf", "segoeuib.ttf", "DejaVuSans-Bold.ttf"] if bold
             else ["arial.ttf", "segoeui.ttf", "DejaVuSans.ttf"])
    for n in names:
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_alert_image(waiting: list, replied: list, employees: list, sent_txt: str, hours: int, out: Path) -> None:
    """waiting = [role, ...]   replied = [(role, 'time'), ...]   employees = [(name, emp id), ...]  -> JPEG at `out`"""
    from PIL import Image, ImageDraw
    W = 1080
    NAVY, RED, AMBER, GREEN = (14, 30, 58), (214, 40, 57), (245, 158, 11), (22, 163, 74)
    INK, MUTE, LINE, PAPER, CARD = (20, 28, 45), (100, 112, 132), (226, 231, 240), (244, 246, 250), (255, 255, 255)
    people = [(r, None) for r in waiting] + list(replied)
    H = 572 + 124 * len(people) + 20 + 50 + 56 * len(employees) + 24 + 50 + 150
    im = Image.new("RGB", (W, H), PAPER)
    d = ImageDraw.Draw(im)

    def fit(text, size, bold, maxw):
        while size > 18 and d.textlength(text, font=_alert_font(size, bold)) > maxw:
            size -= 2
        return _alert_font(size, bold)

    d.rectangle([0, 0, W, 250], fill=NAVY)
    d.rectangle([0, 250, W, 262], fill=RED)
    d.text((60, 52), f"MASTERTRUST  |  ALERT FOR {ALERT_FOR.upper()}", font=fit(f"MASTERTRUST  |  ALERT FOR {ALERT_FOR.upper()}", 26, True, W - 120), fill=(150, 170, 205))
    d.text((60, 100), "EMPLOYEE CLEARANCE", font=_alert_font(62, True), fill="white")
    d.text((60, 178), "Task e-mailer  -  reply overdue", font=_alert_font(34), fill=(255, 190, 196))

    y = 300
    d.rounded_rectangle([60, y, W - 60, y + 190], 28, fill=(255, 238, 240), outline=RED, width=4)
    d.ellipse([92, y + 45, 192, y + 145], fill=RED)
    d.text((142, y + 95), "!", font=_alert_font(78, True), fill="white", anchor="mm")
    d.text((222, y + 38), "YOUR CLEARANCE E-MAIL IS NOT REPLIED BY",
           font=fit("YOUR CLEARANCE E-MAIL IS NOT REPLIED BY", 30, True, W - 60 - 222 - 24), fill=RED)
    names = "  &  ".join(waiting)
    d.text((222, y + 88), names, font=fit(names, 52, True, W - 60 - 222 - 24), fill=INK)
    d.text((222, y + 148), f"{hours} hours since the e-mail was sent (limit {ALERT_AFTER_HOURS} h)",
           font=_alert_font(26), fill=MUTE)

    y = 520
    d.text((60, y), "REPLY STATUS", font=_alert_font(26, True), fill=MUTE)
    y += 52
    for role, when in people:
        done = when is not None
        col = GREEN if done else RED
        d.rounded_rectangle([60, y, W - 60, y + 104], 22, fill=CARD, outline=LINE, width=2)
        d.rounded_rectangle([60, y, 76, y + 104], 8, fill=col)
        d.ellipse([104, y + 24, 160, y + 80], fill=col)
        if done:
            d.line([(119, y + 53), (129, y + 63), (147, y + 41)], fill="white", width=7, joint="curve")
        else:
            d.line([(120, y + 42), (144, y + 62)], fill="white", width=7)
            d.line([(144, y + 42), (120, y + 62)], fill="white", width=7)
        d.text((190, y + 20), role, font=_alert_font(36, True), fill=INK)
        d.text((190, y + 62), ("Replied  " + when) if done else "NOT REPLIED", font=_alert_font(24, not done), fill=col)
        pill = "DONE" if done else "PENDING"
        pw = d.textlength(pill, font=_alert_font(24, True)) + 44
        d.rounded_rectangle([W - 100 - pw, y + 30, W - 100, y + 74], 22, fill=col)
        d.text((W - 100 - pw / 2, y + 52), pill, font=_alert_font(24, True), fill="white", anchor="mm")
        y += 124

    y += 20
    d.text((60, y), f"EMPLOYEES WAITING FOR CLEARANCE  ({len(employees)})", font=_alert_font(26, True), fill=MUTE)
    y += 50
    d.rounded_rectangle([60, y, W - 60, y + 56 * len(employees) + 24], 22, fill=CARD, outline=LINE, width=2)
    for i, (nm, code) in enumerate(employees):
        yy = y + 14 + 56 * i
        d.ellipse([92, yy + 12, 116, yy + 36], fill=AMBER)
        d.text((140, yy + 24), nm, font=fit(nm, 30, True, 620), fill=INK, anchor="lm")
        d.text((W - 92, yy + 24), code, font=_alert_font(26), fill=MUTE, anchor="rm")

    d.rectangle([0, H - 150, W, H], fill=NAVY)
    d.text((60, H - 118), "Please follow up with the people marked PENDING.", font=_alert_font(30, True), fill="white")
    d.text((60, H - 70), f"Mail sent {sent_txt}   |   Auto-generated alert", font=_alert_font(24), fill=(150, 170, 205))
    im.save(out, quality=92)


def overdue_batches(state: dict, force: bool = False) -> list:
    """Batches that need an alert: [{batch, waiting[], replied[], employees[], hours, sent_txt}]"""
    now, out = datetime.now(), []
    for b in state["batches"]:
        if (b.get("alert") or {}).get("done") and not force:
            continue
        age = now - _sent_at(b)
        if age > timedelta(days=REPLY_WATCH_DAYS) or (age < timedelta(hours=ALERT_AFTER_HOURS) and not force):
            continue
        emps = [(state["sent"][k].get("name", k), state["sent"][k].get("emp_code") or k)
                for k in b["keys"] if k in state["sent"] and not state["sent"][k].get("acted")]
        got = {r.get("role"): r.get("received_at", "") for r in b["replies"].values()}
        waiting = [r for r in ALERT_ROLES if r not in got]
        if not emps or not waiting:
            continue
        replied = [(r, datetime.strptime(got[r], "%Y-%m-%d %H:%M:%S").strftime("%d-%b %H:%M")) for r in ALERT_ROLES if r in got]
        out.append({"batch": b, "waiting": waiting, "replied": replied, "employees": emps,
                    "hours": int(age.total_seconds() // 3600), "sent_txt": _sent_at(b).strftime("%d-%b-%Y %H:%M")})
    return out


def send_alert_email(item: dict, image: Path) -> None:
    names = " and ".join(item["waiting"])
    msg = EmailMessage()
    msg["Subject"] = f"Reminder: Employee Clearance - no reply from {', '.join(item['waiting'])}"
    msg["From"], msg["To"] = FROM_ADDR, ", ".join(ALERT_EMAIL_TO)
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=(FROM_ADDR.split("@")[-1] or "localhost"))
    lines = "\n".join(f"  - {n} ({c})" for n, c in item["employees"])
    msg.set_content(f"Dear {ALERT_FOR},\n\nYour Employee Clearance task e-mail sent on {item['sent_txt']} has had no reply "
                    f"from {names} for {item['hours']} hours.\n\nEmployees waiting:\n{lines}\n\n"
                    "Please follow up with them. The picture is attached.\n\nRegards,\nEmployee Clearance task e-mailer\n")
    msg.add_attachment(image.read_bytes(), maintype="image", subtype="jpeg", filename=image.name)
    send_with_retry(msg, recipients=ALERT_EMAIL_TO)


def wa_upload(image: Path) -> str:
    """Uploads the picture once; returns the media id. Raises AppError when it did not work."""
    for attempt in range(1, WA_ATTEMPTS + 1):
        try:
            with open(image, "rb") as fh:
                r = requests.post(WA_UPLOAD_URL, headers={"Authorization": f"Bearer {WA_AUTH_TOKEN}", "accept": "application/json"},
                                  files={"file": ("snapshot.jpg", fh, "image/jpeg")}, timeout=TIMEOUT)
        except requests.RequestException as exc:
            warn(f"WhatsApp upload attempt {attempt} failed ({type(exc).__name__}) - retrying in {RETRY_DELAY}s")
            time.sleep(RETRY_DELAY)
            continue
        if r.status_code in (200, 201):
            body = r.json()
            media = body.get("id") or body.get("mediaHandle")
            if media:
                return media
            raise AppError(f"WhatsApp upload gave no media id: {r.text[:200]}")
        if r.status_code < 500:
            raise AppError(f"WhatsApp upload refused: {r.status_code} {r.text[:200]}")
        time.sleep(RETRY_DELAY)
    raise AppError("WhatsApp upload failed (network / server) - will be tried again on the next run")


def wa_send(media_id: str, number: str, text1: str, text2: str) -> bool:
    payload = {"to": number, "type": "template", "source": "external",
               "template": {"name": WA_TEMPLATE, "language": {"code": WA_LANGUAGE},
                            "components": [{"type": "header", "parameters": [{"type": "image", "image": {"id": media_id}}]},
                                           {"type": "body", "parameters": [{"type": "text", "text": text1},
                                                                          {"type": "text", "text": text2}]}]}}
    for attempt in range(1, WA_ATTEMPTS + 1):
        try:
            r = requests.post(WA_SEND_URL, headers={"Authorization": f"Bearer {WA_AUTH_TOKEN}", "Content-Type": "application/json"},
                              json=payload, timeout=TIMEOUT)
        except requests.RequestException as exc:
            warn(f"WhatsApp to {number}: attempt {attempt} failed ({type(exc).__name__}) - retrying in {RETRY_DELAY}s")
            time.sleep(RETRY_DELAY)
            continue
        if r.status_code == 200:
            return True
        fail(f"WhatsApp to {number} failed: {r.status_code} {r.text[:200]}")
        if r.status_code < 500:
            return False
        time.sleep(RETRY_DELAY)
    return False


def alert_step(state: dict, preview: bool = False, now: bool = False) -> None:
    """preview = draw the picture(s) only, send nothing, change nothing.  now = ignore the 24 h, send, change nothing."""
    items = overdue_batches(state, force=preview or now)
    if not items:
        info(f"Alert: no mail is overdue (limit {ALERT_AFTER_HOURS} h)")
        return
    try:
        import PIL  # noqa: F401
    except ImportError:
        warn("Alert needs Pillow for the picture:  pip install pillow   - alert skipped")
        return
    for it in items:
        b = it["batch"]
        st = b.setdefault("alert", {"email": False, "wa": {}, "done": False})
        image = HERE / f"Reply_Overdue_Alert_{_sent_at(b):%Y%m%d_%H%M}.jpg"
        draw_alert_image(it["waiting"], it["replied"], it["employees"], it["sent_txt"], it["hours"], image)
        log("⏰", f"Mail of {it['sent_txt']}: no reply from {', '.join(it['waiting'])} after {it['hours']} h  -> {image.name}", "y")
        if preview:
            ok(f"PREVIEW only - picture saved, nothing sent: {image}")
            continue
        todo_mail = ALERT_EMAIL_TO and not st["email"]
        todo_wa = [n for n in WA_TO if WA_ENABLED and not st["wa"].get(n)]
        if not ALERT_EMAIL_TO and not (WA_ENABLED and WA_TO):
            warn("Alert: no recipient set (ALERT_EMAIL_TO / WA_TO are empty) - picture saved only")
            continue
        if todo_mail:
            try:
                send_alert_email(it, image)
                st["email"] = True
            except AppError as exc:
                fail(str(exc))
        if todo_wa:
            try:
                media = wa_upload(image)
                for n in todo_wa:
                    if wa_send(media, n, WA_TEXT1, datetime.now().strftime("%d-%m-%Y")):
                        st["wa"][n] = True
                        STATS["alerts"] += 1
                        ok(f"WhatsApp alert sent to {n}")
            except AppError as exc:
                fail(str(exc))
        if not now:
            st["done"] = (not ALERT_EMAIL_TO or st["email"]) and all(st["wa"].get(n) for n in WA_TO if WA_ENABLED)
            st["at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            save_state(state)
        elif "alert" in b and not st.get("at"):
            b.pop("alert")                 # --alert-now leaves no trace: the real 24 h alert still goes out


def alert_step_safe(state: dict, **kw) -> None:
    try:
        alert_step(state, **kw)
    except Exception as exc:               # an alert problem must never break the run
        warn(f"Alert step failed: {type(exc).__name__}: {exc}")


# ==========================================================================
# 🔔 Gentle reminder (after REMINDER_AFTER_HOURS, on the thread of the original mail, to the person who has not replied)
# ==========================================================================
def _records_of(state: dict, b: dict) -> list:
    out = []
    for k in b["keys"]:
        rec = dict((state["sent"].get(k) or {}).get("record") or {})
        for c in COLUMNS:
            rec.setdefault(c, "")
        sd = rec.get("Separation date")
        if isinstance(sd, str):
            try:
                rec["Separation date"] = date.fromisoformat(sd)
            except ValueError:
                pass
        out.append(rec)
    return out


def build_reminder(state: dict, b: dict, addr: str, role: str) -> EmailMessage:
    esc = html.escape
    original = build_message(_records_of(state, b))
    o_html = original.get_body(("html",)).get_content()
    o_text = original.get_body(("plain",)).get_content()
    ask = " ".join(t for g, t in BODY_LINES if role.lower() in g.lower())
    sent_txt = _sent_at(b).strftime("%d %B %Y, %H:%M")
    cc = [c for c in REMINDER_CC if c.lower() != addr.lower()]
    intro = f"This is a gentle reminder - we have not yet received your reply to the mail below (sent on {sent_txt})."
    close = "Kindly share it at the earliest. Thank you."
    sig = list(SIGNATURE)
    head_lines = [("From", FROM_ADDR), ("Sent", sent_txt), ("To", ", ".join(TO)), ("Cc", ", ".join(CC)), ("Subject", SUBJECT)]
    text = (f"Dear {role},\n\n{intro} {ask}\n\n{close}\n\n" + "\n".join(sig) + "\n\n"
            "-----Original Message-----\n" + "\n".join(f"{k}: {v}" for k, v in head_lines if v) + "\n\n" + o_text)
    page = ('<div style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#000;">'
            f'<p style="margin:0 0 12px 0;">Dear {esc(role)},</p>'
            f'<p style="margin:0 0 12px 0;">{esc(intro)} {esc(ask)}</p>'
            f'<p style="margin:0 0 12px 0;">{esc(close)}</p>'
            f'<p style="margin:0 0 16px 0;">{"<br>".join(esc(x) for x in sig)}</p>'
            '<hr style="border:0;border-top:1px solid #bbb;">'
            '<p style="color:#444;font-size:10pt;">' + "<br>".join(f"<b>{esc(k)}:</b> {esc(v)}" for k, v in head_lines if v) + '</p>'
            f'{o_html}</div>')
    msg = EmailMessage()
    msg["Subject"] = "Re: " + SUBJECT
    msg["From"], msg["To"] = FROM_ADDR, addr
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=(FROM_ADDR.split("@")[-1] or "localhost"))
    if b.get("message_id"):                       # this is what keeps the original mail in the same thread
        msg["In-Reply-To"] = b["message_id"]
        msg["References"] = b["message_id"]
    msg.set_content(text)
    msg.add_alternative(page, subtype="html")
    return msg


def due_reminders(state: dict, force: bool = False) -> list:
    now, out = datetime.now(), []
    for b in state["batches"]:
        age = now - _sent_at(b)
        if age > timedelta(days=REPLY_WATCH_DAYS) or not batch_open(state, b):
            continue
        if age < timedelta(hours=REMINDER_AFTER_HOURS) and not force:
            continue
        for addr, role in REPLY_SENDERS.items():
            if addr in b["replies"]:
                continue
            if (b.get("reminders") or {}).get(addr) and not force:
                continue
            out.append((b, addr, role))
    return out


def reminder_step(state: dict, preview: bool = False) -> None:
    """preview = write each reminder as an .eml file (opens in Outlook), send nothing, remember nothing."""
    items = due_reminders(state, force=preview)
    if not items:
        info(f"Reminder: nobody needs one yet (after {REMINDER_AFTER_HOURS} h without a reply)")
        return
    for b, addr, role in items:
        msg = build_reminder(state, b, addr, role)
        label = f"{role} - mail of {_sent_at(b):%d-%b %H:%M}"
        if preview:
            path = HERE / f"Reminder_{role.replace(' ', '_')}_{_sent_at(b):%Y%m%d_%H%M}.eml"
            path.write_bytes(bytes(msg))
            ok(f"PREVIEW only - reminder for {label} saved, nothing sent: {path.name}")
            continue
        send_with_retry(msg, recipients=[addr] + [c for c in REMINDER_CC if c.lower() != addr.lower()])
        b.setdefault("reminders", {})[addr] = {"at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                               "message_id": str(msg["Message-ID"])}
        save_state(state)
        STATS["reminders"] += 1
        ok(f"Gentle reminder sent to {label}")


def reminder_step_safe(state: dict, **kw) -> None:
    try:
        reminder_step(state, **kw)
    except Exception as exc:               # a reminder problem must never break the run
        warn(f"Reminder step failed: {type(exc).__name__}: {exc}")


# ==========================================================================
# ▶️  Main flow
# ==========================================================================
def check_config(dry_run: bool, replies_only: bool = False) -> None:
    if replies_only:
        missing = [n for n, v in (("POP_HOST", POP_HOST), ("POP_USER or SMTP_USER", POP_USER or SMTP_USER),
                                  ("POP_PASSWORD or SMTP_PASSWORD", POP_PASSWORD or SMTP_PASSWORD)) if not v]
    else:
        missing = [n for n, v in (("USERNAME", USERNAME), ("PASSWORD", PASSWORD)) if not v]
        if not dry_run:
            missing += [n for n, v in (("SMTP_HOST", SMTP_HOST), ("FROM_ADDR", FROM_ADDR), ("TO", TO)) if not v]
    if missing:
        raise AppError("Fill these in the CONFIG block first: " + ", ".join(missing))


def process_new(client, state: dict, dry_run: bool, fresh: bool) -> int:
    sent = state["sent"]
    rows = extract_rows(client.get_tasks_html(fresh=fresh))
    ok(f"Task Box has {len(rows)} task row(s)")
    if len(rows) >= PAGE_SIZE_HINT:
        warn(f"{len(rows)} rows = one full page; further pages are NOT read yet (needs the "
             "/tasksApi/LoadTasksByCategory request).")
    records = [parse_row(r) for r in rows]
    new = [r for r in records if r["_key"] and r["_key"] not in sent]
    skipped = len(records) - len(new)
    if skipped:
        info(f"{skipped} already sent earlier - skipped")
    if not new:
        ok("No new records - nothing to send")
        return 0
    info(f"{len(new)} new record(s): " + ", ".join(f"{r['Name']} ({r['Emp ID']})" for r in new))

    enrich(client, new)
    ready = [r for r in new if r["_ok"]]
    if not ready:
        fail("Details could not be fetched for any new record - nothing sent; will retry on the next run")
        return 1

    msg = build_message(ready)
    if dry_run:
        PREVIEW_FILE.write_text(msg.get_body(("html",)).get_content(), encoding="utf-8")
        ok(f"DRY RUN: nothing sent, nothing remembered. Preview saved to {PREVIEW_FILE.name}")
        return 0

    send_with_retry(msg)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for r in ready:
        sent[r["_key"]] = {"name": r["Name"], "emp_code": r["Emp ID"], "sent_at": stamp, "record": record_of(r)}
    state["batches"].append({"message_id": str(msg["Message-ID"]), "sent_at": stamp,
                             "keys": [r["_key"] for r in ready], "replies": {}})
    save_state(state)
    STATS["emailed"] += len(ready)
    ok(f"Remembered {len(ready)} employee(s) as sent")
    return 0 if len(ready) == len(new) else 1


def run(client, dry_run: bool = False, fresh: bool = False) -> int:
    state = load_state()
    sent_before = len(state["sent"])
    pending_before = [k for k, v in state["sent"].items() if not v.get("acted")]
    info(f"{sent_before} employee(s) already e-mailed earlier")
    try:
        rc = process_new(client, state, dry_run, fresh)
    except AppError as exc:
        fail(str(exc))
        rc = 1
    except requests.RequestException as exc:
        fail(f"Network error: {type(exc).__name__}: {exc}")
        rc = 2
    if dry_run:
        return rc
    if len(state["sent"]) > sent_before:  # NEW task(s) were just e-mailed
        try:
            write_excel(state)
        except Exception as exc:
            warn(f"Excel not updated now ({type(exc).__name__}: {exc}) - it is rebuilt on the next run")
        if not pending_before:            # nothing else was open, so no reply can be waiting: stop here
            ok("New task(s) e-mailed - stopping. Replies are checked on your next run")
            return rc
        info("New task(s) e-mailed - the earlier open task(s) are still checked for replies")
    pending = [k for k, v in state["sent"].items() if not v.get("acted")]
    if not pending:                       # nothing new and nothing waiting: do not touch the mailbox at all
        STATS["idle"] = True
        ok("No new task assigned and nothing is waiting for replies - stopping")
        return rc
    info(f"{len(pending)} employee(s) waiting for replies / Act step - checking the mailbox")
    finish_replies(state)
    if ACT_ENABLED and "--no-act" not in ARGS:
        rc = max(rc, act_step_safe(client, state, fresh=False))
    if REMINDER_ENABLED and "--no-reminder" not in ARGS:
        reminder_step_safe(state)
    if ALERT_ENABLED and "--no-alert" not in ARGS:
        alert_step_safe(state)
    return rc


def acquire_lock() -> bool:
    if LOCK_FILE.exists():
        age_h = (time.time() - LOCK_FILE.stat().st_mtime) / 3600
        if age_h < LOCK_STALE_HOURS:
            return False
        warn(f"Ignoring a stale lock file ({age_h:.1f} h old)")
    LOCK_FILE.write_text(str(os.getpid()), encoding="utf-8")
    return True


def mark_done(args: list) -> int:
    """--mark-done EMP_ID [EMP_ID ...] : record employees as already submitted in Darwinbox (done by hand),
    so the Act step never touches them. Needs no login and sends no mail."""
    ids = [a.strip().upper() for a in args if a.strip() and not a.startswith("--")]
    if not ids:
        raise AppError("Usage:  python darwinbox_exit_mailer.py --mark-done M0214691 [M0214700 ...]")
    state = load_state()
    index = {k.upper(): k for k in state["sent"]}
    index.update({str(v.get("emp_code", "")).upper(): k for k, v in state["sent"].items() if v.get("emp_code")})
    changed = 0
    for emp in ids:
        key = index.get(emp)
        if not key:
            warn(f"{emp}: not in {STATE_FILE.name} (was never e-mailed) - nothing marked")
            continue
        v = state["sent"][key]
        if v.get("acted"):
            info(f"{emp} {v.get('name', '')}: already marked as submitted ({v['acted'].get('at')}) - unchanged")
            continue
        v["acted"] = {"at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "total": "", "remarks": "submitted manually"}
        ok(f"{emp} {v.get('name', '')}: marked as submitted - the Act step will skip this employee")
        changed += 1
    if changed:
        save_state(state)
    return 0


def main() -> int:
    if "--mark-done" in ARGS:
        try:
            if not acquire_lock():
                warn("Another run is still in progress (lock file present) - try again in a minute")
                return 0
            try:
                return mark_done(ARGS[ARGS.index("--mark-done") + 1:])
            finally:
                LOCK_FILE.unlink(missing_ok=True)
        except AppError as exc:
            fail(str(exc))
            return 1
    if "--reminder-preview" in ARGS:
        try:
            if not acquire_lock():
                warn("Another run is still in progress (lock file present) - try again in a minute")
                return 0
            try:
                reminder_step_safe(load_state(), preview=True)
                return 0
            finally:
                LOCK_FILE.unlink(missing_ok=True)
        except AppError as exc:
            fail(str(exc))
            return 1
    if "--alert-preview" in ARGS or "--alert-now" in ARGS:
        try:
            if not acquire_lock():
                warn("Another run is still in progress (lock file present) - try again in a minute")
                return 0
            try:
                alert_step_safe(load_state(), preview="--alert-preview" in ARGS, now="--alert-now" in ARGS)
                return 0
            finally:
                LOCK_FILE.unlink(missing_ok=True)
        except AppError as exc:
            fail(str(exc))
            return 1
    dry_run, fresh = "--dry-run" in ARGS, "--fresh" in ARGS
    replies_only, act_only = "--replies-only" in ARGS, "--act-only" in ARGS
    log("📧", "Darwinbox exit clearance mailer started" + ("  [DRY RUN]" if dry_run else "")
        + ("  [REPLIES ONLY]" if replies_only else "") + ("  [ACT STEP ONLY]" if act_only else ""), "b")
    try:
        check_config(dry_run or act_only, replies_only)
        if not acquire_lock():
            warn("Another run is still in progress (lock file present) - exiting")
            return 0
        try:
            if replies_only:
                finish_replies(load_state())
                return 0
            if act_only:                  # no mail, no mailbox: only the Act step, using the replies already stored
                return act_step_safe(Darwinbox(), load_state(), fresh=fresh)
            return run(Darwinbox(), dry_run=dry_run, fresh=fresh)
        finally:
            LOCK_FILE.unlink(missing_ok=True)
    except AppError as exc:
        fail(str(exc))
        return 1
    except requests.RequestException as exc:
        fail(f"Network error: {type(exc).__name__}: {exc}")
        return 2
    except KeyboardInterrupt:
        warn("Stopped by user")
        return 130


def configure(settings: dict) -> None:
    """Fill the credentials / options (from the console's connections and the process settings) before a run.
    settings keys: darwin_user, darwin_password, smtp_user, smtp_password, smtp_host, smtp_port, pop_user, pop_password,
    pop_host, pop_port, wa_token, wa_upload_url, wa_send_url, wa_template, wa_lang, to, cc, signature, act_submit,
    overdue_phones"""
    global USERNAME, PASSWORD, SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, FROM_ADDR, POP_HOST, POP_PORT, POP_USER
    global POP_PASSWORD, WA_AUTH_TOKEN, WA_UPLOAD_URL, WA_SEND_URL, WA_TEMPLATE, WA_LANGUAGE, WA_TO, TO, CC, REMINDER_CC
    global SIGNATURE, ACT_SUBMIT, ACT_ALERT_TO
    s = settings or {}
    g = lambda k, d="": (str(s.get(k)).strip() if s.get(k) not in (None, "") else d)
    USERNAME, PASSWORD = g("darwin_user", USERNAME), g("darwin_password", PASSWORD)
    SMTP_HOST, SMTP_USER, SMTP_PASSWORD = g("smtp_host", SMTP_HOST), g("smtp_user", SMTP_USER), g("smtp_password", SMTP_PASSWORD)
    SMTP_PORT = int(g("smtp_port", str(SMTP_PORT)) or SMTP_PORT)
    FROM_ADDR = g("from_addr") or SMTP_USER or FROM_ADDR
    POP_HOST, POP_USER, POP_PASSWORD = g("pop_host", POP_HOST), g("pop_user", POP_USER), g("pop_password", POP_PASSWORD)
    POP_PORT = int(g("pop_port", str(POP_PORT)) or POP_PORT)
    WA_AUTH_TOKEN, WA_UPLOAD_URL, WA_SEND_URL = g("wa_token", WA_AUTH_TOKEN), g("wa_upload_url", WA_UPLOAD_URL), g("wa_send_url", WA_SEND_URL)
    WA_TEMPLATE, WA_LANGUAGE = g("wa_template", WA_TEMPLATE), g("wa_lang", WA_LANGUAGE)
    if s.get("to"):
        TO = list(s["to"])
    if s.get("cc") is not None:
        CC = list(s["cc"])
    REMINDER_CC = list(CC)
    if s.get("overdue_phones") is not None:
        WA_TO = list(s["overdue_phones"])
    if s.get("signature"):
        SIGNATURE = [ln.strip() for ln in str(s["signature"]).splitlines() if ln.strip()] or SIGNATURE
    if s.get("act_submit") is not None:
        ACT_SUBMIT = bool(s["act_submit"])
    ACT_ALERT_TO = [FROM_ADDR] if FROM_ADDR else []


def run_once(dry_run: bool = False, fresh: bool = False, args: list | None = None) -> dict:
    """One complete run (what the scheduler calls). Returns {"rc": int, "stats": {...}}; raises AppError on a hard stop."""
    global ARGS
    ARGS = list(args or [])
    _reset_stats()
    check_config(dry_run)
    if not acquire_lock():
        raise AppError("Another Darwinbox run is still in progress (lock file present).")
    try:
        rc = run(Darwinbox(), dry_run=dry_run, fresh=fresh)
    finally:
        LOCK_FILE.unlink(missing_ok=True)
    return {"rc": rc, "stats": dict(STATS)}


def _cli_bootstrap() -> None:
    """Manual use: read the same settings the console uses (.env / .env.local) and the process settings."""
    global ARGS
    ARGS = sys.argv[1:]
    try:
        from dotenv import load_dotenv
        load_dotenv(_REPO / ".env.local", override=True)
        load_dotenv(_REPO / ".env", override=False)
        sys.path.insert(0, str(_REPO))
        from core.scheduler import connections as _c
        from core.scheduler.processes import darwinbox_exit_clearance as _p
        configure(_p.settings_from_connections({}))
    except Exception as exc:
        warn(f"Could not read the console settings ({type(exc).__name__}: {exc}) - using what is in this file")


if __name__ == "__main__":
    _cli_bootstrap()
    sys.exit(main())
