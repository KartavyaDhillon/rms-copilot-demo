"""PARTIC ID  ↔  RISK REPORT CP-CODE validation — the reusable parts of PARTIC_ID_AUTOMATION.py.

(mailbox download over POP3 with binary search, reading the partic id file, matching, Excel report)
Settings come from the caller / Scheduler → Connections — nothing is hard-coded here.
"""
from __future__ import annotations

import csv
import email
import os
import poplib
import re
import ssl
import time
from datetime import datetime, timedelta, timezone, time as dtime
from email import policy
from email.parser import BytesParser
from email.utils import parseaddr
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

IST = timezone(timedelta(hours=5, minutes=30))

SEG_MAP = {"NSE": "NSECAP", "BSE": "BSECAP", "NFO": "NSEFNO", "BFO": "BSEFNO", "MCX": "MCXCOM", "NCOM": "NSECOM"}
NOT_ENABLED_SEGMENTS = {"CDS", "BCD", "MFSS"}
PARTIC_STEMS = {"output", "partic_id", "particid", "partic id"}
DATA_EXTS = (".txt", ".csv", ".xlsx", ".xls")


def today_ddmmyyyy() -> str:
    return datetime.now(IST).strftime("%d%m%Y")


# ───────────────────────── find / download the partic id file ─────────────────────────
def find_partic_file(folder: Path) -> Optional[Path]:
    if not folder.is_dir():
        return None
    for p in sorted(folder.iterdir()):
        if p.is_file() and p.suffix.lower() in DATA_EXTS and p.stem.strip().lower() in PARTIC_STEMS:
            return p
    return None


def _relaxed_ssl_context():
    """pop9.netcore.co.in uses a legacy small DH key that Python 3.11+ rejects by default; SECLEVEL=1 allows it
    for this one connection."""
    ctx = ssl.create_default_context()
    ctx.set_ciphers("DEFAULT@SECLEVEL=1")
    return ctx


def _parse_message(lines):
    return BytesParser(policy=policy.default).parsebytes(b"\r\n".join(lines))


def _header_dt_and_msg(pop, i):
    _resp, lines, _o = pop.top(i, 0)
    msg = _parse_message(lines)
    dh = msg["Date"]
    if not dh:
        return None, msg
    try:
        dt = email.utils.parsedate_to_datetime(dh)
    except (TypeError, ValueError):
        return None, msg
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=IST)
    return dt.astimezone(IST), msg


def _dt_at(pop, i, count):
    order = []
    for step in range(0, 4):
        if i + step <= count:
            order.append(i + step)
        if step and i - step >= 1:
            order.append(i - step)
    for j in order:
        try:
            dt, _ = _header_dt_and_msg(pop, j)
        except (TimeoutError, OSError, poplib.error_proto):
            continue
        if dt is not None:
            return dt
    return None


def _find_start_index(pop, count, bound):
    lo, hi = 1, count + 1
    while lo < hi:
        mid = (lo + hi) // 2
        dt = _dt_at(pop, mid, count)
        if dt is None:
            hi = mid
            continue
        if dt < bound:
            lo = mid + 1
        else:
            hi = mid
    return lo


def _pick_attachment(full_msg):
    atts = list(full_msg.iter_attachments())
    for part in atts:
        fname = part.get_filename() or ""
        if "partic" in fname.lower():
            return part.get_payload(decode=True), fname
    for part in atts:
        fname = part.get_filename() or ""
        if os.path.splitext(fname)[0].strip().lower() in PARTIC_STEMS:
            return part.get_payload(decode=True), fname
    if len(atts) == 1:
        return atts[0].get_payload(decode=True), (atts[0].get_filename() or f"PARTIC_ID_{today_ddmmyyyy()}.txt")
    return None, None


def download_partic_file(folder: Path, server: str, port: int, user: str, password: str, sender: str,
                         subject_exact: str, scan_start: dtime, log: Callable[[str], None],
                         timeout: int = 30, max_scan: int = 50000) -> Optional[Path]:
    """Save today's matching e-mail attachment into `folder`; returns the path or None."""
    target = datetime.now(IST).date()
    bound = datetime.combine(target, scan_start, tzinfo=IST)
    log(f"Mailbox search: from '{sender}', subject exactly '{subject_exact}', today from {bound:%H:%M} IST")
    folder.mkdir(parents=True, exist_ok=True)
    pop = poplib.POP3_SSL(server, port, context=_relaxed_ssl_context(), timeout=timeout)
    saved = None
    try:
        pop.user(user)
        pop.pass_(password)
        count, _size = pop.stat()
        log(f"Mailbox login OK — {count} message(s).")
        if count == 0:
            return None
        start = _find_start_index(pop, count, bound)
        if start > count:
            log("No message on/after the start time yet.")
            return None
        scanned = 0
        for i in range(start, count + 1):
            if scanned >= max_scan:
                break
            scanned += 1
            try:
                dt, msg = _header_dt_and_msg(pop, i)
            except (TimeoutError, OSError, poplib.error_proto):
                continue
            if dt is None:
                continue
            if dt.date() > target:
                break
            if dt.date() < target:
                continue
            if (msg["Subject"] or "").strip().lower() != subject_exact.lower():
                continue
            allowed = {s.strip().lower() for s in re.split(r"[,;\s]+", sender or "") if s.strip()}
            if allowed and parseaddr(msg["From"] or "")[1].strip().lower() not in allowed:
                continue
            try:
                _r, full_lines, _o = pop.retr(i)
            except (TimeoutError, OSError, poplib.error_proto):
                continue
            data, name = _pick_attachment(_parse_message(full_lines))
            if data is None:
                continue
            saved = folder / os.path.basename(name)
            saved.write_bytes(data)
            log(f"Downloaded {saved.name} (mail received {dt:%H:%M:%S} IST)")
            break
        if saved is None:
            log("No matching e-mail for today (yet).")
        return saved
    finally:
        try:
            pop.quit()
        except Exception:
            pass


# ───────────────────────── read the files ─────────────────────────
def _rows_from_table(raw_rows) -> List[Tuple[str, str, str]]:
    rows = []
    if not raw_rows:
        return rows
    lower = [str(c).strip().lower() for c in raw_rows[0]]
    ic, isg, ip, start = 0, 1, 2, 0
    if any(h in lower for h in ("client_id", "clientid", "client id")):
        def find(*names):
            for n in names:
                if n in lower:
                    return lower.index(n)
            return None
        ic = find("client_id", "clientid", "client id")
        isg = find("segment", "seg", "exchange")
        ip = find("partic_id", "particid", "partic id", "output_partic_id")
        ic, isg, ip = (0 if ic is None else ic), (1 if isg is None else isg), (2 if ip is None else ip)
        start = 1
    for r in raw_rows[start:]:
        cells = [str(c).strip() for c in r]
        if not any(cells) or max(ic, isg, ip) >= len(cells):
            continue
        cid, seg, pid = cells[ic], cells[isg], cells[ip]
        if not (cid or seg or pid):
            continue
        if cid.lower() == "client_id" and seg.lower() == "segment":
            continue
        rows.append((cid, seg, pid))
    return rows


def read_partic_rows(path: Path):
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = [["" if c is None else c for c in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
        return _rows_from_table(raw)
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            return _rows_from_table([list(r) for r in csv.reader(f, dialect)])
    raw = []
    with open(path, encoding="utf-8-sig") as f:
        for ln in f:
            ln = ln.rstrip("\r\n")
            if not ln.strip():
                continue
            if "\t" in ln:
                parts = [p.strip() for p in ln.split("\t")]
            elif "," in ln:
                parts = [p.strip() for p in ln.split(",")]
            else:
                parts = ln.split()
            raw.append(parts)
    return _rows_from_table(raw)


def read_risk_cpcodes(path: Path) -> Dict[str, str]:
    risk = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f)
        header = [h.strip() for h in next(rd)]
        if "ACCOUNTID" not in header or "CPCode" not in header:
            raise RuntimeError("Risk report has no ACCOUNTID / CPCode column.")
        ai, ci = header.index("ACCOUNTID"), header.index("CPCode")
        for fields in rd:
            if len(fields) > ci:
                risk[fields[ai].strip()] = fields[ci].strip()
    return risk


# ───────────────────────── matching ─────────────────────────
def parse_cp(cp: str) -> Dict[str, List[str]]:
    d: Dict[str, List[str]] = {}
    for tok in cp.split("$"):
        tok = tok.strip()
        if not tok or "-" not in tok:
            continue
        pref, pid = tok.split("-", 1)
        pref, pid = pref.strip().upper(), pid.strip()
        d.setdefault(pref, [])
        if pid not in d[pref]:
            d[pref].append(pid)
    return d


def run_matching(rows, risk):
    matched, mismatched, notfound = [], [], []

    def elsewhere(parsed, pid):
        keys = sorted(k for k, v in parsed.items() if pid in v)
        return ("YES - " + ", ".join(keys)) if keys else "NO"

    for cid, seg, pid in rows:
        if cid not in risk:
            notfound.append({"Client_ID": cid, "Segment": seg, "Output_Partic_ID": pid,
                             "Reason": "CLIENT ID NOT FOUND IN RISK REPORT"})
            continue
        cp = risk[cid]
        if not cp:
            notfound.append({"Client_ID": cid, "Segment": seg, "Output_Partic_ID": pid,
                             "Reason": "CPCode is BLANK in risk report"})
            continue
        parsed = parse_cp(cp)
        pref = SEG_MAP.get(seg.upper())
        base = {"Client_ID": cid, "Segment": seg, "Output_Partic_ID": pid,
                "Partic_ID_seen_elsewhere": elsewhere(parsed, pid), "Full_CPCode": cp}
        if pref is None:
            reason = (f"Segment {seg} is not enabled on the trading platform" if seg.upper() in NOT_ENABLED_SEGMENTS
                      else f"Segment {seg} has no matching CP code type")
            mismatched.append({**base, "CP_Segment": "(none for this segment)", "Risk_Partic_ID": "-", "Reason": reason})
        elif pref not in parsed:
            mismatched.append({**base, "CP_Segment": pref, "Risk_Partic_ID": "-",
                               "Reason": f"No {pref} entry in this client's CP code"})
        elif pid in parsed[pref]:
            matched.append({"Client_ID": cid, "Segment": seg, "Output_Partic_ID": pid, "CP_Segment": pref,
                            "Risk_Partic_ID": "/".join(parsed[pref])})
        else:
            mismatched.append({**base, "CP_Segment": pref, "Risk_Partic_ID": "/".join(parsed[pref]),
                               "Reason": f"Partic_ID MISMATCH in {pref}"})
    return matched, mismatched, notfound


# ───────────────────────── Excel report ─────────────────────────
def build_report(matched, mismatched, notfound, rows, risk, out_path: Path, partic_name: str, risk_name: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    real = [m for m in mismatched if m["Reason"].startswith("Partic_ID MISMATCH")]
    NAVY, GREEN, RED, ORANGE = "1F3864", "2E7D32", "C0392B", "E67E22"
    GREEN_L, RED_L, ORANGE_L, GREY, WHITE = "E8F5E9", "FDECEA", "FDF2E6", "808080", "FFFFFF"
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def header(ws, n, color):
        for c in ws[1]:
            c.font = Font(bold=True, color=WHITE, size=11)
            c.fill = PatternFill("solid", fgColor=color)
            c.alignment = center
            c.border = border
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(n)}1"

    def fit(ws, heads, records, keys, cap=60):
        for i, (h, k) in enumerate(zip(heads, keys), start=1):
            longest = max([len(h)] + [len(str(r.get(k, ""))) for r in records])
            ws.column_dimensions[get_column_letter(i)].width = min(longest + 3, cap)

    def sheet(wb, title, heads, keys, recs, color, light, status=None, red_col=None, cap=60):
        ws = wb.create_sheet(title)
        ws.append(heads)
        header(ws, len(heads), color)
        fill = PatternFill("solid", fgColor=light)
        for rec in recs:
            ws.append([rec.get(k, "") for k in keys] + ([status] if status else []))
            for c in ws[ws.max_row]:
                c.fill, c.border, c.alignment = fill, border, Alignment(vertical="center")
            if red_col is not None:
                ws[ws.max_row][red_col].font = Font(bold=True, color=color)
        fit(ws, heads, recs, keys, cap)
        return ws

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A2"] = "CP CODE  vs  PARTIC ID  RECONCILIATION"
    ws["A2"].font = Font(bold=True, size=20, color=NAVY)
    ws.merge_cells("A2:D2")
    ws["A3"] = (f"Source of truth: {partic_name}     |     Checked against: {risk_name}     |     "
                f"Report date: {datetime.now(IST):%d-%b-%Y}")
    ws["A3"].font = Font(italic=True, size=10, color=GREY)
    ws.merge_cells("A3:D3")
    total_c = len({r[0] for r in rows})
    found_c = len({r[0] for r in rows if r[0] in risk})
    r = 5
    for label, value, color in [("Total rows checked", len(rows), NAVY), ("MATCHED", len(matched), GREEN),
                                ("MISMATCHED", len(mismatched), RED), ("Not Found / Blank", len(notfound), ORANGE)]:
        ws.merge_cells(f"A{r}:B{r}")
        c = ws[f"A{r}"]
        c.value = label
        c.font = Font(bold=True, size=13, color=WHITE)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws[f"B{r}"].fill = PatternFill("solid", fgColor=color)
        n = ws[f"C{r}"]
        n.value = value
        n.font = Font(bold=True, size=16, color=color)
        n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 26
        r += 1
    ws["A10"] = "Notes"
    ws["A10"].font = Font(bold=True, size=12, color=NAVY)
    notes = [f"Clients in {partic_name}: {total_c}  |  found in risk report: {found_c}  |  not found: {total_c - found_c}",
             f"Genuine Partic_ID conflicts (same segment, different ID): {len(real)}.",
             "Segment mapping used:  " + ",  ".join(f"{k}->{v}" for k, v in SEG_MAP.items()) + ".",
             "CDS / BCD / MFSS etc. are not enabled on the trading platform, so they are listed as MISMATCHED with that "
             "remark. The 'Partic_ID seen elsewhere' column shows when the same Partic_ID appears under another CP "
             "segment for that client."]
    rr = 11
    for n in notes:
        ws.merge_cells(f"A{rr}:H{rr}")
        c = ws[f"A{rr}"]
        c.value = n
        c.font = Font(size=10, color="404040")
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if len(n) > 120:
            ws.row_dimensions[rr].height = 46
        rr += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14
    ws.sheet_view.showGridLines = False

    mh = ["Client_ID", "Segment", "Output Partic_ID", "CP Segment", "Risk Partic_ID", "Status"]
    mk = ["Client_ID", "Segment", "Output_Partic_ID", "CP_Segment", "Risk_Partic_ID"]
    ws2 = sheet(wb, "Matched", mh, mk, matched, GREEN, GREEN_L, status="MATCHED", red_col=5)
    ws2.column_dimensions["F"].width = 12
    sheet(wb, "Mismatched",
          ["Client_ID", "Segment", "Output Partic_ID", "CP Segment", "Risk Partic_ID", "Reason",
           "Partic_ID seen elsewhere", "Full CPCode (risk report)"],
          ["Client_ID", "Segment", "Output_Partic_ID", "CP_Segment", "Risk_Partic_ID", "Reason",
           "Partic_ID_seen_elsewhere", "Full_CPCode"], mismatched, RED, RED_L, red_col=5, cap=70)
    sheet(wb, "Not Found - Blank", ["Client_ID", "Segment", "Output Partic_ID", "Reason"],
          ["Client_ID", "Segment", "Output_Partic_ID", "Reason"], notfound, ORANGE, ORANGE_L, red_col=3)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return real
