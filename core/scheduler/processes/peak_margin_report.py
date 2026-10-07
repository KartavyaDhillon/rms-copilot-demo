"""PEAK MARGIN REPORT — NSE peak-margin snapshots i01–i04 (FO + CM SA01 files) → cumulative Peak1..Peak4 report, the same
images as the console's Peak Margin box (core/pipelines/report_image.py — no Excel needed), sent on WhatsApp and e-mailed
(from Peak_Margin_Final.py).

Each run handles the NEXT snapshot of the day (i01 → i04) unless 'Snapshot' is set to a specific one.  Schedule it at
10:30, 12:30, 13:45 and 15:30 (it waits for the files).  Its own state lives in this process's state folder, so it never
disturbs the console's Peak Margin box.
"""
import datetime as dt
import gzip
import io
import json

from core.scheduler import connections, helpers as H, nse, paths
from core.scheduler.catalog import Param, register
from core.scheduler.processes import _kit as K

FOLDER = "PEAK MARGIN REPORT"
INDEXES = ["i01", "i02", "i03", "i04"]
PEAKS = ["Peak1", "Peak2", "Peak3", "Peak4"]
COLS = ["Date", "Code", "Peak1", "Peak2", "Peak3", "Peak4", "Other"]


def _state_file():
    return paths.proc_dir(FOLDER, "state") / "peak_margin_report.json"


def _empty(today, fo_codes):
    seg = lambda: {"Code": None, **{p: 0.0 for p in PEAKS}}
    return {"date": today, "done": [], "fo_clients": {c: {p: 0.0 for p in PEAKS} for c in fo_codes},
            "cm_segment": seg(), "fo_segment": seg()}


def _load(today, fo_codes):
    try:
        rep = json.loads(_state_file().read_text(encoding="utf-8"))
        if rep.get("date") == today:
            for c in fo_codes:
                rep["fo_clients"].setdefault(c, {p: 0.0 for p in PEAKS})
            return rep
    except Exception:
        pass
    return _empty(today, fo_codes)


def _frame(gz_path):
    pd = H.need("pandas")
    with gzip.open(gz_path, "rt") as f:
        df = pd.read_csv(io.StringIO(f.read()), header=None)
    df = df.iloc[:, :7]
    df.columns = COLS[:df.shape[1]]
    return df


def _val(df, codes):
    pd = H.need("pandas")
    m = df[df["Code"].isin(codes)]
    if m.empty:
        return None, 0.0
    v = pd.to_numeric(m.iloc[0]["Peak1"], errors="coerce")
    return m.iloc[0]["Code"], float(0.0 if v != v else v)


def _xlsx(rep, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    fill = PatternFill(start_color="C4D79B", end_color="C4D79B", fill_type="solid")
    thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    wb = Workbook()
    ws = wb.active
    r = 1

    def seg(title, d):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(r, 1, title)
        c.font, c.fill, c.alignment = Font(bold=True), fill, Alignment(horizontal="center", vertical="center")
        for col in range(1, 6):
            ws.cell(r, col).border = Border(bottom=Side(style="thick"))
        r += 1
        for j, h in enumerate(["Client Code"] + PEAKS, 1):
            x = ws.cell(r, j, h)
            x.font, x.fill, x.border, x.alignment = Font(bold=True), fill, thin, Alignment(horizontal="center", vertical="center")
        r += 1
        for j, v in enumerate([d["Code"]] + [d[p] for p in PEAKS], 1):
            x = ws.cell(r, j, v)
            x.alignment, x.number_format, x.border = Alignment(horizontal="right", vertical="center"), "#,##0.00", thin
        r += 2
    seg("CM Segment", rep["cm_segment"])
    seg("FO Segment", rep["fo_segment"])
    ws.cell(r, 1, "TOTAL (CM + FO)").font = Font(bold=True)
    ws.cell(r, 1).fill = fill
    for j, p in enumerate(PEAKS, 2):
        x = ws.cell(r, j, rep["combined_totals"][p])
        x.number_format, x.border, x.alignment = "#,##0.00", thin, Alignment(horizontal="right", vertical="center")
    for j, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(len(str(c.value)) if c.value else 0 for c in col) + 5
    wb.save(path)


@H.logged(FOLDER)
def _run(params, ctx):
    H.need("PIL", "pillow")
    H.need("openpyxl")
    to, cc = K.recipients(params)
    f_phones, c_phones = H.parse_phones(params.get("phones_fo")), H.parse_phones(params.get("phones_cm"))
    K.require_trading_day(ctx)
    fo_codes = H.parse_list(params.get("fo_codes"))
    cm_codes = H.parse_list(params.get("cm_codes"))
    if not fo_codes or not cm_codes:
        raise ValueError("Set the FO client codes and the CM / FO segment codes (Settings of this process).")
    today = H.today_ddmmyyyy()
    mem = connections.require("nse", "member_code")["member_code"]
    rep = _load(today, fo_codes)

    pick = (params.get("snapshot") or "auto").strip()
    if pick == "auto":
        idx = next((i for i in INDEXES if i not in rep["done"]), None)
        if not idx:
            raise H.Skipped("All four snapshots (i01–i04) are already done today. Choose a snapshot to repeat one.")
    else:
        idx = pick
    col = PEAKS[INDEXES.index(idx)]
    ctx.log(f"Snapshot {idx} → column {col}.")

    inp, out = paths.proc_dir(FOLDER, "inputs"), paths.proc_dir(FOLDER, "outputs")
    wait = K.wait_of(params)
    f_gz = nse.download_wait("member", "FO", "Reports", f"F_SA01_P_{mem}_{today}_{idx}.lis.gz", inp, ctx, wait)
    c_gz = nse.download_wait("member", "CM", "Reports", f"C_SA01_P_{mem}_{today}_{idx}.lis.gz", inp, ctx, wait)
    fdf, cdf = _frame(f_gz), _frame(c_gz)

    for code in fo_codes:
        m = fdf[fdf["Code"] == code]
        v = H.need("pandas").to_numeric(m.iloc[0]["Peak1"], errors="coerce") if not m.empty else 0.0
        rep["fo_clients"][code][col] = float(0.0 if v != v else v)
    cm_code, cm_val = _val(cdf, cm_codes)
    fo_code, fo_val = _val(fdf, cm_codes)
    if cm_code is not None:
        rep["cm_segment"]["Code"] = cm_code
    rep["cm_segment"][col] = cm_val
    if fo_code is not None:
        rep["fo_segment"]["Code"] = fo_code
    rep["fo_segment"][col] = fo_val
    if idx not in rep["done"]:
        rep["done"].append(idx)
    _state_file().write_text(json.dumps(rep, indent=2), encoding="utf-8")
    rep["combined_totals"] = {p: rep["cm_segment"][p] + rep["fo_segment"][p] for p in PEAKS}

    from PIL import Image
    from core.pipelines.report_image import build_report_images
    imgs = build_report_images(rep, str(out))
    jp = {}
    for k, name in (("fo_image", "F_filtered.jpg"), ("cm_image", "C_filtered.jpg")):
        Image.open(imgs[k]).convert("RGB").save(out / name, "JPEG", quality=95)
        jp[k] = out / name
    xl = out / "06436_Peak_Margin.xlsx"
    _xlsx(rep, xl)

    d = H.today_dash()
    failed = []
    for ph in f_phones:
        try:
            H.whatsapp_send_image(jp["fo_image"], ph, "BJDA_Peak Margin", d, ctx)
        except Exception as e:
            failed.append(ph[-4:])
            ctx.log(f"WhatsApp (FO) to …{ph[-4:]} failed: {e}")
    for ph in c_phones:
        try:
            H.whatsapp_send_image(jp["cm_image"], ph, "06436_Peak_Margin", d, ctx)
        except Exception as e:
            failed.append(ph[-4:])
            ctx.log(f"WhatsApp (CM) to …{ph[-4:]} failed: {e}")

    body = (f"<html><body>Dear Sir,<br><br>Please find the Peak Margin Report dated {d}<br><br>"
            f"<img src=\"cid:peak\"><br><br>Best regards,<br>{K.sig(params)}</body></html>")
    H.send_email("PEAK MARGIN REPORT", body, to, cc, [xl], ctx, html=True, inline_images={"peak": imgs["cm_image"]})

    # keep this snapshot's files, then free the inputs
    H.backup_files(FOLDER, [f_gz, c_gz, jp["fo_image"], jp["cm_image"]], ctx)
    if idx == "i04":
        for f in inp.glob("*"):
            f.unlink(missing_ok=True)
    if failed:
        raise RuntimeError(f"Report e-mailed, but WhatsApp failed for number(s) ending {', '.join(sorted(set(failed)))}.")
    return f"{idx} done — e-mailed; WhatsApp sent"


register(
    key="peak_margin_report",
    name="PEAK MARGIN REPORT",
    description="Takes the next NSE peak-margin snapshot (i01–i04), updates the day's Peak1–Peak4 report, and sends the "
                "images on WhatsApp and by e-mail. Same images as the console's Peak Margin box.",
    run=_run,
    params=K.mail_params("rks@mastertrust.co.in,sukhbirsingh@mastertrust.co.in", "rms_del@mastertrust.co.in",
                         wait_default=60) + [
        Param("phones_fo", "WhatsApp — FO image to", "text", fmt="phones", scope="process",
              default="7541827310, 9891419393"),
        Param("phones_cm", "WhatsApp — CM/Total image to", "text", fmt="phones", scope="process", default="7541827310"),
        Param("fo_codes", "FO client codes", "text", scope="process", default="JDA563, JDA937, JDA1209, JDA1210, JDA1260"),
        Param("cm_codes", "CM / FO segment codes", "text", scope="process", default="06436, PRO_06436, 6436"),
        Param("snapshot", "Snapshot", "choice", choices=["auto", "i01", "i02", "i03", "i04"], default="auto",
              help="auto = the next one of the day. Pick one to repeat it."),
    ],
    retry_safe=False, parallel_safe=False, category="Reports", folder=FOLDER, uses=("nse", "whatsapp", "smtp"),
)
