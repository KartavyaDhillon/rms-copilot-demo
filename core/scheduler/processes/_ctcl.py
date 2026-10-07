"""Helpers for the processes that read the CTCL vs NMASS run-1 output (FULL_FILE_CTCL_VS_NMASS_01.xlsx)."""
from __future__ import annotations

from pathlib import Path

from core.scheduler import helpers as H
from core.scheduler.catalog import Param


def wait_param(default=120):
    return Param("wait_minutes", "Wait for CTCL vs NMASS run 1 (minutes)", "int", default=default,
                 help="The process starts when scheduled and keeps checking every 30 s until today's run-1 output "
                      "(FULL_FILE_CTCL_VS_NMASS_01.xlsx) exists, for up to this many minutes.")


def get_run1(params, ctx) -> Path:
    n = max(0, int(params.get("wait_minutes") or 0))
    return H.wait_for_file(lambda: H.ctcl_run1_file(ctx), n, ctx,
                           "Today's CTCL vs NMASS run-1 output (FULL_FILE_CTCL_VS_NMASS_01.xlsx)")


def read_run1(path: Path, cols):
    """DataFrame of the run-1 workbook: finds the header row (within the first 10 rows) that holds all `cols`."""
    pd = H.need("pandas")
    raw = pd.read_excel(path, header=None, dtype=str)
    want = [c.upper() for c in cols]
    for i in range(min(10, len(raw))):
        row = [str(v).strip().upper() for v in raw.iloc[i].tolist()]
        if all(w in row for w in want):
            df = pd.read_excel(path, header=i)
            df.columns = [str(c).strip().upper() for c in df.columns]
            return df
    raise RuntimeError(f"{path.name}: could not find the column(s) {', '.join(cols)} in the first 10 rows.")


def num(col):
    pd = H.need("pandas")
    return (col.astype(str).str.replace(",", "", regex=False).str.replace("₹", "", regex=False).str.strip()
            .pipe(pd.to_numeric, errors="coerce"))


def _font(size, bold=False):
    from PIL import ImageFont
    for n in (("arialbd.ttf" if bold else "arial.ttf"),
              ("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(n, size)
        except Exception:
            pass
    return ImageFont.load_default()


def save_xlsx(path: Path, title: str, header, rows, footer=None, fmt2=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    n = len(header)
    ws.append([title] + [""] * (n - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(1, 1)
    c.font, c.fill = Font(bold=True, size=14), PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.append(list(header))
    for j in range(1, n + 1):
        ws.cell(2, j).font, ws.cell(2, j).alignment = Font(bold=True), Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(list(r))
    if fmt2:
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=2, max_col=2):
            row[0].number_format = "0.00"
    if footer:
        ws.append([footer] + [""] * (n - 1))
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=n)
        g = ws.cell(ws.max_row, 1)
        g.font, g.alignment = Font(bold=True, size=12), Alignment(horizontal="center", vertical="center")
    thin = Side(border_style="thin", color="000000")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=n):
        for cell in row:
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
    for j in range(1, n + 1):
        L = max([len(str(ws.cell(r, j).value)) for r in range(2, ws.max_row + 1) if ws.cell(r, j).value] or [8])
        ws.column_dimensions[get_column_letter(j)].width = L + 5
    wb.save(path)


def render_equal(title, header, rows, footer, out: Path):
    """Amit_final.py look: 900 px wide, equal columns, grid lines, footer box (e.g. GOOD MORNING SIR)."""
    from PIL import Image, ImageDraw
    data = [[title] + [""] * (len(header) - 1), list(header)] + [list(r) for r in rows]
    W, rh, th, lw = 900, 75, 90, 3
    n = len(header)
    cw = [W // n] * n
    nrows = len(data) - 1
    H_ = th + rh * nrows + (rh if footer else 0)
    img = Image.new("RGB", (W, H_), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, th], fill="#90EE90")
    d.text((W // 2, th // 2), title, font=_font(40, True), fill="black", anchor="mm")
    y = th
    for i, row in enumerate(data[1:]):
        x = 0
        for j, cell in enumerate(row[:n]):
            w = cw[j]
            if i == 0:
                d.rectangle([x, y, x + w, y + rh], fill="#D9D9D9")
                d.text((x + w // 2, y + rh // 2), cell, font=_font(26, True), fill="black", anchor="mm")
            else:
                d.rectangle([x, y, x + w, y + rh], fill="white")
                d.text((x + w // 2, y + rh // 2), cell, font=_font(24), fill="black", anchor="mm")
            x += w
        y += rh
    table_end = y
    yl = th
    for _ in range(len(data)):
        d.line([0, yl, W, yl], fill="black", width=lw)
        yl += rh
    d.line([0, yl, W, yl], fill="black", width=lw)
    bottom = table_end
    if footer:
        top = table_end
        d.rectangle([0, top, W, top + rh], fill="white")
        for a in ([0, top, W, top], [0, top + rh, W, top + rh], [0, top, 0, top + rh], [W, top, W, top + rh]):
            d.line(a, fill="black", width=3)
        d.text((W // 2, top + rh // 2), footer, font=_font(26, True), fill="black", anchor="mm")
        bottom = top + rh
    d.line([0, th, 0, bottom], fill="black", width=3)
    d.line([cw[0], th, cw[0], table_end], fill="black", width=3)
    d.line([W, th, W, bottom], fill="black", width=3)
    d.line([0, th, W, th], fill="black", width=3)
    img.save(out)
    return out


def render_auto(title, header, rows, out: Path):
    """NEW_AMIT_SIR_FINAL.py look: column widths follow the data, boxed cells, auto-fitting title."""
    from PIL import Image, ImageDraw
    data = [[title] + [""] * (len(header) - 1), list(header)] + [list(r) for r in rows]
    th, rh = 90, 65
    cws = [max(len(r[c]) for r in data[1:]) * 22 for c in range(len(header))]
    W = sum(cws)
    img = Image.new("RGB", (W, th + rh * (len(data) - 1)), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, th], fill="#90EE90")
    fs = 42
    while fs > 10:
        f = _font(fs, True)
        bb = d.textbbox((0, 0), title, font=f)
        if bb[2] - bb[0] <= W - 40:
            break
        fs -= 1
    d.text((W // 2, th // 2), title, font=f, fill="black", anchor="mm")
    y = th
    for i, row in enumerate(data[1:]):
        x = 0
        for j, cell in enumerate(row):
            w = cws[j]
            if i == 0:
                d.rectangle([x, y, x + w, y + rh], fill="#D9D9D9", outline="black", width=2)
                d.text((x + w // 2, y + rh // 2), cell, font=_font(26, True), fill="black", anchor="mm")
            else:
                d.rectangle([x, y, x + w, y + rh], fill="white", outline="black", width=1)
                d.text((x + w // 2, y + rh // 2), cell, font=_font(26), fill="black", anchor="mm")
            x += w
        y += rh
    img.save(out)
    return out
