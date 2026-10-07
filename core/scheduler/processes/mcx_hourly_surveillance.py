"""MCX HOURLY SURVEILLANCE — logs in to MCX eclear (RSA soft-token generated from the seed), reads the member's row of
"Margin Utilization View" (RSK335), appends it (in crores) to the monthly master workbook, builds the snapshot image,
e-mails it inline with the workbook and sends it on WhatsApp (from MCX_HOURLY_SURVEILLANCE_AUTOMATION.py, automated
mode only — the manual 'Search Results' mode and the password pop-up are not part of the scheduler).

Schedule it hourly (e.g. 09:00–00:30).  The 'Run only between' window is a second safety net.
"""
import base64
import csv
import datetime as dt
import os
import re
import shutil
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timedelta, timezone
from functools import lru_cache

from core.scheduler import connections, helpers as H, paths
from core.scheduler.catalog import Param, register
from core.scheduler.processes import _kit as K

FOLDER = "MCX HOURLY SURVEILLANCE"
BASE_URL = "https://eclear.mcxccl.com"
MEMBER_ROW = "56720"

COLUMN_MAP = [("DEPOSIT MCX", "TM / CP Margin Limit"), ("CM PRO/ MCX", "Margin Utilization"),
              ("TM Client/MCX", "Total MU (Rs)"), ("CLI MARGIN 90%", "Cli / TM Mrgn 90% (Rs)"),
              ("Used by %", "Margin %")]
FINAL_HEADERS = ["Time"] + [c[0] for c in COLUMN_MAP]
SUBTITLE = "Margin Utilization Monitoring   •   ₹ in Crores"
NAVY, GOLD, HEADP, ALT = "0F2C4D", "C8A24B", "1F4E78", "EEF3F8"
RED, GREEN, AMBER = "C0392B", "1E7D34", "B9770E"
IMG_SCALE, JPEG_QUALITY, WA_MAX_BYTES = 2, 95, 5 * 1024 * 1024


# ───────────────────────── run window ─────────────────────────
def _hm(s, dflt):
    try:
        h, m = (s or dflt).strip().split(":")
        return int(h) * 60 + int(m)
    except Exception:
        raise ValueError(f"'{s}' is not a time like 09:00")


def _in_window(now, first, last, sat_cut):
    wd = now.weekday()
    cur = now.hour * 60 + now.minute
    if wd == 6:
        return False
    if wd == 5 and cur > sat_cut:
        return False                                   # Saturday only for Friday's after-midnight carry-over
    return (first <= cur <= last) if first <= last else (cur >= first or cur <= last)


# ───────────────────────── RSA soft token ─────────────────────────
def _token(ctx):
    sec = H.need("securid", "python-securid")
    from securid.sdtid import SdtidFile
    c = connections.get("mcx_eclear")
    path = c.get("rsa_sdtid")
    if not (path and os.path.exists(path)):
        zp = c.get("rsa_zip")
        if not zp or not os.path.exists(zp):
            raise RuntimeError("MCX eclear: the RSA token zip / .sdtid path is not set (or the file is missing on this "
                               "PC) — set it in Connections → MCX eclear.")
        with zipfile.ZipFile(zp) as z:
            names = [n for n in z.namelist() if n.lower().endswith(".sdtid")]
            if not names:
                raise RuntimeError("No .sdtid seed found inside the RSA zip.")
            tmp = tempfile.mkdtemp(prefix="rsa_")
            z.extract(names[0], tmp)
            path = os.path.join(tmp, names[0])
    return SdtidFile(path).get_token()


# ───────────────────────── crypto / login ─────────────────────────
def _encrypt(value):
    H.need("cryptography")
    from cryptography.hazmat.primitives import padding
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    key, iv = os.urandom(16), os.urandom(16)
    p = padding.PKCS7(128).padder()
    e = Cipher(algorithms.AES(key), modes.CBC(iv)).encryptor()
    enc = e.update(p.update(value.encode()) + p.finalize()) + e.finalize()
    return base64.b64encode(enc).decode(), base64.b64encode(key + iv).decode()


def _entity(text):
    m = re.search(r"<EntityList>.*</EntityList>", text, re.DOTALL)
    if not m:
        raise RuntimeError("Entity missing")
    root = ET.fromstring(m.group(0))
    for e in root.findall(".//EntityDetail"):
        code, desc = e.find("entityCode").text, e.find("shortDescription").text
        if code != "DefaultEntity":
            return code, f"{code}:{desc},"
    raise RuntimeError("No valid entity")


AJAX = {"User-Agent": "Mozilla/5.0", "Accept": "*/*", "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "X-Requested-With": "XMLHttpRequest", "Origin": BASE_URL,
        "Referer": BASE_URL + "/Bancs/common/dologin.jsp?method=authMode&JSLoadingMode=null"}
LOGINH = {"User-Agent": "Mozilla/5.0", "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
          "Content-Type": "application/x-www-form-urlencoded", "Origin": BASE_URL,
          "Referer": BASE_URL + "/Bancs/common/dologin.jsp?method=authMode&JSLoadingMode=null",
          "Upgrade-Insecure-Requests": "1"}


def login(ctx):
    requests = H.need("requests")
    c = connections.require("mcx_eclear", "user_id", "password")
    uid, pwd = c["user_id"], c["password"]
    s = requests.Session()
    s.get(BASE_URL, timeout=60)
    s.post(BASE_URL + "/Bancs/login.do?method=authMode", headers=AJAX, timeout=60)
    enc, ig = _encrypt(pwd)
    r = s.post(BASE_URL + "/Bancs/login.do?method=authenticatePassward", headers=AJAX, timeout=60,
               data={"userId": uid + "\t", "password": enc, "lng": "1", "ignoreFlag": ig})
    if "INVALID_LOGIN_OR_PASSWORD" in r.text or "error_div1" in r.text or "EntityList" not in r.text:
        raise RuntimeError("MCX eclear refused the password — update it in Connections → MCX eclear "
                           "(not retried, to avoid locking the user).")
    tok = _token(ctx)
    code, phase, resp2, entity, codes = tok.now(), "current", None, None, None
    for att in range(1, 7):
        te, ti = _encrypt(code)
        resp2 = s.post(BASE_URL + "/Bancs/login.do?method=authenticateToken", headers=AJAX, timeout=60,
                       data={"userId": uid + "\t", "password": te, "lng": "1", "ignoreFlag": ti})
        try:
            entity, codes = _entity(resp2.text)
            ctx.log(f"RSA token accepted (attempt {att}).")
            break
        except Exception:
            if att == 6:
                raise RuntimeError("MCX eclear: the RSA code was rejected 6 times.")
            if phase == "current":
                code, phase = tok.at(datetime.now(timezone.utc) + timedelta(seconds=tok.interval)), "next"
            else:
                time.sleep(tok.time_left(datetime.now(timezone.utc)) + 2)
                code, phase = tok.now(), "current"
    fe, fi = _encrypt(pwd)
    r3 = s.post(BASE_URL + "/Bancs/login.do?method=Login", headers=LOGINH, timeout=60, data={
        "method": "Login", "availHeight": "744", "availWidth": "1522", "isArchival": "false",
        "userEntityCodes": codes, "lng": "1", "userId": uid, "ignoreFlag": fi, "ArchivalignoreFlag": "",
        "password": fe, "dummy": "FltM/sA2fPfqsjtM55oLOQ==", "entity": entity})
    m = re.search(r"name=['\"]rndaak['\"][^>]*value=['\"]([^'\"]+)", r3.text) or \
        re.search(r"rndaak['\"]?\s*:\s*['\"]([^'\"]+)", r3.text)
    if not m:
        raise RuntimeError("MCX eclear login did not finish (rndaak not found).")
    rnd = m.group(1)
    common = {"User-Agent": "Mozilla/5.0", "Content-Type": "application/x-www-form-urlencoded", "Origin": BASE_URL,
              "Referer": BASE_URL + "/Bancs/login.do?method=Login&availHeight=744&availWidth=1522",
              "Sec-Fetch-Site": "same-origin"}
    nav = {**common, "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
           "Sec-Fetch-Mode": "navigate", "Sec-Fetch-Dest": "iframe", "Upgrade-Insecure-Requests": "1"}
    aj2 = {**common, "X-Requested-With": "XMLHttpRequest", "Accept": "*/*", "Sec-Fetch-Mode": "cors",
           "Sec-Fetch-Dest": "empty", "Referer": BASE_URL + "/Bancs/common/welcome.jsp"}
    s.post(BASE_URL + "/Bancs/common/CommonImpl.jsp", headers=nav, data={"rndaak": rnd}, timeout=60)
    s.post(BASE_URL + "/Bancs/common/welcome.jsp", headers=nav, data={"rndaak": rnd}, timeout=60)
    s.post(BASE_URL + "/Bancs/common/layout.save", headers=aj2, data={"events": "", "rndaak": rnd}, timeout=60)
    s.post(BASE_URL + "/Bancs/common/welcome.jsp", headers=nav, data={"rndaak": rnd}, timeout=60)
    ctx.log("MCX eclear login OK.")
    return s, rnd


def fetch_rows(s, rnd, ctx):
    bs = H.need("bs4", "beautifulsoup4")
    RSK, TAB = BASE_URL + "/Bancs/RSK/RSK335.do", "panel2"
    s.post(RSK, headers={"User-Agent": "Mozilla/5.0", "Content-Type": "application/x-www-form-urlencoded",
                         "Origin": BASE_URL, "Referer": BASE_URL + "/Bancs/common/PageLoader.html"}, timeout=60, data={
        "app": "bancsapp?prefix=/RSK", "page": "/RSK335.do", "confViewSize": "50", "reqType": "0", "mode": "DEF_MODE",
        "parentName": "blank", "browserType": "IE55", "GuiWinName": "RSK335", "compname": "RSK",
        "over_ride_button": "Cancel", "AvlHt": "744", "AvlWd": "1481", "tblWd": "1465", "GuiBrowserInst": "0",
        "theme": "black", "pageId": "RSK335", "wintitle": "Margin Utilization View", "availHeight": "816",
        "availWidth": "1536", "tabId": TAB, "rndaak": rnd})
    payload = {
        "propertyMap(MCB_SearchWC_wca_bpid_Cmb)": "DFLT", "MCB_SearchWC_wca_bpid": "",
        "propertyMap(MCB_SearchWC_wca_bpid_op)": "=", "MCB_SearchWC_wca_actype": "",
        "propertyMap(MCB_SearchWC_wca_actype_op)": "=", "propertyMap(MCB_SearchWC_wca_associatedtm_Cmb)": "DFLT",
        "MCB_SearchWC_wca_associatedtm": "", "propertyMap(MCB_SearchWC_wca_associatedtm_op)": "=",
        "propertyMap(MCB_SearchWC_wca_associatedcm_Cmb)": "DFLT", "MCB_SearchWC_wca_associatedcm": "",
        "propertyMap(MCB_SearchWC_wca_associatedcm_op)": "=", "MCB_SearchWC_wca_valuedate": "",
        "propertyMap(MCB_SearchWC_wca_valuedate_op)": "=", "MCB_SearchWC_wca_category": "",
        "propertyMap(MCB_SearchWC_wca_category_op)": "=", "MCB_SearchWC_wca_CMName": "DFLT",
        "propertyMap(MCB_SearchWC_wca_CMName_op)": "=", "MCB_SearchWC_wca_TMName": "DFLT",
        "propertyMap(MCB_SearchWC_wca_TMName_op)": "=", "MCB_SearchWC_wca_datedummy1": "",
        "propertyMap(MCB_SearchWC_wca_datedummy1_op)": "=", "MCB_SearchWC_wca_datedummy2": "",
        "propertyMap(MCB_SearchWC_wca_datedummy2_op)": "=", "GRID_RESPONSE": "RSK335_Table", "operationType": "ET",
        "searchClicked": "true", "PageNum": "0", "service": "qryMarginUtilView", "windowName": "RSK335",
        "MODE": "DEF_MODE", "compName": "RSK", "lgbinst": "0", "WRITE_APP": "true", "REQ_TYPE": "0", "reqType": "1",
        "sQuery": "Client Code  Equals  DFLT AND TM / CP  Equals  DFLT AND CM  Equals  DFLT",
        "AvlWd": "1472", "tblWd": "1450",
        "MCBrowserEventInformation": "sourceElementName~|~Search#*#eventType~|~click#*#butId~|~null",
        "IXHRts": str(int(time.time() * 1000)), "XHR": "true", "rndaak": rnd, "theme": "black", "tabId": TAB,
        "mode": "DEF_MODE", "GuiBrowserInst": "0", "parentName": "blank", "browserType": "IE55",
        "GuiWinName": "RSK335", "compname": "RSK", "over_ride_button": "Cancel", "confViewSize": "50"}
    res = s.post(RSK, headers={"User-Agent": "Mozilla/5.0", "Accept": "*/*",
                               "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
                               "X-Requested-With": "XMLHttpRequest", "Origin": BASE_URL, "Referer": RSK},
                 data=payload, timeout=90)
    if "RSK335_Table" not in res.text:
        raise RuntimeError("RSK335 search failed: " + res.text[:200])
    marker = "RSK335_Table#*#"
    html = res.text.split(marker, 1)[1] if marker in res.text else res.text
    soup = bs.BeautifulSoup(html, "html.parser")
    table = soup.find("table", id="RSK335_Table") or soup.find("table")
    if not table:
        return []
    keep, heads = [], []
    for i, th in enumerate(table.select("thead th")):
        if "hide" in (th.get("class") or []):
            continue
        heads.append(th.get_text().replace("\xa0", " ").strip())
        keep.append(i)
    rows = []
    for tr in table.select("tbody tr"):
        tds = tr.find_all("td")
        if tds:
            rows.append({h: (tds[i].get_text().replace("\xa0", " ").strip() if i < len(tds) else "")
                         for h, i in zip(heads, keep)})
    return rows


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace(",", "").replace("\xa0", "").strip()
    if t in ("", "-", "&nbsp;", "nan", "NaN", "None"):
        return None
    try:
        return float(t)
    except ValueError:
        return None


# ───────────────────────── master workbook ─────────────────────────
def _border():
    from openpyxl.styles import Border, Side
    s = Side(border_style="thin", color="D0D7E2")
    return Border(top=s, bottom=s, left=s, right=s)


def ensure_master(path):
    if os.path.exists(path):
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Surveillance"
    n = len(FINAL_HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    t = ws.cell(1, 1, "MCX HOURLY SURVEILLANCE  —  MASTER CAPITAL SERVICES LTD")
    t.font, t.fill = Font(bold=True, size=15, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    s = ws.cell(2, 1, SUBTITLE)
    s.font, s.alignment = Font(italic=True, size=10, color="555555"), Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18
    for c, h in enumerate(FINAL_HEADERS, start=1):
        cell = ws.cell(3, c, h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEADP)
        cell.alignment, cell.border = Alignment(horizontal="center", vertical="center"), _border()
    ws.row_dimensions[3].height = 22
    for c, w in enumerate([20, 22, 20, 22, 20, 12], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


def _sep(v):
    return isinstance(v, str) and v.strip().startswith("─")


def append_master(path, time_hm, values):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    ensure_master(path)
    wb = load_workbook(path)
    ws = wb.active
    n = len(FINAL_HEADERS)
    today = datetime.now().date()
    ws.cell(2, 1).value = SUBTITLE
    last = None
    for r in range(ws.max_row, 3, -1):
        v = ws.cell(r, 1).value
        if _sep(v):
            try:
                last = datetime.strptime(v.strip().strip("─").strip(), "%A, %d %B %Y").date()
            except ValueError:
                pass
            break
    if last != today:
        r = ws.max_row + 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        sep = ws.cell(r, 1, f"─────  {datetime.now().strftime('%A, %d %B %Y')}  ─────")
        sep.font, sep.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=GOLD)
        sep.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 20
    cnt = 0
    for r in range(ws.max_row, 3, -1):
        v = ws.cell(r, 1).value
        if _sep(v):
            break
        if v not in (None, ""):
            cnt += 1
    r = ws.max_row + 1
    for c, val in enumerate([time_hm] + [values[h] for h in FINAL_HEADERS[1:]], start=1):
        cell = ws.cell(r, c, val)
        cell.border = _border()
        if cnt % 2:
            cell.fill = PatternFill("solid", fgColor=ALT)
        if c == 1:
            cell.font, cell.alignment = Font(bold=True, color="222222"), Alignment(horizontal="center", vertical="center")
        elif c == n:
            pct = val if isinstance(val, (int, float)) else to_num(val)
            col = NAVY if pct is None else RED if pct >= 90 else AMBER if pct >= 80 else GREEN
            cell.font, cell.alignment = Font(bold=True, color=col), Alignment(horizontal="center", vertical="center")
            cell.number_format = "#,##0.00"
        else:
            cell.alignment, cell.number_format = Alignment(horizontal="right", vertical="center"), "#,##0.00"
    wb.save(path)


def todays_rows(path):
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    sep_row = 3
    for r in range(ws.max_row, 3, -1):
        if _sep(ws.cell(r, 1).value):
            sep_row = r
            break
    out = []
    for r in range(sep_row + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or _sep(v):
            continue
        out.append([ws.cell(r, c).value for c in range(1, len(FINAL_HEADERS) + 1)])
    return out


# ───────────────────────── image (same look as the original) ─────────────────────────
@lru_cache(maxsize=64)
def _font(sz, bold=False):
    from PIL import ImageFont
    sz = max(1, int(round(sz)))
    cands = ([r"C:\Windows\Fonts\segoeuib.ttf", r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\calibrib.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"] if bold else
             [r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"])
    for p in cands:
        try:
            return ImageFont.truetype(p, sz)
        except Exception:
            pass
    try:
        return ImageFont.load_default(size=sz)
    except Exception:
        return ImageFont.load_default()


def _fmt(v):
    if v is None or v == "":
        return "-"
    return f"{v:,.2f}" if isinstance(v, (int, float)) else str(v)


def build_image(stub, rows, ctx):
    from PIL import Image, ImageDraw
    hx = lambda h: "#" + h
    S = IMG_SCALE
    col_w = [w * S for w in (210, 235, 225, 235, 215, 150)]
    W, pad = sum(col_w), 28 * S
    title_h, sub_h, head_h, row_h, foot_h = 88 * S, 40 * S, 54 * S, 48 * S, 54 * S
    n = max(len(rows), 1)
    Ht = pad + title_h + sub_h + head_h + row_h * n + foot_h + pad
    img = Image.new("RGB", (W + pad * 2, Ht), "white")
    d = ImageDraw.Draw(img)
    x0, y = pad, pad
    d.rectangle([x0, y, x0 + W, y + title_h], fill=hx(NAVY))
    d.rectangle([x0, y + title_h - 6 * S, x0 + W, y + title_h], fill=hx(GOLD))
    d.text((x0 + 24 * S, y + (title_h - 6 * S) // 2), "MCX HOURLY SURVEILLANCE", font=_font(34 * S, True), fill="white", anchor="lm")
    y += title_h
    d.rectangle([x0, y, x0 + W, y + sub_h], fill="#F4F6F9")
    d.text((x0 + 24 * S, y + sub_h // 2), f"Margin Utilization (in Crores)  |  {datetime.now().strftime('%A, %d %B %Y')}",
           font=_font(18 * S, True), fill=hx(NAVY), anchor="lm")
    y += sub_h
    cx = x0
    for i, h in enumerate(FINAL_HEADERS):
        d.rectangle([cx, y, cx + col_w[i], y + head_h], fill=hx(HEADP))
        d.text((cx + col_w[i] // 2, y + head_h // 2), h, font=_font(16 * S, True), fill="white", anchor="mm")
        cx += col_w[i]
    y += head_h
    for ri in range(n):
        row = rows[ri] if ri < len(rows) else None
        d.rectangle([x0, y, x0 + W, y + row_h], fill=hx(ALT) if ri % 2 else "white")
        if row is not None:
            cx = x0
            for i in range(len(FINAL_HEADERS)):
                val = row[i] if i < len(row) else ""
                f, color = _font(16 * S, i == 0), "black"
                if i == len(FINAL_HEADERS) - 1:
                    pv = val if isinstance(val, (int, float)) else to_num(val)
                    if pv is not None:
                        color = hx(RED) if pv >= 90 else hx(AMBER) if pv >= 80 else hx(GREEN)
                        f = _font(17 * S, True)
                    txt = _fmt(val)
                else:
                    txt = str(val) if i == 0 else _fmt(val)
                d.text((cx + col_w[i] // 2, y + row_h // 2), txt, font=f, fill=color, anchor="mm")
                cx += col_w[i]
        cx = x0
        for i in range(len(FINAL_HEADERS) - 1):
            cx += col_w[i]
            d.line([cx, y, cx, y + row_h], fill="#D0D7E2", width=max(1, S // 2))
        y += row_h
    d.rectangle([x0, y, x0 + W, y + foot_h], fill=hx(NAVY))
    d.text((x0 + 24 * S, y + foot_h // 2), f"Generated {datetime.now().strftime('%d-%m-%Y %H:%M')} IST",
           font=_font(14 * S), fill="white", anchor="lm")
    d.rectangle([x0, pad, x0 + W, Ht - pad], outline=hx(NAVY), width=2 * S)
    png, jpg = stub.with_suffix(".png"), stub.with_suffix(".jpg")
    img.save(png, format="PNG", optimize=True)
    q = JPEG_QUALITY
    while True:
        img.save(jpg, format="JPEG", quality=q, subsampling=0, optimize=True, dpi=(144, 144))
        if jpg.stat().st_size <= WA_MAX_BYTES or q <= 70:
            break
        q -= 5
    ctx.log(f"Image built {img.size[0]}x{img.size[1]} px.")
    return png, jpg


# ───────────────────────── the process ─────────────────────────
@H.logged(FOLDER)
def _run(params, ctx):
    H.need("openpyxl"), H.need("PIL", "pillow")
    to, cc = K.recipients(params)
    phones = H.parse_phones(params.get("phones"))
    now = datetime.now()
    first, last = _hm(params.get("first_run"), "09:00"), _hm(params.get("last_run"), "00:30")
    cut = _hm(params.get("saturday_cutoff"), "04:00")
    if not _in_window(now, first, last, cut):
        raise H.Skipped(f"{now:%H:%M} is outside the run window ({params.get('first_run') or '09:00'}–"
                        f"{params.get('last_run') or '00:30'}, Sundays never) — no MCX login attempted.")
    out, tmp = paths.proc_dir(FOLDER, "outputs"), paths.proc_dir(FOLDER, "inputs")
    master = out / f"MAIN_MCX_HOURLY_SURVEILLANCE_{now:%B}_{now:%Y}.xlsx".upper().replace("XLSX", "xlsx")

    s, rnd = login(ctx)
    rows = fetch_rows(s, rnd, ctx)
    if not rows:
        raise RuntimeError("The Margin Utilization search returned no rows.")
    rec = next((r for r in rows if r.get("TM / CP Code", "").strip() == MEMBER_ROW), rows[0])
    raw = tmp / f"RAW_MCX_{now:%Y%m%d_%H%M%S}.csv"
    with open(raw, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(list(rec.keys()))
        w.writerow(list(rec.values()))

    values = {}
    for final, src in COLUMN_MAP:
        num = to_num(rec.get(src, ""))
        if num is not None and final != "Used by %":
            num /= 1e7
        values[final] = num if num is not None else rec.get(src, "")
    ctx.log("Values: " + " | ".join(f"{k}={_fmt(v)}" for k, v in values.items()))
    time_hm = now.strftime("%H:%M")
    append_master(master, time_hm, values)
    png, jpg = build_image(tmp / f"MCX_SNAPSHOT_{now:%Y%m%d_%H%M%S}", todays_rows(master), ctx)

    used = _fmt(values["Used by %"])
    body = ('<div style="font-family:Segoe UI,Arial,sans-serif;color:#1a1a1a;font-size:14px;"><p>Dear Sir/Madam,</p>'
            f"<p>Please find below the <b>MCX Margin Utilization</b> snapshot captured at <b>{now:%d-%m-%Y} {time_hm} IST</b>. "
            f"The current margin usage stands at <b>{used}%</b>. The running master workbook is attached for reference.</p>"
            '<img src="cid:snapshot" alt="MCX Hourly Surveillance" style="display:block;max-width:100%;border:1px solid '
            '#d0d7e2;border-radius:4px;margin:12px 0;"/>'
            f'<p style="margin-top:16px;">Best regards,<br>{K.sig(params)}</p></div>')
    H.send_email("MCX Hourly Surveillance", body, to, cc, [master], ctx, html=True, inline_images={"snapshot": png})
    failed = []
    for ph in phones:
        try:
            H.whatsapp_send_image(jpg, ph, "MCX HOURLY SURVEILLANCE", f"{now:%d-%m-%Y %H:%M}", ctx)
        except Exception as e:
            failed.append(ph[-4:])
            ctx.log(f"WhatsApp to …{ph[-4:]} failed: {e}")

    day = paths.proc_dir(FOLDER, "backup") / H.today_dash()
    day.mkdir(parents=True, exist_ok=True)
    n = max([int(x.name) for x in day.iterdir() if x.is_dir() and x.name.isdigit()] or [0]) + 1
    run = day / f"{n:02d}"
    run.mkdir()
    shutil.copy2(master, run / master.name)
    for f in (png, jpg, raw):
        if f.is_file():
            f.replace(run / f.name)
    if failed:
        raise RuntimeError(f"E-mail sent, but WhatsApp failed for number(s) ending {', '.join(failed)}.")
    return f"snapshot {time_hm} sent — margin used {used}%"


register(
    key="mcx_hourly_surveillance",
    name="MCX HOURLY SURVEILLANCE",
    description="Hourly: reads MCX margin utilisation (eclear login with RSA token), updates the master workbook, "
                "builds the snapshot image and sends it by e-mail and WhatsApp.",
    run=_run,
    params=K.mail_params("rks@mastertrust.co.in",
                         "va@mastertrust.co.in,hs@mastertrust.co.in,sandeep.kumar2@mastertrust.co.in,sukhbirsingh@mastertrust.co.in,"
                         "sumit.aggarwal@mastertrust.co.in,maninders@mastertrust.co.in,rms_del@mastertrust.co.in", with_wait=False) + [
        Param("phones", "WhatsApp numbers", "text", fmt="phones", scope="process",
              help="10-digit numbers, optional #name after each."),
        Param("first_run", "Run only from (HH:MM)", "str", default="09:00"),
        Param("last_run", "…until (HH:MM, may be after midnight)", "str", default="00:30"),
        Param("saturday_cutoff", "Saturday carry-over until (HH:MM)", "str", default="04:00"),
    ],
    retry_safe=False, parallel_safe=False, category="Surveillance", folder=FOLDER,
    uses=("mcx_eclear", "whatsapp", "smtp"),
)
