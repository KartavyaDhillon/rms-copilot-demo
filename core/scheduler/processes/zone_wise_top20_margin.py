"""ZONE WISE TOP 20 MARGIN — collects the day's "Zone wise Client margin <90%" snapshot mails (ClientMargin_FO_*.xlsx,
before the cut-off) from the second mailbox, ranks the top 20 unique clients per zone by Net Cash Supported FROM PRO and
e-mails a password-protected Excel (from CLIENT MARGIN TOP 20\\new.py).  Schedule it after the cut-off (e.g. 16:00)."""
import email
import io
import poplib
import time
from collections import defaultdict
from datetime import datetime, time as dtime, timedelta, timezone
from email import policy
from email.parser import BytesParser

from core.scheduler import connections, helpers as H, paths, partic_lib
from core.scheduler.catalog import Param, register
from core.scheduler.processes import _kit as K

FOLDER = "ZONE WISE TOP 20 MARGIN"
IST = timezone(timedelta(hours=5, minutes=30))
ZONES = ["DELHI", "MUMBAI", "LUDHIANA", "CHANDIGARH"]
PREFIX = "ClientMargin_FO"


def _zone(raw):
    return {"ZOMUMB": "MUMBAI", "ZOLUDH": "LUDHIANA", "ZOCHND": "CHANDIGARH"}.get(raw, "DELHI")


def _hhmm(s, d):
    try:
        h, m = (s or d).strip().split(":")
        return dtime(int(h), int(m))
    except Exception:
        raise ValueError(f"'{s}' is not a time like 15:55")


def fetch(subject_filter, cutoff, ctx):
    """[(received_ist, name, bytes)] earliest first — today's matching mails only."""
    c = connections.require("pop3b", "server", "port", "username", "password")
    target = datetime.now(IST).date()
    pop = poplib.POP3_SSL(c["server"], int(c["port"]), context=partic_lib._relaxed_ssl_context(), timeout=30)
    out = []
    try:
        pop.user(c["username"])
        pop.pass_(c["password"])
        count, _ = pop.stat()
        for i in range(count, max(0, count - 3000), -1):          # newest → oldest, headers first
            try:
                _, lines, _ = pop.top(i, 0)
            except Exception:
                continue
            msg = BytesParser(policy=policy.default).parsebytes(b"\r\n".join(lines))
            try:
                rdt = email.utils.parsedate_to_datetime(msg["Date"])
            except Exception:
                continue
            rdt = (rdt.replace(tzinfo=IST) if rdt.tzinfo is None else rdt).astimezone(IST)
            if rdt.date() < target:
                break
            if rdt.date() > target or subject_filter.lower() not in (msg["Subject"] or "").lower():
                continue
            if rdt.time() >= cutoff:
                continue
            _, full, _ = pop.retr(i)
            m = BytesParser(policy=policy.default).parsebytes(b"\r\n".join(full))
            for part in m.iter_attachments():
                fn = part.get_filename() or ""
                if fn.upper().startswith(PREFIX.upper()) and fn.lower().endswith(".xlsx"):
                    out.append((rdt, fn, part.get_payload(decode=True)))
                    break
    finally:
        try:
            pop.quit()
        except Exception:
            pass
    out.sort(key=lambda x: x[0])
    return out


@H.logged(FOLDER)
def _run(params, ctx):
    openpyxl = H.need("openpyxl")
    msoff = H.need("msoffcrypto", "msoffcrypto-tool")
    to, cc = K.recipients(params)
    if H.already_sent(FOLDER, "mail", ctx):
        raise H.Skipped(H.skip_message(FOLDER, ["mail"], "e-mailed"))
    pw = connections.require("client_margin", "source_password", "output_password")
    src_pw, out_pw = pw["source_password"], pw["output_password"]
    sub = (params.get("subject_filter") or "Zone wise Client margin <90%").strip()
    cutoff = _hhmm(params.get("cutoff"), "15:55")
    need_n = max(1, int(params.get("min_files") or 4))
    wait_s = K.wait_of(params) * 60
    t0 = time.time()
    while True:
        found = fetch(sub, cutoff, ctx)
        ctx.log(f"{len(found)} matching mail(s) before {cutoff:%H:%M} (need {need_n}).")
        if len(found) >= need_n or time.time() - t0 >= wait_s or getattr(ctx, "cancelled", False):
            break
        time.sleep(60)
    if len(found) < need_n and not params.get("proceed_on_shortfall"):
        raise RuntimeError(f"Only {len(found)} snapshot mail(s) found, {need_n} needed — stopped. "
                           "Tick 'Proceed with fewer files' to allow a shorter report.")
    if not found:
        raise FileNotFoundError("No matching snapshot mail found today.")

    inp, out = paths.proc_dir(FOLDER, "inputs"), paths.proc_dir(FOLDER, "outputs")
    files, per_file = [], {}
    for seq, (rdt, name, data) in enumerate(found, 1):
        p = inp / f"{name.rsplit('.', 1)[0]}_{seq:02d}.xlsx"
        p.write_bytes(data)
        files.append((p, rdt))
        of = msoff.OfficeFile(io.BytesIO(data))
        of.load_key(password=src_pw)
        buf = io.BytesIO()
        of.decrypt(buf)
        buf.seek(0)
        ws = openpyxl.load_workbook(buf, data_only=True)["Sheet1"]
        zr = defaultdict(list)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] is None or row[2] is None:
                continue
            z = _zone(row[1])
            zr[z].append({"cm": row[0], "raw": row[1], "mz": z, "client": row[2], "net": row[14] or 0, "src": p.name})
        per_file[p.name] = {z: sorted(zr.get(z, []), key=lambda r: r["net"], reverse=True)[:20] for z in ZONES}

    final = {}
    for z in ZONES:
        best = {}
        for fr in per_file.values():
            for r in fr[z]:
                if r["client"] not in best or r["net"] > best[r["client"]]["net"]:
                    best[r["client"]] = r
        final[z] = sorted(best.values(), key=lambda r: r["net"], reverse=True)[:20]

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    HF, ZF = PatternFill("solid", start_color="1F4E78", end_color="1F4E78"), PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    disp = {p.name: f"{p.name} ({rdt:%I:%M %p})" for p, rdt in files}

    def block(ws, r, zone, rows, heads, vals, vcol):
        n = len(heads)
        for c in range(1, n + 1):
            ws.cell(r, c).fill = ZF
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        zc = ws.cell(r, 1, f"ZONE: {zone}")
        zc.font, zc.alignment = Font(name="Arial", bold=True, size=11), Alignment(horizontal="center", vertical="center")
        r += 1
        for c, h in enumerate(heads, 1):
            x = ws.cell(r, c, h)
            x.font, x.fill, x.border = Font(name="Arial", bold=True, color="FFFFFF", size=11), HF, bd
            x.alignment = Alignment(horizontal="center", vertical="center")
        r += 1
        for i, row in enumerate(rows, 1):
            for c, v in enumerate(vals(i, row), 1):
                x = ws.cell(r, c, v)
                x.font, x.border = Font(name="Arial", size=10), bd
                if c == vcol:
                    x.number_format = "#,##0"
            r += 1
        return r + 1

    def fit(ws):
        w = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    w[c.column_letter] = max(w.get(c.column_letter, 0), len(str(c.value)))
        for k, v in w.items():
            ws.column_dimensions[k].width = max(8, min(45, v + 4))

    wb = Workbook()
    ws = wb.active
    ws.title = "Final Top20 per Zone"
    r = 1
    for z in ZONES:
        r = block(ws, r, z, final[z], ["Rank", "Zone", "CM Code", "Client Code", "Net Cash Supported FROM PRO",
                                       "Source File & Snapshot Time"],
                  lambda i, x: (i, x["raw"], x["cm"], x["client"], round(x["net"]), disp.get(x["src"], "")), 5)
    fit(ws)
    for idx, (p, _) in enumerate(files, 1):
        w2 = wb.create_sheet(f"File{idx} Top20 per Zone")
        r = 1
        for z in ZONES:
            r = block(w2, r, z, per_file[p.name][z], ["Rank", "Zone", "Zone Code (raw)", "CM Code", "Client Code",
                                                     "Net Cash Supported FROM PRO"],
                      lambda i, x: (i, x["mz"], x["raw"], x["cm"], x["client"], round(x["net"])), 6)
        fit(w2)

    day = datetime.now(IST)
    name = f"Top20_Client_Margin_Report_{day:%d-%m-%Y}_{day:%H%M}.xlsx"
    plain, prot = out / (name + ".tmp"), out / name
    wb.save(plain)
    with open(plain, "rb") as f:
        msoff.OfficeFile(f).encrypt(password=out_pw, outfile=open(prot, "wb"))
    plain.unlink(missing_ok=True)

    body = f"Respected Sir,\n\nPlease find attached the Zone-wise UNIQUE Top 20 Client Margin report for {day:%d-%m-%Y}.\n\nBest regards,\n{K.sig_text(params)}"
    H.send_email(f"Zone Wise UNIQUE Top 20 Client Margin Report - {day:%d-%m-%Y}", body, to, cc, [prot], ctx, group="smtpb")
    H.mark_sent(FOLDER, "mail")
    H.backup_files(FOLDER, [prot] + [p for p, _ in files], ctx)
    return f"report e-mailed from {len(files)} snapshot file(s)"


register(
    key="zone_wise_top20_margin",
    name="ZONE WISE TOP 20 MARGIN",
    description="Collects today's 'Zone wise Client margin <90%' snapshot mails, ranks the top 20 unique clients per zone "
                "and e-mails a password-protected Excel. Uses the second mailbox (Connections).",
    run=_run,
    params=K.mail_params("rks@mastertrust.co.in",
                         "va@mastertrust.co.in,deepak01@mastertrust.co.in,sumit.aggarwal@mastertrust.co.in,sandeep.kumar2@mastertrust.co.in,"
                         "vipin.rai@mastertrust.co.in,priyanshu.karna@mastertrust.co.in") + [
        Param("subject_filter", "Mail subject contains", "str", default="Zone wise Client margin <90%"),
        Param("cutoff", "Only mails before (HH:MM)", "str", default="15:55"),
        Param("min_files", "Snapshot mails needed", "int", default=4),
        Param("proceed_on_shortfall", "Proceed with fewer files", "bool", default=False,
              help="Off = stop (run fails) when fewer than the needed mails arrived."),
    ],
    retry_safe=False, parallel_safe=False, category="Reports", folder=FOLDER, uses=("pop3b", "smtpb", "client_margin"),
)
