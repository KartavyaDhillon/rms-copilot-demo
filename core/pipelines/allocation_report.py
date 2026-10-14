"""Allocation report — console engine around ALLOCATION_REPORT.py.

The proven script (mail fetch, parsing, opening balances, report) is NOT rewritten: copy it, unchanged, next to this
file as   core/pipelines/allocation_report_src.py   (or point ALLOC_REPORT_SCRIPT at it).  This module loads it,
re-points its folders into the console's data folder, and drives it in four modes:

    close   nightly (Scheduler): window = last close  ->  now.  E-mails, backs up, and moves the day marker.
    sofar   manual:  today 00:00 -> now
    until   manual:  today 00:00 -> HH:MM
    custom  manual:  any date+time  ->  any date+time (past days allowed)

Manual modes NEVER move the nightly marker.  They can be limited to some client codes, and are either downloaded
(files kept in <folder>/Manual/<stamp>/) or e-mailed to the addresses given.

Run as a subprocess:   python -c "from core.pipelines.allocation_report import run_cli; run_cli('<json options>')"
Lines starting "[FRIENDLY_MSG]" show in the console strip; "ALLOC_RESULT {...}" is the machine-readable result.

Optional settings in .env (otherwise the values inside the copied script are used):
    ALLOC_REPORT_FOLDER      working / backup / state folder (default: <console>/data/scheduler/ALLOCATION REPORT)
    ALLOC_REPORT_SCRIPT      full path of the copied script
    ALLOC_POP_SERVER / ALLOC_POP_PORT / ALLOC_POP_USER / ALLOC_POP_PASS      mailbox the allocation logs arrive in
    ALLOC_SMTP_SERVER / ALLOC_SMTP_PORT / ALLOC_SMTP_USER / ALLOC_SMTP_PASS  mailbox the report is sent from
    ALLOC_MAIL_TO / ALLOC_MAIL_CC                                            nightly recipients (comma separated)
    CLI_FOLDER               the CLI box's folder: opening balances are read from <CLI_FOLDER>\\BACKUPFILES AUTOMATED
"""
from __future__ import annotations

import datetime as dt
import importlib.util
import json
import os
import re
import shutil
import sys
import time
import traceback
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SRC_DEFAULT = Path(__file__).with_name("allocation_report_src.py")
LOCK_MAX_AGE = 4 * 3600
STEPS = 5
_MOD = None


# ───────────────────────────── folders / state ─────────────────────────────
def base_dir() -> Path:
    v = os.getenv("ALLOC_REPORT_FOLDER", "").strip()
    if v:
        return Path(v)
    try:
        from core.scheduler import paths
        return Path(paths.sched_root()) / "ALLOCATION REPORT"
    except Exception:
        return REPO / "data" / "scheduler" / "ALLOCATION REPORT"


def src_path() -> Path:
    v = os.getenv("ALLOC_REPORT_SCRIPT", "").strip()
    return Path(v) if v else SRC_DEFAULT


def cli_backup_root() -> str:
    return os.path.join(os.getenv("CLI_FOLDER", r"D:\CLI_PRUDENTIAL_LIMIT"), "BACKUPFILES AUTOMATED")


def _jread(p: Path):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None


def _jwrite(p: Path, data):
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    except Exception:
        pass


def log(msg=""):
    print(msg, flush=True)


def progress(step, msg, ok=None):
    _jwrite(base_dir() / "progress.json",
            {"step": step, "total": STEPS, "msg": msg, "ok": ok, "at": dt.datetime.now().isoformat(timespec="seconds")})


# ───────────────────────────── lock (nightly vs manual) ─────────────────────────────
def _lock_file() -> Path:
    return base_dir() / ".run.lock"


def lock_info():
    p = _lock_file()
    try:
        if p.is_file() and time.time() - p.stat().st_mtime < LOCK_MAX_AGE:
            return _jread(p) or {"mode": "?", "by": ""}
    except OSError:
        pass
    return None


def _acquire(info: dict) -> bool:
    if lock_info():
        return False
    _jwrite(_lock_file(), {**info, "pid": os.getpid(), "since": dt.datetime.now().isoformat(timespec="seconds")})
    return True


def _release():
    try:
        _lock_file().unlink()
    except OSError:
        pass


# ───────────────────────────── status (for the UI) ─────────────────────────────
def status() -> dict:
    b = base_dir()
    now = dt.datetime.now()
    lk = lock_info()
    marker = ""
    ws = _jread(b / "window_state.json")
    if isinstance(ws, dict):
        marker = str(ws.get("last_window_end", ""))
    return {
        "ok": True, "script_ok": src_path().is_file(), "script": str(src_path()), "folder": str(b),
        "running": bool(lk), "lock": lk or {}, "progress": _jread(b / "progress.json") or {},
        "marker": marker, "now": now.isoformat(timespec="seconds"), "today": now.strftime("%Y-%m-%d"),
        "last_close": _jread(b / "last_close.json") or {}, "last_manual": _jread(b / "last_manual.json") or {},
    }


def served_file(name: str):
    """A file of the last manual run (the UI's download buttons)."""
    last = _jread(base_dir() / "last_manual.json") or {}
    d = last.get("dir")
    if not d or name not in (last.get("files") or []):
        return None
    p = Path(d) / os.path.basename(name)
    return p if p.is_file() else None


# ───────────────────────────── loading the proven script ─────────────────────────────
def _load():
    global _MOD
    if _MOD is not None:
        return _MOD
    p = src_path()
    if not p.is_file():
        raise FileNotFoundError(
            f"The allocation script is not in the console yet. Copy ALLOCATION_REPORT.py to {p} and run again.")
    spec = importlib.util.spec_from_file_location("alloc_report_src", str(p))
    m = importlib.util.module_from_spec(spec)
    sys.modules["alloc_report_src"] = m
    spec.loader.exec_module(m)
    _MOD = m
    return m


def _env_list(key):
    return [x.strip() for x in re.split(r"[,;\s]+", os.getenv(key, "")) if x.strip()]


def _configure(m, base: Path):
    """Point every folder the script uses into the console's folder (functions read these names at call time)."""
    b = str(base)
    m.BASE_FOLDER = m.BASE_DIR = b
    m.WORK_FOLDER = os.path.join(b, "analysis_work")
    m.DOWNLOAD_DIR = os.path.join(b, "downloaded_files")
    m.ALLOC_CATEGORY_DIR = os.path.join(m.DOWNLOAD_DIR, "ALLOCATION LOGS")
    m.DOWNUP_CATEGORY_DIR = os.path.join(m.DOWNLOAD_DIR, "DOWNWARDS_UPWARDS LOGS")
    m.MCX_CATEGORY_DIR = os.path.join(m.DOWNLOAD_DIR, "MCX ALLOCATION LOGS")
    m.OUTPUT_DIR = os.path.join(b, "output_files")
    m.BACKUP_DIR = os.path.join(b, "BACKUPFILES AUTOMATED")
    m.STATE_FILE = os.path.join(b, "processed_uidls.json")
    m.WINDOW_STATE_FILE = os.path.join(b, "window_state.json")
    m.OPEN_WINDOW_FILE = os.path.join(b, "open_window_uidls.json")
    m.UNPARSED_LOG = os.path.join(b, "unparsed_lines.log")
    m.RUN_LOG = os.path.join(b, "run_log.txt")
    m.EMAIL_MAPPING_XLSX = os.path.join(b, "Email_Mapping_Log.xlsx")
    m._CATEGORY_DIR = {"MASTER": m.ALLOC_CATEGORY_DIR, "MASTERTRUST": m.DOWNUP_CATEGORY_DIR, "MCX": m.MCX_CATEGORY_DIR}
    m.BACKUP_ROOT = cli_backup_root()
    g = os.getenv
    if g("ALLOC_POP_SERVER"):
        m.POP_SERVER = g("ALLOC_POP_SERVER").strip()
    if g("ALLOC_POP_PORT"):
        m.POP_PORT = int(g("ALLOC_POP_PORT"))
    if g("ALLOC_POP_USER"):
        m.EMAIL_ID = g("ALLOC_POP_USER").strip()
    if g("ALLOC_POP_PASS"):
        m.PASSWORD = g("ALLOC_POP_PASS")
    if g("ALLOC_SMTP_SERVER"):
        m.SMTP_SERVER = g("ALLOC_SMTP_SERVER").strip()
    if g("ALLOC_SMTP_PORT"):
        m.SMTP_PORT = int(g("ALLOC_SMTP_PORT"))
    if g("ALLOC_SMTP_USER"):
        m.SMTP_EMAIL = g("ALLOC_SMTP_USER").strip()
    if g("ALLOC_SMTP_PASS"):
        m.SMTP_PASSWORD = g("ALLOC_SMTP_PASS")
    if _env_list("ALLOC_MAIL_TO"):
        m.EMAIL_TO = _env_list("ALLOC_MAIL_TO")
        m.EMAIL_CC = _env_list("ALLOC_MAIL_CC")
    os.makedirs(b, exist_ok=True)


# ───────────────────────────── helpers ─────────────────────────────
def _dt(s: str) -> dt.datetime:
    return dt.datetime.strptime(s.strip(), "%Y-%m-%d %H:%M")


def _clients(v) -> list:
    out = []
    for c in re.split(r"[\s,;]+", str(v or "")):
        c = c.strip().upper()
        if c and re.fullmatch(r"[A-Z0-9_\-]{1,20}", c) and c not in out:
            out.append(c)
    return out


def _filter_xlsx(path: str, keep: set) -> int:
    """Keep only the rows of these client codes (the file's "ClientCode" column). Returns rows kept (file removed at 0)."""
    from openpyxl import Workbook, load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header or "ClientCode" not in header:
        wb.close()
        return -1
    idx = list(header).index("ClientCode")
    kept = [list(r) for r in it if r and idx < len(r) and str(r[idx] or "").strip().upper() in keep]
    wb.close()
    if not kept:
        os.remove(path)
        return 0
    out = Workbook()
    w = out.active
    w.title = "Data"
    w.append(list(header))
    for i, h in enumerate(header, start=1):
        w.column_dimensions[w.cell(row=1, column=i).column_letter].width = max(14, len(str(h)) + 2)
    for r in kept:
        w.append(r)
    out.save(path)
    return len(kept)


def _opening_date(mode: str, start: dt.datetime, end: dt.datetime) -> dt.date:
    if mode == "close":
        return (end if end.hour >= 6 else end - dt.timedelta(days=1)).date()
    return start.date()


# ───────────────────────────── the run ─────────────────────────────
def run(opts: dict) -> int:
    mode = (opts.get("mode") or "").strip()
    by = (opts.get("by") or "").strip()
    base = base_dir()
    base.mkdir(parents=True, exist_ok=True)
    if mode not in ("close", "sofar", "until", "custom"):
        log(f"[FRIENDLY_MSG] Unknown mode '{mode}'.")
        return 1
    if not _acquire({"mode": mode, "by": by}):
        who = lock_info() or {}
        log(f"[FRIENDLY_MSG] Another allocation report run is in progress ({who.get('mode', '?')}"
            f"{', started by ' + who['by'] if who.get('by') else ''}). Wait for it to finish.")
        return 1
    m = None
    work = None
    try:
        progress(1, "Reading the mailbox")
        m = _load()
        _configure(m, base)
        work = m.WORK_FOLDER
        m.setup_logging()
        now = dt.datetime.now()
        log("=" * 70)
        log(f"  ALLOCATION REPORT  mode={mode}  started {now:%d-%b-%Y %H:%M:%S}" + (f"  by {by}" if by else ""))
        log("=" * 70)

        # ---- the window -------------------------------------------------------
        if mode == "close":
            start = m.get_window_state() or dt.datetime.combine(now.date(), dt.time.min)
            end = _dt(opts["as_of"]) if opts.get("as_of") else now
            if opts.get("as_of") and start >= end:
                start = dt.datetime.combine(end.date(), dt.time.min)
                log(f"  [!] marker is already past {end:%d-%b %H:%M}; re-closing {end:%d-%b-%Y} from 00:00.")
        elif mode == "sofar":
            start, end = dt.datetime.combine(now.date(), dt.time.min), now
        elif mode == "until":
            hh, mm = str(opts.get("until", "")).split(":")
            start = dt.datetime.combine(now.date(), dt.time.min)
            end = dt.datetime.combine(now.date(), dt.time(int(hh), int(mm)))
        else:
            start, end = _dt(opts["from"]), _dt(opts["range_to"])
        if end > now:
            end = now
        if start >= end:
            log("[FRIENDLY_MSG] The start of the range is not before its end.")
            progress(1, "Invalid range", ok=False)
            return 1
        clients = _clients(opts.get("clients"))
        keep = set(clients)
        log(f"  window : {start:%d-%b-%Y %H:%M}  ->  {end:%d-%b-%Y %H:%M}")
        if clients:
            log(f"  clients: {', '.join(clients)}")
        if mode == "close" and start.date() != end.date() and (end - start) > dt.timedelta(hours=30):
            log("  [!] this window spans more than one day (a night was missed?) — opening balances are those of the end day.")

        log_offsets = m.log_start_offsets() if mode == "close" else None
        m.clean_folder(m.WORK_FOLDER)

        # ---- STEP 1: fetch the mails and build the three log files ---------------
        session = m._Pop3Session()
        try:
            matches = m.fetch_window_messages(session, start, end)
        finally:
            session.close()
        bucket = m.bucket_live(matches, start, end)
        label = "EOD" if mode == "close" else f"Intraday {end:%H:%M}"
        log(f"\n  emails in window: {len(bucket)}")
        if not bucket:
            msg = "No allocation e-mails in this window."
            log(f"[FRIENDLY_MSG] {msg}")
            _finish(base, mode, {"ok": True, "empty": True, "msg": msg}, start, end, clients, by, [], None)
            progress(STEPS, msg, ok=True)
            return 0

        progress(2, f"Building the log files ({len(bucket)} e-mails)")
        grabbed = {}
        real_send = m.send_email_with_attachments

        def capture(subject, body, paths):
            grabbed["files"] = []
            for p in paths:
                dest = os.path.join(m.WORK_FOLDER, os.path.basename(p))
                shutil.copy2(p, dest)
                grabbed["files"].append(dest)

        m.send_email_with_attachments = capture
        try:
            m.process_bucket(bucket, start, end, label, permanent=False)
        finally:
            m.send_email_with_attachments = real_send
        produced = grabbed.get("files", [])
        pick = lambda pre: next((p for p in produced if os.path.basename(p).startswith(pre)), None)
        alloc_log, downwards, upwards = pick("Allocation_Logs"), pick("DOWNWARDS_Logs"), pick("UPWARDS_Logs")
        if not alloc_log:
            raise RuntimeError("No Allocation_Logs file was produced, so the report can't be built.")

        # ---- STEP 2: opening balances (from the CLI box's backup) ------------------
        progress(3, "Collecting opening balances")
        opening_used = []
        odate = _opening_date(mode, start, end)
        log(f"\n  opening balances: looking for the {odate:%d-%m-%Y} folder in {m.BACKUP_ROOT}")
        dfolder = m.find_date_folder(odate)
        if dfolder:
            srcf = m.oldest_time_subfolder(dfolder)
            if srcf:
                for p in m.pick_opening_files(srcf).values():
                    shutil.copy2(p, os.path.join(m.WORK_FOLDER, os.path.basename(p)))
                    opening_used.append(p)
        if not opening_used:
            log("[FRIENDLY_MSG] No opening-balance files found for this day — they are treated as 0.")

        # ---- STEP 3: the report ------------------------------------------------------
        progress(4, "Building the report")
        all_rows, log_date = m.read_log(alloc_log)
        if not all_rows:
            raise RuntimeError("No data rows found in the log file.")
        missing_clients = []
        if keep:
            seen = {str(r.get("ClientCode", "")).strip().upper() for r in all_rows}
            missing_clients = sorted(keep - seen)
            all_rows = [r for r in all_rows if str(r.get("ClientCode", "")).strip().upper() in keep]
            for p in (downwards, upwards):
                if p and os.path.exists(p):
                    n = _filter_xlsx(p, keep)
                    log(f"  {os.path.basename(p)}: {n if n >= 0 else '?'} row(s) for these clients")
            downwards = downwards if downwards and os.path.exists(downwards) else None
            upwards = upwards if upwards and os.path.exists(upwards) else None
            if missing_clients:
                log(f"[FRIENDLY_MSG] No records in this window for: {', '.join(missing_clients)}")
            if not all_rows:
                msg = "None of these client codes has any record in this window."
                log(f"[FRIENDLY_MSG] {msg}")
                _finish(base, mode, {"ok": True, "empty": True, "msg": msg, "missing_clients": missing_clients},
                        start, end, clients, by, [], None)
                progress(STEPS, msg, ok=True)
                return 0
        rows = [r for r in all_rows if m.is_success(r)]
        excluded = [r for r in all_rows if not m.is_success(r)]
        log(f"  SUCCESS rows kept: {len(rows)}   excluded: {len(excluded)}")
        if not rows:
            raise RuntimeError("No SUCCESS rows found — nothing to report.")
        report = m.build_report(m.WORK_FOLDER, rows, excluded, log_date)
        attachments = [p for p in (downwards, upwards, report) if p and os.path.exists(p)]

        # ---- STEP 4: deliver ------------------------------------------------------------
        progress(5, "Delivering")
        delivered = ""
        if mode == "close":
            to, cc = _clean_mails(opts.get("to")), _clean_mails(opts.get("cc"))
            if to:
                m.EMAIL_TO, m.EMAIL_CC = to, cc
            m.send_email_with_attachments(m.build_email_subject(label, end), m.build_email_body(label, start, end, {}), attachments)
            delivered = ", ".join(m.EMAIL_TO)
            log(f"✅ Email sent successfully to {delivered}")
        elif opts.get("deliver") == "mail":
            to = _clean_mails(opts.get("mail_to"))
            cc_list = []
            if not to:                                   # box left empty -> use the saved nightly To / CC
                to = _clean_mails(os.environ.get("ALLOC_MAIL_TO"))
                cc_list = _clean_mails(os.environ.get("ALLOC_MAIL_CC"))
            if not to:
                raise RuntimeError("E-mail delivery was chosen but no valid address was given.")
            m.EMAIL_TO, m.EMAIL_CC = to, cc_list
            subj = m.build_email_subject(label, end) + (f" [{', '.join(clients)}]" if clients else "")
            m.send_email_with_attachments(subj, m.build_email_body(label, start, end, {}), attachments)
            delivered = ", ".join(to)
            log(f"✅ Email sent successfully to {delivered}")

        # ---- STEP 5: keep the files / backup ------------------------------------------------
        files, mdir = [], None
        if mode == "close":
            dest = os.path.join(m.BACKUP_DIR, f"{end:%Y-%m-%d}", f"{end:%H.%M}")
            os.makedirs(dest, exist_ok=True)
            for p in produced + [report]:
                if p and os.path.exists(p):
                    shutil.copy2(p, os.path.join(dest, os.path.basename(p)))
            if opening_used:
                od = os.path.join(dest, "opening_balance_files")
                os.makedirs(od, exist_ok=True)
                for p in opening_used:
                    try:
                        shutil.copy2(p, os.path.join(od, os.path.basename(p)))
                    except Exception as e:
                        log(f"  [!] couldn't back up {os.path.basename(p)}: {e}")
            m.save_log_slices(dest, log_offsets)
            log(f"  everything copied to: {dest}")
            files = [os.path.basename(p) for p in attachments]
            mdir = dest
        else:
            mdir = str(base / "Manual" / f"{now:%Y%m%d_%H%M%S}")
            os.makedirs(mdir, exist_ok=True)
            for p in attachments:
                shutil.copy2(p, os.path.join(mdir, os.path.basename(p)))
                files.append(os.path.basename(p))
            log(f"  files kept in: {mdir}")
        for root, _d, _f in os.walk(m.DOWNLOAD_DIR, topdown=False):
            if root != m.DOWNLOAD_DIR:
                try:
                    if not os.listdir(root):
                        os.rmdir(root)
                except OSError:
                    pass
        shutil.rmtree(m.WORK_FOLDER, ignore_errors=True)
        if mode == "close":
            m.close_open_window()
            m.save_window_state(end)
            log(f"  day closed; marker moved to {end:%d-%b-%Y %H:%M}")
        else:
            log("  manual run — the nightly marker was not touched")
        extra = {"ok": True, "emails": len(bucket), "delivered": delivered, "missing_clients": missing_clients,
                 "success_rows": len(rows), "excluded_rows": len(excluded), "opening_files": len(opening_used)}
        _finish(base, mode, extra, start, end, clients, by, files, mdir)
        progress(STEPS, "Done", ok=True)
        log("\n✅ Allocation report completed")
        return 0
    except SystemExit as e:                      # the script's own "stop" messages
        log(f"[FRIENDLY_MSG] The report stopped: {e}")
        progress(STEPS, f"Stopped: {e}", ok=False)
        return 1
    except Exception as e:
        traceback.print_exc()
        log(f"[FRIENDLY_MSG] Allocation report failed: {type(e).__name__}: {e}")
        progress(STEPS, f"{type(e).__name__}: {e}", ok=False)
        if m is not None and mode == "close":
            try:
                m.send_failure_alert(traceback.format_exc()[-3000:])
            except Exception:
                pass
        return 1
    finally:
        if work:
            shutil.rmtree(work, ignore_errors=True)
        _release()


def _clean_mails(v) -> list:
    out = []
    for x in re.split(r"[,;\s]+", str(v or "")):
        x = x.strip()
        if x and re.fullmatch(r"[A-Za-z0-9._%+\-']+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", x) and x not in out:
            out.append(x)
    return out


def _finish(base, mode, extra, start, end, clients, by, files, mdir):
    res = {"mode": mode, "finished": dt.datetime.now().strftime("%d-%m-%Y %H:%M:%S"), "by": by,
           "from": start.strftime("%Y-%m-%d %H:%M"), "to": end.strftime("%Y-%m-%d %H:%M"),
           "clients": clients, "files": files, "dir": mdir, **extra}
    _jwrite(base / ("last_close.json" if mode == "close" else "last_manual.json"), res)
    print("ALLOC_RESULT " + json.dumps(res, ensure_ascii=False), flush=True)


def run_cli(arg_json: str):
    try:
        opts = json.loads(arg_json)
    except Exception:
        print("[FRIENDLY_MSG] Bad options for the allocation report.", flush=True)
        sys.exit(1)
    sys.exit(run(opts))


if __name__ == "__main__":
    run_cli(sys.argv[1] if len(sys.argv) > 1 else "{}")
