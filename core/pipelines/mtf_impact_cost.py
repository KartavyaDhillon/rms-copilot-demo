"""
core/pipelines/mtf_impact_cost.py — MTF <-> Security-Category Impact-Cost reconciliation
(one on-demand run from the dashboard; logic ported from reconcile_mtf_impact_cost.py and
download_nse_cm_impact_cost.py in D:\\SECURITY CAT IMPACT COST).

ONE base folder (env MTF_IMPACT_FOLDER, default D:\\SECURITY CAT IMPACT COST) holds everything:
    TradeBook_<DDMMYYYY>.csv|xlsx|xls      uploaded by the operator
    C_CATG_<MON><YYYY>.T01                  downloaded from NSE (or uploaded if the download failed)
    Reconciliation_MTF_ImpactCost_<DDMMYYYY>.xlsx / Reconciliation_MTF_Dashboard_<DDMMYYYY>.png   (outputs)
    Backup\\<dd-mm-yyyy>\\<HHMMSS>\\          every file used + generated, after a COMPLETED run
    Backup\\last_run.json, Backup\\progress.json   small state files read by the dashboard

HOLDING_<DDMMYYYY>.csv is read from COMMON FILES (SA_COMMON_DIR) and is NEVER moved or deleted.

After a run completes: the TradeBook, the impact-cost file and both outputs are copied to
Backup (size verified), and only those verified originals are then deleted from the base folder.
Nothing is deleted when the run fails.

Secrets: WhatsApp token = env MTFIC_WA_TOKEN, else WA_AUTH_TOKEN; recipients = env MTFIC_WA_NUMBERS
(comma separated, E.164). Nothing is hard-coded here.
"""
from __future__ import annotations

import datetime
import glob
import json
import os
import re
import shutil
import sys
import traceback
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
OK_CATEGORY = 1
TOTAL_STEPS = 5
_BASE_URL = "https://www.nseindia.com"
_REPORTS_PAGE = f"{_BASE_URL}/all-reports"
_API_URL = f"{_BASE_URL}/api/monthly-reports"
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")
_WA_UPLOAD = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/media"
_WA_SEND = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/messages"
_WA_TEMPLATE = "daily_report"
_WA_LABEL = "MTF_RECONCILIATION"


def log(msg=""):
    print(msg, flush=True)


# ───────────────────────────── folders / state ──────────────────────────────
def base_dir() -> Path:
    v = os.getenv("MTF_IMPACT_FOLDER", "").strip()
    return Path(v) if v else Path(r"D:\SECURITY CAT IMPACT COST")


def common_dir() -> Path:
    return Path(os.getenv("SA_COMMON_DIR", r"C:\COMMON FILES"))


def backup_root() -> Path:
    return base_dir() / "Backup"


def ensure_dirs():
    backup_root().mkdir(parents=True, exist_ok=True)


def _read_json(p: Path, default):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json(p: Path, data):
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        tmp = p.with_suffix(p.suffix + ".tmp")
        tmp.write_text(json.dumps(data, indent=1), encoding="utf-8")
        os.replace(tmp, p)
    except Exception:
        pass


def last_run() -> dict:
    return _read_json(backup_root() / "last_run.json", {})


def progress() -> dict:
    return _read_json(backup_root() / "progress.json", {})


_LINES: list = []


def _prog(step=0, msg="", running=True, ok=None):
    """Live progress for the dashboard modal (it polls /mtfic/status)."""
    if msg:
        _LINES.append(f"{datetime.datetime.now():%H:%M:%S} {msg}")
    st = progress() if (running and step == 0) else {}
    data = {"running": running, "step": step or st.get("step", 0), "total": TOTAL_STEPS,
            "msg": msg or st.get("msg", ""), "ok": ok, "lines": _LINES[-25:],
            "updated": datetime.datetime.now().isoformat(timespec="seconds")}
    _write_json(backup_root() / "progress.json", data)


def dmy(d: datetime.date) -> str:
    return d.strftime("%d-%m-%Y")


def parse_date(s) -> datetime.date | None:
    s = (s or "").strip()
    for fmt in ("%d-%m-%Y", "%d%m%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ───────────────────────────── file discovery ───────────────────────────────
_EXTS = (".csv", ".xlsx", ".xls")


def find_dated(folder: Path, prefix: str, d: datetime.date) -> Path | None:
    tag = d.strftime("%d%m%Y")
    for ext in _EXTS:
        p = folder / f"{prefix}_{tag}{ext}"
        if p.exists():
            return p
    return None


def detect_date() -> datetime.date:
    """Date of the TradeBook sitting in the base folder (newest if several), else today."""
    best = None
    for p in base_dir().glob("TradeBook_*"):
        m = re.match(r"^tradebook_(\d{8})\.(csv|xlsx|xls)$", p.name, re.I)
        if not m:
            continue
        d = parse_date(m.group(1))
        if d and (best is None or p.stat().st_mtime > best[1]):
            best = (d, p.stat().st_mtime)
    return best[0] if best else datetime.date.today()


def _month_tag(d: datetime.date) -> str:
    return d.strftime("%b%Y").upper()          # SEP2026


def find_catg(d: datetime.date):
    """(path, note). Order: this month's file in the base folder → newest C_CATG in the base
    folder (flagged) → this month's copy from a previous run's Backup (flagged). (None, '') if none."""
    base = base_dir()
    tag = _month_tag(d)
    exact = base / f"C_CATG_{tag}.T01"
    if exact.exists():
        return exact, ""
    others = sorted(base.glob("C_CATG_*.T01"), key=lambda p: p.stat().st_mtime, reverse=True)
    if others:
        return others[0], f"{others[0].name} is not this month's file ({tag})"
    cands = sorted(backup_root().glob(f"*/*/C_CATG_{tag}.T01"), key=lambda p: p.stat().st_mtime, reverse=True)
    if cands:
        return cands[0], f"copy from backup {cands[0].parent.parent.name}"
    return None, ""


def status(date_str: str = "") -> dict:
    """Everything the dashboard modal needs in one call."""
    ensure_dirs()
    d = parse_date(date_str) or detect_date()
    tb = find_dated(base_dir(), "TradeBook", d)
    hd = find_dated(common_dir(), "HOLDING", d)
    cg, cnote = find_catg(d)
    items = [
        {"key": "tradebook", "label": "TradeBook", "need": f"TradeBook_{d:%d%m%Y}.csv", "ok": bool(tb),
         "file": tb.name if tb else "", "where": "this folder (upload here)"},
        {"key": "holding", "label": "HOLDING", "need": f"HOLDING_{d:%d%m%Y}.csv", "ok": bool(hd),
         "file": hd.name if hd else "", "where": "COMMON FILES (read only, never deleted)"},
        {"key": "catg", "label": "Impact-cost file", "need": f"C_CATG_{_month_tag(d)}.T01",
         "ok": bool(cg) and not cnote.endswith(f"({_month_tag(d)})"),
         "file": cg.name if cg else "", "note": cnote,
         "where": "this folder (downloaded from NSE, or upload here)"},
    ]
    pr = progress()
    running = False
    try:
        from core.pipeline_box import _PB
        running = bool(_PB.get("mtfic", {}).get("running"))
    except Exception:
        pass
    lr = last_run()
    return {"date": dmy(d), "iso": d.isoformat(), "folder": str(base_dir()), "common_folder": str(common_dir()),
            "items": items,
            "ready": bool(tb) and bool(hd),      # the NSE file is fetched by the run itself if missing
            "running": running, "progress": pr,
            "last": lr,
            "wa_numbers": len(_wa_numbers()), "wa_token": bool(_wa_token())}


# ───────────────────────────── uploads ──────────────────────────────────────
def route_uploads(paths) -> dict:
    """paths = temp files saved by the API. TradeBook_<DDMMYYYY> and C_CATG_<MON><YYYY>.T01 go to
    the base folder (same name replaces). Anything else is rejected with a reason."""
    ensure_dirs()
    ok, rejected, notes = [], [], []
    for fp in paths:
        name = os.path.basename(fp)
        low = name.lower()
        try:
            if os.path.getsize(fp) == 0:
                rejected.append(f"{name}: the file is empty")
            elif re.match(r"^tradebook_\d{8}\.(csv|xlsx|xls)$", low):
                dest = base_dir() / name
                _replace(fp, dest)
                ok.append(f"{name} → TradeBook")
            elif re.match(r"^c_catg_[a-z]{3}\d{4}\.t01$", low):
                dest = base_dir() / (name[:-4].upper() + ".T01")
                _replace(fp, dest)
                ok.append(f"{dest.name} → impact-cost file")
            elif re.match(r"^holding_\d{8}", low):
                rejected.append(f"{name}: HOLDING is read from COMMON FILES — upload it with the dashboard's "
                                f"Common files button, not here")
            else:
                rejected.append(f"{name}: not recognised (expected TradeBook_DDMMYYYY.csv or C_CATG_MONYYYY.T01)")
        except Exception as e:
            rejected.append(f"{name}: {e}")
        finally:
            try:
                if os.path.exists(fp):
                    os.remove(fp)
            except OSError:
                pass
    return {"ok": ok, "rejected": rejected, "notes": notes}


def _replace(src, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        dest.unlink()
    shutil.move(str(src), str(dest))


# ───────────────────────────── NSE download ─────────────────────────────────
def _flatten(obj) -> list:
    out = []
    if isinstance(obj, dict):
        for v in obj.values():
            out.extend(_flatten(v))
    elif isinstance(obj, list):
        for v in obj:
            out.extend(_flatten(v))
    elif isinstance(obj, str):
        out.append(obj)
    return out


def _entry_url(entry: dict):
    fp = entry.get("filePath")
    fn = entry.get("fileActlName") or entry.get("fileName")
    if isinstance(fp, str) and fp.startswith("http") and fn:
        return f"{fp.rstrip('/')}/{fn}"
    for v in _flatten(entry):
        if v.startswith("http") and re.search(r"\.(zip|csv|dat|txt|t01|pdf|xls|xlsx)(\?.*)?$", v, re.I):
            return v
    return None


def download_catg(d: datetime.date | None = None) -> Path:
    """Fetch the CM 'Security Category Impact Cost' (.T01) file from nseindia.com into the base
    folder. Restarts the whole fetch on network errors (core/net_retry.py)."""
    import requests
    from core import net_retry
    d = d or datetime.date.today()
    ensure_dirs()

    def once() -> Path:
        s = requests.Session()
        s.headers.update({"User-Agent": _UA, "Accept-Language": "en-US,en;q=0.9"})
        import time
        s.get(_BASE_URL, timeout=15).raise_for_status()
        time.sleep(1)
        s.get(_REPORTS_PAGE, timeout=15).raise_for_status()
        time.sleep(1)
        r = s.get(_API_URL, params={"key": "CM"}, timeout=20,
                  headers={"Accept": "*/*", "Referer": _REPORTS_PAGE, "X-Requested-With": "XMLHttpRequest"})
        r.raise_for_status()
        data = r.json()
        if isinstance(data, dict):
            for k in ("data", "reports", "result"):
                if k in data:
                    data = data[k]
                    break
        if not isinstance(data, list):
            raise RuntimeError("NSE report list came back in an unexpected format.")
        matches = [e for e in data if "impact cost" in json.dumps(e).lower()]
        if not matches:
            raise RuntimeError("NSE has no 'Impact Cost' entry in its CM report list right now.")
        tag = _month_tag(d)
        pick = None
        for e in matches:
            if tag in json.dumps(e).upper():
                pick = e
                break
        if pick is None:
            pick = matches[0]
            log(f"⚠ NSE's entry does not mention {tag}; using the newest one it lists.")
        url = _entry_url(pick)
        if not url:
            raise RuntimeError("NSE listed the report but gave no file link.")
        name = os.path.basename(url.split("?")[0]) or f"C_CATG_{tag}.T01"
        dest = base_dir() / name
        part = dest.with_suffix(dest.suffix + ".part")
        log(f"Downloading {url}")
        with s.get(url, headers={"Referer": _REPORTS_PAGE}, stream=True, timeout=30) as resp:
            resp.raise_for_status()
            with open(part, "wb") as f:
                for chunk in resp.iter_content(chunk_size=8192):
                    f.write(chunk)
        if part.stat().st_size == 0:
            part.unlink()
            raise RuntimeError("NSE sent an empty file.")
        if dest.exists():
            dest.unlink()
        os.replace(part, dest)
        log(f"Saved: {dest.name} ({dest.stat().st_size:,} bytes)")
        return dest

    return net_retry.run(once, "NSE impact-cost download", log)


def run_download(date_str: str = ""):
    """Stand-alone 'Download NSE file' button."""
    d = parse_date(date_str) or detect_date()
    ensure_dirs()
    try:
        p = download_catg(d)
        log(f"✅ Impact-cost file ready: {p.name}")
    except SystemExit:
        raise
    except Exception as e:
        log(f"❌ Download failed: {e}")
        log("   Upload the file by hand with the Upload files button if NSE keeps refusing.")
        raise SystemExit(1)


# ───────────────────────────── reconciliation ───────────────────────────────
def _pd():
    import pandas as pd
    return pd


def read_impact_cost_file(path: Path):
    pd = _pd()
    df = pd.read_csv(path, header=None,
                     names=["RecordType", "Symbol", "Series", "ISIN", "Category", "ImpactCost"],
                     dtype=str, skip_blank_lines=True, engine="python", encoding="utf-8-sig")
    df = df[df["RecordType"].astype(str).str.strip() == "20"].copy()
    df["Symbol"] = df["Symbol"].str.strip().str.upper()
    df["ISIN"] = df["ISIN"].str.strip().str.upper()
    df["Series"] = df["Series"].str.strip().str.upper()
    df["Category"] = pd.to_numeric(df["Category"], errors="coerce")
    return df.reset_index(drop=True)


def build_isin_category_map(impact_df):
    """ISIN -> {Category, Series, ImpactCost}; prefers the EQ-series row (MTF trades only in EQ)."""
    pd = _pd()
    cmap = {}
    for isin, group in impact_df.groupby("ISIN"):
        rows = group.to_dict("records")
        eq = [r for r in rows if r["Series"] == "EQ"]
        ch = eq[0] if eq else rows[0]
        cmap[isin] = {"Category": int(ch["Category"]) if pd.notna(ch["Category"]) else None,
                      "Series": ch["Series"], "ImpactCost": ch["ImpactCost"]}
    return cmap


def read_table_auto(path: Path):
    pd = _pd()
    if path.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path, dtype=str)
    else:
        df = pd.read_csv(path, dtype=str, encoding="utf-8-sig", low_memory=False)
    drop = [c for c in df.columns if str(c).startswith("Unnamed") and df[c].isna().all()]
    return df.drop(columns=drop)


def get_col(df, wanted: str) -> str:
    w = wanted.strip().lower()
    for c in df.columns:
        if str(c).strip().lower() == w:
            return c
    raise KeyError(f"Column '{wanted}' not found. Available columns: {list(df.columns)}")


def build_holding_lookup(path: Path):
    """Symbol -> ISIN (first wins) and the Symbols that map to more than one ISIN."""
    pd = _pd()
    excel = path.suffix.lower() in (".xlsx", ".xls")
    head = pd.read_excel(path, nrows=0) if excel else pd.read_csv(path, nrows=0, encoding="utf-8-sig")
    isin_col, sym_col = get_col(head, "ISIN"), get_col(head, "Symbol")
    if excel:
        chunks = [pd.read_excel(path, usecols=[isin_col, sym_col], dtype=str)]
    else:
        chunks = pd.read_csv(path, usecols=[isin_col, sym_col], dtype=str, encoding="utf-8-sig",
                             chunksize=200_000, low_memory=False)
    lookup, seen = {}, {}
    for ch in chunks:
        ch[sym_col] = ch[sym_col].astype(str).str.strip().str.upper()
        ch[isin_col] = ch[isin_col].astype(str).str.strip().str.upper()
        for sym, isin in zip(ch[sym_col], ch[isin_col]):
            if not sym or sym.lower() == "nan" or not isin or isin == "NAN":
                continue
            seen.setdefault(sym, set()).add(isin)
            lookup.setdefault(sym, isin)
    return lookup, {s: sorted(v) for s, v in seen.items() if len(v) > 1}


def _write_sheet(wb, title, df, status_col=None, ok_value=None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws = wb.create_sheet(title=title[:31])
    if df.empty:
        ws.append(["(no rows)"])
        return
    headers = list(df.columns)
    ws.append(headers)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    sidx = headers.index(status_col) if status_col in headers else None
    okf, badf = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for _, row in df.iterrows():
        vals = ["" if _pd().isna(v) else v for v in row.tolist()]
        ws.append(vals)
        r = ws.max_row
        fill = (okf if vals[sidx] == ok_value else badf) if sidx is not None else None
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=ci)
            cell.border = border
            if fill:
                cell.fill = fill
    ws.freeze_panes = "A2"
    tbl = Table(displayName="Tbl_" + "".join(ch if ch.isalnum() else "_" for ch in title)[:200],
                ref=f"A1:{get_column_letter(len(headers))}{ws.max_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
    ws.add_table(tbl)
    for ci, h in enumerate(headers, start=1):
        sample = df[h].astype(str).tolist()[:500]
        ws.column_dimensions[get_column_letter(ci)].width = min(max([len(str(h))] + [len(v) for v in sample]) + 3, 45)


def _safe_save_fig(fig, path: Path):
    try:
        fig.savefig(path, facecolor=fig.get_facecolor(), pad_inches=0)
        return path
    except OSError as e:
        alt = path.with_name(f"{path.stem}_{datetime.datetime.now():%H%M%S}{path.suffix}")
        log(f"  Could not write {path.name} ({e}) — saved as {alt.name} instead")
        fig.savefig(alt, facecolor=fig.get_facecolor(), pad_inches=0)
        return alt


def render_dashboard_image(out_path: Path, generated_at, mtf_trades, ok_count, check_count,
                           no_holding, isin_missing, wrong_cat, details, sources):
    """One-page PNG report card sized for WhatsApp. details = [(Symbol, ISIN, Reason, Trades)]."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyBboxPatch
    except ImportError:
        log("  (PNG skipped: matplotlib not installed — pip install matplotlib)")
        return None

    INK, INK2, MUTED, SURF, PAGE, BORD = "#0b0b0b", "#52514e", "#898781", "#ffffff", "#f4f5f7", "#dcdfe4"
    GOOD, CRIT, NAVY, ACC, GRID = "#0ca30c", "#d03b3b", "#0f1b2d", "#3987e5", "#e6e7ea"
    W, pad = 1080, 56
    all_clear = check_count == 0
    HEADER_H, ACCENT_H, GAP_H = 176, 4, 40
    BANNER_H, GAP_B, TILE_H, GAP_T, PANEL_H, GAP_P = 84, 32, 160, 32, 232, 32
    MAXD = 10
    nu = len(details)
    show_details = nu > 0
    nd = min(nu, MAXD)
    more = nu > MAXD
    DETAIL_H = 30 + 30 + nd * 28 + (24 if more else 0) + 26
    H = (HEADER_H + ACCENT_H + GAP_H + BANNER_H + GAP_B + TILE_H + GAP_T
         + (0 if all_clear else PANEL_H + GAP_P) + (DETAIL_H + 32 if show_details else 0) + 34 + 32 + 32)

    fig, ax = plt.subplots(figsize=(W / 100, H / 100), dpi=100)
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.invert_yaxis()
    ax.axis("off")
    fig.patch.set_facecolor(PAGE)

    def card(xy, w, h, r=8):
        x, y = xy
        ax.add_patch(FancyBboxPatch((x + 3, y + 4), w, h, boxstyle=f"round,pad=0,rounding_size={r}",
                                    linewidth=0, facecolor=(0.043, 0.055, 0.09, 0.07)))
        ax.add_patch(FancyBboxPatch(xy, w, h, boxstyle=f"round,pad=0,rounding_size={r}",
                                    linewidth=1, edgecolor=BORD, facecolor=SURF))

    ax.add_patch(plt.Rectangle((0, 0), W, HEADER_H, facecolor=NAVY, edgecolor="none"))
    ax.add_patch(plt.Rectangle((0, HEADER_H), W, ACCENT_H, facecolor=ACC, edgecolor="none"))
    ax.text(pad, 46, "D A I L Y   R E C O N C I L I A T I O N   R E P O R T", fontsize=11.5,
            fontweight="bold", color=ACC, va="top")
    ax.text(pad, 74, "MTF Impact-Cost Reconciliation", fontsize=29, fontweight="bold", color="#ffffff", va="top")
    ax.text(pad, 128, f"Generated on {generated_at:%d %B %Y} · {generated_at:%I:%M %p}", fontsize=14.5,
            color="#a9b6c8", va="top")
    tw, th = 214, 32
    tx, ty = W - pad - tw, 40
    ax.add_patch(FancyBboxPatch((tx, ty), tw, th, boxstyle="round,pad=0,rounding_size=16", linewidth=1,
                                edgecolor="#3d5170", facecolor="none"))
    ax.text(tx + tw / 2, ty + th / 2, "INTERNAL USE ONLY", fontsize=10, fontweight="bold", color="#a9b6c8",
            ha="center", va="center")

    y = HEADER_H + ACCENT_H + GAP_H
    sc = GOOD if all_clear else CRIT
    card((pad, y), W - 2 * pad, BANNER_H)
    ax.add_patch(plt.Rectangle((pad, y), 6, BANNER_H, facecolor=sc, edgecolor="none"))
    icx, icy = pad + 46, y + BANNER_H / 2
    ax.add_patch(plt.Circle((icx, icy), 13, facecolor=sc, edgecolor="none"))
    if all_clear:
        ax.plot([icx - 6.5, icx - 1, icx + 7], [icy - 0.6, icy + 5.5, icy - 5], color="#ffffff", linewidth=2.6,
                solid_capstyle="round", solid_joinstyle="round")
    else:
        ax.plot([icx, icx], [icy - 6.5, icy + 1.5], color="#ffffff", linewidth=2.6, solid_capstyle="round")
        ax.add_patch(plt.Circle((icx, icy + 6.5), 1.2, facecolor="#ffffff", edgecolor="none"))
    ax.text(icx + 34, y + 32, "All Clear" if all_clear else "Review Required", fontsize=18, fontweight="bold",
            color=INK, va="center")
    ax.text(icx + 34, y + 60,
            "Every MTF trade matches its Category 1 impact-cost record" if all_clear
            else f"{check_count:,} of {mtf_trades:,} MTF trades need attention", fontsize=13.5, color=INK2,
            va="center")
    y += BANNER_H + GAP_B

    gap = 24
    tile_w = (W - 2 * pad - 2 * gap) / 3
    for i, (label, val, ncol, acol) in enumerate([
            ("MTF TRADES", mtf_trades, INK, ACC), ("OK", ok_count, GOOD, GOOD),
            ("CHECK REQUIRED", check_count, CRIT if check_count else INK, CRIT if check_count else BORD)]):
        x = pad + i * (tile_w + gap)
        card((x, y), tile_w, TILE_H)
        ax.add_patch(plt.Rectangle((x, y), tile_w, 4, facecolor=acol, edgecolor="none"))
        ax.text(x + tile_w / 2, y + 76, f"{val:,}", fontsize=38, fontweight="bold", color=ncol, ha="center",
                va="center")
        ax.text(x + tile_w / 2, y + 128, label, fontsize=12.5, color=MUTED, ha="center", va="center")
    y += TILE_H + GAP_T

    if not all_clear:
        card((pad, y), W - 2 * pad, PANEL_H)
        py = y + 32
        ax.text(pad + 28, py, "CHECK REQUIRED — BREAKDOWN", fontsize=13, fontweight="bold", color=INK, va="top")
        py += 38
        for label, n in [("No Holding record for Symbol", no_holding), ("ISIN not found in Impact-Cost file", isin_missing),
                         ("Category not 1 (wrong group)", wrong_cat)]:
            ax.text(pad + 28, py, label, fontsize=14, color=INK2, va="top")
            ax.text(W - pad - 28, py, f"{n:,}", fontsize=14, color=INK, fontweight="bold", ha="right", va="top")
            py += 32
        py += 8
        ax.plot([pad + 28, W - pad - 28], [py, py], color=GRID, linewidth=1)
        py += 24
        ax.text(pad + 28, py, f"Total Check Required — out of {mtf_trades:,} MTF trades", fontsize=14.5,
                fontweight="bold", color=INK, va="top")
        ax.text(W - pad - 28, py, f"{check_count:,}", fontsize=15, fontweight="bold", color=sc, ha="right", va="top")
        y += PANEL_H + GAP_P

    if show_details:
        card((pad, y), W - 2 * pad, DETAIL_H)
        dy = y + 28
        ax.text(pad + 28, dy, "CHECK REQUIRED — DETAILS (unique Symbol / ISIN)", fontsize=13, fontweight="bold",
                color=INK, va="top")
        dy += 30
        cs, ci, cr, ct = pad + 28, pad + 236, pad + 426, W - pad - 28
        for x, t, ha in [(cs, "SYMBOL", "left"), (ci, "ISIN", "left"), (cr, "REASON", "left"), (ct, "TRADES", "right")]:
            ax.text(x, dy, t, fontsize=10.5, fontweight="bold", color=MUTED, ha=ha, va="top")
        dy += 20
        ax.plot([pad + 28, W - pad - 28], [dy, dy], color=GRID, linewidth=1)
        dy += 10
        for sym, isin, reason, n in details[:MAXD]:
            reason = str(reason)
            if reason.startswith("Category ="):
                reason = reason.replace("Category = ", "Category ").replace(" (expected 1)", " (need 1)")
            elif reason == "No Holding record for this Symbol":
                reason = "No Holding record"
            elif reason == "ISIN not found in Impact-Cost file":
                reason = "ISIN not in Impact-Cost file"
            ax.text(cs, dy, str(sym), fontsize=13, fontweight="bold", color=INK, va="top")
            ax.text(ci, dy, str(isin), fontsize=12, color=INK2, va="top", family="monospace")
            ax.text(cr, dy, reason, fontsize=11.5, color=INK2, va="top")
            ax.text(ct, dy, f"{n:,}", fontsize=13, fontweight="bold", color=CRIT, ha="right", va="top")
            dy += 28
        if more:
            ax.text(pad + 28, dy, f"+ {nu - MAXD} more — see the Check_Required sheet in the Excel file",
                    fontsize=11, color=MUTED, va="top")
        y += DETAIL_H + 32

    y += 34
    ax.text(pad, y, "Sources: " + "  ·  ".join(sources), fontsize=11, color=MUTED, va="top")
    try:
        return _safe_save_fig(fig, out_path)
    except OSError as e:
        log(f"  Fallback save also failed: {e}")
        return None
    finally:
        plt.close(fig)


# ───────────────────────────── WhatsApp ─────────────────────────────────────
def _wa_token() -> str:
    return (os.getenv("MTFIC_WA_TOKEN") or os.getenv("WA_AUTH_TOKEN") or "").strip()


def _wa_numbers() -> list:
    return [n.strip() for n in re.split(r"[,;\s]+", os.getenv("MTFIC_WA_NUMBERS", "")) if n.strip()]


def send_whatsapp(image: Path, report_date: str) -> dict:
    """Never raises; returns {'sent': n, 'failed': n, 'skipped': reason|''}."""
    import requests
    nums, tok = _wa_numbers(), _wa_token()
    if not nums:
        log("  WhatsApp skipped: no numbers set (Credentials → WhatsApp numbers).")
        return {"sent": 0, "failed": 0, "skipped": "no numbers set"}
    if not tok:
        log("  WhatsApp skipped: no bearer token set (Credentials → WhatsApp token).")
        return {"sent": 0, "failed": 0, "skipped": "no token set"}
    try:
        with open(image, "rb") as f:
            up = requests.post(_WA_UPLOAD, headers={"Authorization": f"Bearer {tok}", "accept": "application/json"},
                               files={"file": (image.name, f, "image/png")}, timeout=30)
    except requests.RequestException as e:
        log(f"  WhatsApp upload failed (network error): {e}")
        return {"sent": 0, "failed": len(nums), "skipped": ""}
    if up.status_code not in (200, 201):
        log(f"  WhatsApp upload failed: HTTP {up.status_code}")
        return {"sent": 0, "failed": len(nums), "skipped": ""}
    j = up.json()
    media = j.get("id") or j.get("mediaHandle")
    if not media:
        log("  WhatsApp upload gave no media id.")
        return {"sent": 0, "failed": len(nums), "skipped": ""}
    sent = failed = 0
    for n in nums:
        body = {"to": n, "type": "template", "source": "external",
                "template": {"name": _WA_TEMPLATE, "language": {"code": "en"}, "components": [
                    {"type": "header", "parameters": [{"type": "image", "image": {"id": media}}]},
                    {"type": "body", "parameters": [{"type": "text", "text": _WA_LABEL},
                                                    {"type": "text", "text": report_date}]}]}}
        try:
            r = requests.post(_WA_SEND, headers={"Authorization": f"Bearer {tok}", "Content-Type": "application/json"},
                              json=body, timeout=30)
            if r.status_code == 200:
                sent += 1
                log(f"  WhatsApp sent to {'•' * max(0, len(n) - 4)}{n[-4:]}")
            else:
                failed += 1
                log(f"  WhatsApp send failed (HTTP {r.status_code}) for ••••{n[-4:]}")
        except requests.RequestException as e:
            failed += 1
            log(f"  WhatsApp send failed (network error) for ••••{n[-4:]}: {e}")
    return {"sent": sent, "failed": failed, "skipped": ""}


# ───────────────────────────── backup + clean ───────────────────────────────
def backup_and_clean(files: list, d: datetime.date) -> tuple:
    """Copy every file to Backup\\<dd-mm-yyyy>\\<HHMMSS>\\, verify size, then delete ONLY the verified
    originals from the base folder. Returns (backup_dir, [names backed up], [problems])."""
    dest = backup_root() / dmy(d) / f"{datetime.datetime.now():%H%M%S}"
    dest.mkdir(parents=True, exist_ok=True)
    done, problems = [], []
    base = base_dir().resolve()
    for src in files:
        src = Path(src)
        if not src.exists():
            continue
        try:
            tgt = dest / src.name
            shutil.copy2(src, tgt)
            if tgt.stat().st_size != src.stat().st_size:
                raise OSError("size mismatch after copy")
            done.append(src.name)
            if src.resolve().parent == base:          # only ever delete from the base folder
                try:
                    src.unlink()
                except OSError as e:
                    problems.append(f"{src.name}: backed up, but could not be removed ({e})")
        except Exception as e:
            problems.append(f"{src.name}: NOT backed up, left in place ({e})")
    return dest, done, problems


# ───────────────────────────── the run ──────────────────────────────────────
def _fail(msg):
    log(f"❌ {msg}")
    _prog(0, msg, running=False, ok=False)
    raise SystemExit(1)


def run_all(date_str: str = "", wa: str = "1"):
    try:
        _run_all(date_str, str(wa) not in ("0", "false", "False", "no", ""))
    except SystemExit:
        raise
    except Exception as e:
        traceback.print_exc()
        _fail(f"Run stopped: {e}")


def _run_all(date_str, send_wa):
    pd = _pd()
    ensure_dirs()
    _LINES.clear()
    d = parse_date(date_str) or detect_date()
    log(f"MTF impact-cost reconciliation for {dmy(d)}")

    # 1 ─ inputs
    _prog(1, "Checking input files")
    tb_path = find_dated(base_dir(), "TradeBook", d)
    hold_path = find_dated(common_dir(), "HOLDING", d)
    miss = []
    if not tb_path:
        miss.append(f"TradeBook_{d:%d%m%Y} (upload it with the Upload files button)")
    if not hold_path:
        miss.append(f"HOLDING_{d:%d%m%Y} in COMMON FILES ({common_dir()})")
    if miss:
        for m in miss:
            log(f"[MISSING_FILE] {m}")
        _fail("Missing input file(s): " + "; ".join(miss))
    log(f"TradeBook : {tb_path.name}")
    log(f"HOLDING   : {hold_path} (read only — stays in COMMON FILES)")

    # 2 ─ impact-cost file
    _prog(2, "Getting the NSE impact-cost file")
    catg, note = find_catg(d)
    if catg is None or (note and "not this month's" in note):
        log("Impact-cost file for this month is not in the folder — downloading from NSE…")
        try:
            catg = download_catg(d)
            note = ""
        except Exception as e:
            log(f"⚠ NSE download failed: {e}")
            if catg is None:
                _fail("No impact-cost file: the NSE download failed and there is none in the folder or in Backup. "
                      "Download it yourself and use Upload files.")
            log(f"   Using {catg.name} instead ({note}).")
    elif note:
        log(f"Using impact-cost file: {note}")
    log(f"Impact-cost file: {catg.name}")
    in_base_catg = catg.resolve().parent == base_dir().resolve()

    # 3 ─ reconcile + Excel + PNG
    _prog(3, "Reconciling MTF trades")
    impact_df = read_impact_cost_file(catg)
    cmap = build_isin_category_map(impact_df)
    cat1 = impact_df[impact_df["Category"] == OK_CATEGORY].drop(columns=["RecordType"]).reset_index(drop=True)
    log(f"  impact-cost: {len(impact_df)} rows, {len(cat1)} with Category == {OK_CATEGORY}")

    tb = read_table_auto(tb_path)
    pcol, symcol = get_col(tb, "Product"), get_col(tb, "Symbol")
    get_col(tb, "ClientID")
    tb[pcol] = tb[pcol].astype(str).str.strip().str.upper()
    mtf = tb[tb[pcol] == "MTF"].copy().reset_index(drop=True)
    log(f"  TradeBook: {len(tb)} trades, {len(mtf)} MTF")
    if len(mtf) == 0:
        log("⚠ The TradeBook has no MTF rows — the report will show 0 trades.")

    lookup, ambiguous = build_holding_lookup(hold_path)
    log(f"  HOLDING: {len(lookup)} symbols, {len(ambiguous)} map to more than one ISIN")

    mtf[symcol] = mtf[symcol].astype(str).str.strip().str.upper()
    isins, cats, series, stat, reasons = [], [], [], [], []
    for sym in mtf[symcol]:
        isin = lookup.get(sym)
        if not isin:
            isins.append("N/A"); cats.append("N/A"); series.append("")
            stat.append("Check Required"); reasons.append("No Holding record for this Symbol")
            continue
        info = cmap.get(isin)
        isins.append(isin)
        if info is None:
            cats.append("N/A"); series.append("")
            stat.append("Check Required"); reasons.append("ISIN not found in Impact-Cost file")
        else:
            c = info["Category"]
            cats.append(c if c is not None else "N/A")
            series.append(info["Series"])
            if c == OK_CATEGORY:
                stat.append("OK"); reasons.append("")
            else:
                stat.append("Check Required"); reasons.append(f"Category = {c} (expected {OK_CATEGORY})")
    mtf["ISIN (from Holding)"] = isins
    mtf["Impact Cost Category"] = cats
    mtf["Series Used"] = series
    mtf["Reconciliation Status"] = stat
    mtf["Reason"] = reasons

    ok_rows = mtf[mtf["Reconciliation Status"] == "OK"].reset_index(drop=True)
    chk = mtf[mtf["Reconciliation Status"] == "Check Required"].reset_index(drop=True)
    log(f"Result: {len(ok_rows)} OK, {len(chk)} need checking")

    details = []
    if len(chk):
        src = chk[[symcol, "ISIN (from Holding)", "Reason"]].copy()
        src.columns = ["Symbol", "ISIN", "Reason"]
        g = (src.groupby(["Symbol", "ISIN", "Reason"], dropna=False).size().reset_index(name="Trades")
             .sort_values(["Trades", "Symbol"], ascending=[False, True]))
        details = [(a, b, c, int(n)) for a, b, c, n in g[["Symbol", "ISIN", "Reason", "Trades"]].itertuples(index=False, name=None)]

    relevant = set(mtf[symcol])
    amb_rows = [{"Symbol": s, "ISINs Held": ", ".join(v)} for s, v in ambiguous.items() if s in relevant]
    n_noh = int((mtf["Reason"] == "No Holding record for this Symbol").sum())
    n_nof = int((mtf["Reason"] == "ISIN not found in Impact-Cost file").sum())
    n_wrong = int(mtf["Reason"].astype(str).str.startswith("Category =").sum())

    summary = [
        {"Metric": "Date", "Value": str(d)},
        {"Metric": "Generated at", "Value": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"Metric": "Total trades in TradeBook", "Value": len(tb)},
        {"Metric": "MTF trades", "Value": len(mtf)},
        {"Metric": "OK (Category = 1)", "Value": len(ok_rows)},
        {"Metric": f"Check Required (total, out of {len(mtf)} MTF trades)", "Value": len(chk)},
        {"Metric": "  - No Holding record for Symbol", "Value": n_noh},
        {"Metric": "  - ISIN not found in Impact-Cost file", "Value": n_nof},
        {"Metric": "  - Category not 1 (wrong group)", "Value": n_wrong},
        {"Metric": "Symbols mapping to >1 ISIN in Holding (relevant today)", "Value": len(amb_rows)},
        {"Metric": "Impact-cost file rows (total)", "Value": len(impact_df)},
        {"Metric": "Impact-cost Category 1 rows", "Value": len(cat1)},
    ]
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Summary", pd.DataFrame(summary))
    _write_sheet(wb, "Reconciliation", mtf, "Reconciliation Status", "OK")
    _write_sheet(wb, "Check_Required", chk, "Reconciliation Status", "OK")
    _write_sheet(wb, "ImpactCost_Category1", cat1)
    _write_sheet(wb, "Ambiguous_Holding_Keys", pd.DataFrame(amb_rows))
    tag = d.strftime("%d%m%Y")
    xlsx = base_dir() / f"Reconciliation_MTF_ImpactCost_{tag}.xlsx"
    try:
        wb.save(xlsx)
    except OSError:
        xlsx = xlsx.with_name(f"{xlsx.stem}_{datetime.datetime.now():%H%M%S}{xlsx.suffix}")
        wb.save(xlsx)
    log(f"Saved: {xlsx.name}")

    png = render_dashboard_image(
        base_dir() / f"Reconciliation_MTF_Dashboard_{tag}.png", datetime.datetime.now(), len(mtf), len(ok_rows),
        len(chk), n_noh, n_nof, n_wrong, details, (catg.name, tb_path.name, hold_path.name))
    if png:
        log(f"Saved: {png.name}")

    # 4 ─ WhatsApp
    wa_res = {"sent": 0, "failed": 0, "skipped": "turned off for this run"}
    if not send_whatsapp_enabled(send_wa, png):
        log("WhatsApp: not sent (turned off for this run)" if not send_wa else "WhatsApp: no image to send")
    else:
        _prog(4, "Sending the WhatsApp image")
        wa_res = send_whatsapp(png, dmy(d))
    _prog(4, "WhatsApp step finished")

    # 5 ─ backup, then clean the base folder
    _prog(5, "Backing up files and cleaning the folder")
    used = [tb_path, xlsx] + ([png] if png else []) + ([catg] if in_base_catg else [])
    bdir, done, problems = backup_and_clean(used, d)
    log(f"Backed up {len(done)} file(s) to {bdir}")
    for p in problems:
        log(f"⚠ {p}")
    log("HOLDING was left untouched in COMMON FILES.")

    top = [{"symbol": a, "isin": b, "reason": c, "trades": n} for a, b, c, n in details[:15]]
    result = {
        "date": dmy(d), "finished": datetime.datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
        "mtf": int(len(mtf)), "ok": int(len(ok_rows)), "check": int(len(chk)),
        "no_holding": n_noh, "isin_missing": n_nof, "wrong_category": n_wrong,
        "unique_symbols": len(details), "top": top, "wa": wa_res,
        "sources": {"impact": catg.name, "tradebook": tb_path.name, "holding": hold_path.name},
        "backup_dir": str(bdir),
        "files": {"xlsx": xlsx.name, "png": png.name if png else ""},
        "backup_problems": problems,
    }
    _write_json(backup_root() / "last_run.json", result)
    _prog(5, f"Completed: {len(ok_rows)} OK, {len(chk)} check required", running=False, ok=True)
    log("✅ MTF impact-cost run completed")


def send_whatsapp_enabled(send_wa: bool, png) -> bool:
    return bool(send_wa and png)


def served_file(kind: str):
    """Absolute path of the last run's output (from Backup), or None. kind = xlsx | png."""
    lr = last_run()
    name = (lr.get("files") or {}).get(kind)
    bdir = lr.get("backup_dir")
    if not name or not bdir:
        return None
    p = Path(bdir) / name
    try:
        root = backup_root().resolve()
        if p.exists() and root in p.resolve().parents:
            return p
    except Exception:
        pass
    return None
