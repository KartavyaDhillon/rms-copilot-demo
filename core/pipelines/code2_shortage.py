"""
code2_shortage.py — Tuesday/Thursday Shortage (CODE2), wrapped as a function
=============================================================================
Document 11 (shortage_tuesday_thursday.PY), wrapped in run_code2_shortage().
Reads FULL_FILE_CTCL_VS_NMASS_* + Update Limits from CODE2_FOLDER, computes
shortage (NMASS vs CTCL expiry-margin), filters, builds Output.xlsx with styling,
renders an image, sends WhatsApp, backs up + cleans.

ONLY changes from the original:
  • folder_path        → .env (CODE2_FOLDER)
  • WhatsApp auth_token → .env (WA_AUTH_TOKEN — reuses your existing key)
  • recipient phone     → .env (CODE2_WA_RECIPIENT)
  • wrapped in def run_code2_shortage(); everything else identical.

NOTE: uses xlwings (Excel) for the image step → Windows + Excel required.
"""
import os
import glob
import re
import time
import shutil
from datetime import datetime
import polars as pl
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from dotenv import load_dotenv

load_dotenv()

folder_path = os.getenv("CODE2_FOLDER", r"D:\TUESDAY THURSDAY")


def run_code2_shortage():
    """Tue/Thu shortage: full file + update limits → shortage → image → WhatsApp → backup."""
    full_matches = glob.glob(os.path.join(folder_path, "FULL_FILE_CTCL_VS_NMASS_*.xlsx"))
    if not full_matches:
        print("❌ FULL_FILE_CTCL_VS_NMASS_* not found")
        return {"status": "missing_files", "missing": ["FULL_FILE_CTCL_VS_NMASS_*.xlsx"]}
    full_file = full_matches[0]
    update_files = glob.glob(os.path.join(folder_path, "*Update*Limits*.xlsx"))
    if not update_files:
        print("❌ Update Limits file not found")
        return {"status": "missing_files", "missing": ["*Update*Limits*.xlsx"]}
    update_file = sorted(update_files, reverse=True)[0]
    print(f"Using Full File: {full_file}")
    print(f"Using Update File: {update_file}")

    df_full = pl.read_excel(full_file)
    df_update = pl.read_excel(update_file)

    df_update = df_update.select([
        pl.col("Entity").alias("ACCOUNTID"),
        pl.col("ExpiryMargin").alias("EXPIRY MARGIN")
    ])
    df = df_full.join(df_update, on="ACCOUNTID", how="inner")
    df = df.with_columns([
        (pl.col("DEPOSIT NMASS") - pl.col("EXPIRY MARGIN")).alias("Shortage as per NMASS"),
        (pl.col("DEPOSIT CTCL") - pl.col("EXPIRY MARGIN")).alias("Shortage as per CTCL")
    ])
    df = df.with_columns([
        pl.when((pl.col("Shortage as per NMASS") < 0) & (pl.col("Shortage as per CTCL") < 0))
        .then(pl.min_horizontal(["Shortage as per NMASS", "Shortage as per CTCL"]))
        .otherwise(pl.max_horizontal(["Shortage as per NMASS", "Shortage as per CTCL"]))
        .alias("SHORTAGE")
    ])
    df = df.filter(
        (pl.col("SHORTAGE") < 0) &
        (pl.col("REMARKS") != "VIRTUAL ID") &
        (pl.col("CPCode").is_null() | (pl.col("CPCode") == ""))
    )
    df_final = df.select(["Zone", "ACCOUNTID", "DEPOSIT NMASS", "DEPOSIT CTCL", "EXPIRY MARGIN", "SHORTAGE"])
    df_final = df_final.with_columns([pl.col(pl.Float64).round(2)])
    df_final = df_final.sort("SHORTAGE")

    HAS_DATA = df_final.height > 0
    if not HAS_DATA:
        print("⚠️ No shortage records found.")

    output_path = os.path.join(folder_path, "Output.xlsx")
    df_final.write_excel(output_path)
    print(f"✅ Output saved at: {output_path}")

    wb = load_workbook(output_path)
    ws = wb.active
    ws.freeze_panes = "A2"
    last_row = ws.max_row + 2
    notes = [
        "Shortage is calculated based on the higher value between NMASS and CTCL.",
        "Please verify shortages in the CTCL system before taking any further action.",
        "Shortages related to other server codes may not be accurate and should be verified separately."
    ]
    for i, note in enumerate(notes):
        row_num = last_row + i
        cell = ws.cell(row=row_num, column=1, value=note)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=ws.max_column)

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))
    data_end_row = last_row - 2
    for row in ws.iter_rows(min_row=1, max_row=data_end_row, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col[:data_end_row]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 25)
    for cell in ws[1]:
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    for cell in ws[1]:
        cell.fill = header_fill
    stripe_fill = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=data_end_row), start=2):
        if i % 2 == 0:
            for cell in row:
                cell.fill = stripe_fill
    for cell in ws[1]:
        cell.border = thin_border
    for row in ws.iter_rows(min_row=2, max_row=data_end_row, min_col=1, max_col=6):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
    for row in ws.iter_rows(min_row=2, max_row=data_end_row, min_col=3, max_col=6):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    for row in ws.iter_rows(min_row=2, max_row=data_end_row, min_col=6, max_col=6):
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = Font(color="FF0000", bold=True)
    wb.save(output_path)

    def excel_to_image(excel_file, image_file):
        try:
            import xlwings as xw
            from PIL import ImageGrab
            app = xw.App(visible=False)
            wb2 = app.books.open(excel_file)
            sheet = wb2.sheets[0]
            sheet.api.Rows(1).Font.Bold = True
            sheet.api.Application.ActiveWindow.SplitRow = 1
            sheet.api.Application.ActiveWindow.FreezePanes = True
            sheet.autofit()
            sheet.api.Application.ActiveWindow.Zoom = 85
            try:
                sheet.api.AutoFilterMode = False
                for tbl in sheet.api.ListObjects:
                    tbl.Unlist()
            except Exception:
                pass
            used_range = sheet.used_range
            used_range.api.CopyPicture(Format=2)
            time.sleep(1)
            img = ImageGrab.grabclipboard()
            if img:
                img.save(image_file, "JPEG")
                print(f"✅ Shortage image created: {image_file}")
            else:
                raise Exception("No image in clipboard")
            wb2.close()
            app.quit()
        except Exception as e:
            print(f"❌ Excel to Image failed: {e}")
            raise

    SHORTAGE_IMAGE = os.path.join(folder_path, "Shortage.jpg")
    if HAS_DATA:
        excel_to_image(output_path, SHORTAGE_IMAGE)
    else:
        SHORTAGE_IMAGE = None

    # WhatsApp send
    auth_token = os.getenv("WA_AUTH_TOKEN", "")
    recipient_phone_number = os.getenv("CODE2_WA_RECIPIENT", "")
    upload_url = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/media"
    send_message_url = "https://wb.omni.tatatelebusiness.com/whatsapp-cloud/messages"

    if HAS_DATA and SHORTAGE_IMAGE and os.path.exists(SHORTAGE_IMAGE):
        with open(SHORTAGE_IMAGE, 'rb') as file:
            files = {'file': ('Shortage.jpg', file, 'image/jpeg')}
            headers = {'Authorization': f'Bearer {auth_token}', 'accept': 'application/json'}
            upload_response = requests.post(upload_url, headers=headers, files=files)
        if upload_response.status_code in [200, 201]:
            media_id = upload_response.json().get("id") or upload_response.json().get("mediaHandle")
            message_data = {
                "to": recipient_phone_number, "type": "template", "source": "external",
                "template": {
                    "name": "daily_report", "language": {"code": "en"},
                    "components": [
                        {"type": "header", "parameters": [{"type": "image", "image": {"id": media_id}}]},
                        {"type": "body", "parameters": [
                            {"type": "text", "text": "SHORTAGE REPORT"},
                            {"type": "text", "text": datetime.now().strftime("%d-%m-%Y")}
                        ]}
                    ]
                }
            }
            send_response = requests.post(send_message_url,
                headers={'Authorization': f'Bearer {auth_token}', 'Content-Type': 'application/json'},
                json=message_data)
            if send_response.status_code == 200:
                print("✅ Shortage report sent successfully")
            else:
                print(f"❌ Failed to send: {send_response.text}")
        else:
            print(f"❌ Upload failed: {upload_response.text}")
    else:
        print("❌ Shortage image not found")

    # Backup + cleanup
    try:
        print("📦 Starting backup process...")
        today_folder = datetime.now().strftime("%d-%m-%Y")
        backup_base = os.path.join(folder_path, "BACKUPFILES AUTOMATED")
        backup_date_path = os.path.join(backup_base, today_folder)
        os.makedirs(backup_date_path, exist_ok=True)
        match = re.search(r'_(\d+)\.xlsx$', os.path.basename(full_file))
        suffix = match.group(1) if match else "00"
        suffix_folder = os.path.join(backup_date_path, suffix)
        os.makedirs(suffix_folder, exist_ok=True)
        files_to_backup = [full_file, update_file, output_path]
        if HAS_DATA and SHORTAGE_IMAGE:
            files_to_backup.append(SHORTAGE_IMAGE)
        for file in files_to_backup:
            if os.path.exists(file):
                shutil.copy2(file, os.path.join(suffix_folder, os.path.basename(file)))
                print(f"📁 Backed up: {os.path.basename(file)}")
            else:
                print(f"⚠️ Missing file (skip backup): {file}")
        print(f"✅ Backup completed: {suffix_folder}")
        print("🗑️ Starting cleanup of base folder...")
        for file in files_to_backup:
            try:
                if os.path.exists(file):
                    if file.endswith(".py") or file.endswith(".bat"):
                        continue
                    os.remove(file)
                    print(f"🗑️ Deleted: {os.path.basename(file)}")
            except Exception as e:
                print(f"⚠️ Could not delete {file}: {e}")
        print("✅ Cleanup completed")
    except Exception as e:
        print(f"❌ Backup/Cleanup failed: {e}")

    return {"status": "ok", "had_data": HAS_DATA}


if __name__ == "__main__":
    print(run_code2_shortage())
