"""
core/knowledge.py — the Enterprise Knowledge Engine (Phase 3).

Three worlds, one index, all in a local SQLite database (knowledge.db):

  LEVEL 1 · SOURCE CODE   — every repo file's functions, classes, imports,
                            FastAPI routes and .env keys it reads.
  LEVEL 2 · PROCESSES     — curated + live-enriched metadata for the four
                            production pipelines (purpose, steps, inputs,
                            outputs, emails, failure modes, recovery steps,
                            live success rate / avg runtime).
  LEVEL 3 · OUTPUTS       — every file the pipelines produce, with Excel/CSV
                            files parsed for real business structure: sheet
                            names, column headers, inferred column semantics
                            (client code / margin / exposure / …) and the
                            distinct CLIENT CODES found inside — so "find
                            every report containing client X" is a database
                            lookup, not a guess.

Design rules:
  * Index once, update incrementally (keyed on mtime+size), scan in a single
    background daemon thread — never on a request path, never at server boot.
  * Deep-parse only recent files (KB_DEEP_DAYS, default 180) and only formats
    we can read safely (openpyxl read_only / csv module); older files keep
    shallow metadata (name/date/size/folder) which is still searchable.
  * retrieve(question) returns (context_text, sources[]) — the sources are
    ROWS FROM THIS INDEX, handed to the UI verbatim. The model never invents
    citations because it never produces them.
  * Everything degrades gracefully: no openpyxl, a locked file, a corrupt
    workbook — the file is recorded shallow and the loop moves on.
"""
import os
import re
import csv
import json
import time
import sqlite3
import datetime
import threading

REPO    = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(REPO, "knowledge.db")

DEEP_DAYS      = int(os.getenv("KB_DEEP_DAYS", "180") or 180)
RESCAN_SECS    = int(os.getenv("KB_RESCAN_SECS", "900") or 900)
MAX_XLSX_MB    = 25
MAX_ROWS_SCAN  = 2500
MAX_CLIENTS    = 3000
_SKIP_DIRS     = {".git", ".venv", "venv", "__pycache__", "node_modules", "logs",
                  "certs", ".idea", ".vscode"}

_LOCK   = threading.Lock()
_THREAD = None
_STATE  = {"phase": "idle", "files_indexed": 0, "code_indexed": 0,
           "deep_parsed": 0, "last_cycle": None, "started": None, "errors": 0}

# ── column semantics (Level 3: business meaning, not just header text) ───────
_SEMANTICS = [
    (r"ucc|client\s*code|client\s*id|clientcode|account[\s_]*(no|code|number|id)|acc(ount)?[\s_]?cd", "client_code",   "Client / UCC code"),
    (r"client\s*name|account\s*name|name\s*of\s*client",                                        "client_name",   "Client name"),
    (r"peak\s*margin",                                                                          "peak_margin",   "Peak margin"),
    (r"margin(?!\s*shortage)",                                                                  "margin",        "Margin amount"),
    (r"short(age|fall)",                                                                        "shortage",      "Margin shortage / shortfall"),
    (r"ledger",                                                                                 "ledger",        "Ledger balance"),
    (r"collateral",                                                                             "collateral",    "Collateral value"),
    (r"exposure",                                                                               "exposure",      "Exposure"),
    (r"deposit",                                                                                "deposit",       "Deposit"),
    (r"allocat",                                                                                "allocation",    "Allocation amount"),
    (r"branch",                                                                                 "branch",        "Branch"),
    (r"dealer|rm\b|relationship",                                                               "dealer",        "Dealer / RM"),
    (r"segment|exch(ange)?\b",                                                                  "segment",       "Segment / exchange"),
    (r"expiry",                                                                                 "expiry",        "Expiry date"),
    (r"strike",                                                                                 "strike",        "Strike price"),
    (r"option\s*type|opt\s*type|ce/pe",                                                         "option_type",   "Option type"),
    (r"settle",                                                                                 "settlement",    "Settlement"),
    (r"brokerage",                                                                              "brokerage",     "Brokerage"),
    (r"turnover",                                                                               "turnover",      "Turnover"),
    (r"qty|quantity",                                                                           "quantity",      "Quantity"),
    (r"amount|amt\b|value|val\b",                                                               "amount",        "Amount / value"),
    (r"date|time",                                                                              "datetime",      "Date / time"),
    (r"limit",                                                                                  "limit",         "Limit"),
    (r"mtm|mark\s*to\s*market",                                                                 "mtm",           "Mark-to-market"),
    (r"span",                                                                                   "span",          "SPAN margin"),
]
_SEM_COMPILED = [(re.compile(p, re.I), tag, label) for p, tag, label in _SEMANTICS]
SEM_LABELS = {tag: label for _, tag, label in _SEMANTICS}

_CLIENT_VAL_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_\-/]{1,24}$")
# values that pass the shape test but are clearly NOT client codes
_CLIENT_NOT_RE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,4}$|^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}$"
                            r"|^(total|grand|nil|na|n/a|none)$", re.I)

# ── Level 2: curated process knowledge (live stats merged in at read time) ───
PROCESS_KB = {
    "ctcl": {
        "name": "CTCL vs NMASS", "runs_per_day": 5, "window_env": "CTCL_RUN{i}_WINDOW",
        "purpose": "Reconciles CTCL terminal limits against NSE NMASS collateral/limits, builds zone-wise "
                   "difference reports, Margin Top20 and Payout check, and emails each zone its report.",
        "code": ["core/pipelines/ctcl_vs_nmass.py", "core/pipelines/margin_top20.py",
                 "core/pipelines/payout_download.py", "core/pipelines/payout_check.py", "ctcl_codes/"],
        "inputs": ["NMASS collateral download (connect2nse, captcha-protected)",
                   "NCDEX/NCFE files", "MCX 90% (RSK335)", "RISK_REPORT from C:\\COMMON FILES"],
        "outputs": ["Zone-wise CTCL vs NMASS Excel reports (MARGIN_TOP20_FOLDER)",
                    "Margin Top20 report", "Payout check report"],
        "sends": "Zone emails via CTCL_MAIN/CHND/LUDH/MUMB/DELH_TO/CC; Top20 to CTCL_TOP20_* lists.",
        "failure_modes": ["NMASS captcha failure / wrong password ([NMASS_AUTH_FAIL])",
                          "NCFE locked ([NCFE_LOCKED])", "MCX90 login failure ([MCX90_AUTH_FAIL])",
                          "stale/missing RISK_REPORT in COMMON FILES"],
        "recovery": "Retry the run; fix credentials via the box's 🔑 editor; upload the missing file via "
                    "⬆️ Upload; Top20/Payout have their own standalone retry buttons.",
    },
    "sa": {
        "name": "Short Allocation", "runs_per_day": 4, "window_env": "SA_RUN{i}_WINDOW",
        "purpose": "Downloads 5 exchange files (NSE/MCX/NCDEX + MCX Deposit RSK383), computes short "
                   "allocation per client, runs LAS funding analysis, then emails zone reports and sends "
                   "zone WhatsApp images.",
        "code": ["core/short_allocation.py", "sa_codes/ (downloader, processor, LAS4, WhatsApp sender)",
                 "sa_codes/mcx_deposit_downloader.py"],
        "inputs": ["NSE collateral (captcha)", "MCX SFTP files", "NCDEX files",
                   "MCX Deposit RSK383 (TCS BANCS, RSA token)", "RISK_REPORT / LAS debit emails"],
        "outputs": ["Short Allocation Excel per run (SA_OUTPUT_DIR)", "LAS funding reports",
                    "Scenario/Tentative/Top100/Top15 reports", "WhatsApp report images"],
        "sends": "Zone emails (SA2_MAIL_*/SA_MAIL_*), WhatsApp to SA_WA_* zone numbers.",
        "failure_modes": ["captcha rejection", "MCX Deposit TCS BANCS session busy (shared with CLI's "
                          "MCX Web Allocation — file lock)", "RSA token not entered", "missing input file"],
        "recovery": "Answer captcha/token popups promptly; if MCX session shows busy, wait or Stop the "
                    "other user's download; upload missing files via ⬆️ Upload files.",
    },
    "ms15": {
        "name": "Margin Shortage (15-min)", "runs_per_day": 9, "window_env": "MS15_RUN{i}_WINDOW",
        "purpose": "Every ~15-minute window during market hours: downloads NMASS/NCDEX (auto if missing), "
                   "computes client margin shortage vs collateral+deposit, emails zone reports and sends "
                   "each shortfall client a personalized WhatsApp (Gupshup template) with app links.",
        "code": ["core/pipelines/margin_shortage_15min.py"],
        "inputs": ["NMASS collateral", "NCDEX files", "MCX90 (cross-copied from CTCL)", "RISK_REPORT",
                   "client mobile numbers via Mastertrust webpayout API"],
        "outputs": ["Shortage_vs_CC_Deposit Excel per run (MARGIN_SHORTAGE_FOLDER)",
                    "margin_whatsapp_log table in Postgres"],
        "sends": "Zone emails (MS15_*_TO/CC — currently pointed at kartavya for testing); "
                 "client WhatsApp via Gupshup (test-mode switch available).",
        "failure_modes": ["NMASS/NCDEX captcha", "webpayout OAuth failure", "Gupshup API errors"],
        "recovery": "Retry the slot; use Test Run to verify without sending WhatsApp; check "
                    "whatsapp-report for per-client send status.",
    },
    "cli": {
        "name": "Client Prudential Limit", "runs_per_day": 2, "window_env": "CLI_RUN{i}_WINDOW",
        "purpose": "Builds CLI_PRUDENTIAL_LIMIT.txt from Allocation Detail (+ optional RISK_REPORT/Kambala "
                   "adjustment in the NEW variant), merges CM/FO/CO/NCDEX/MCX Web Allocation downloads, "
                   "uploads the result to FTP and backs up locally + Postgres. Run OLD is password-gated "
                   "(CLI_OLD_RUN_PASS).",
        "code": ["core/pipelines/cli_allocation.py", "core/pipelines/cli_allocation_new.py",
                 "core/pipelines/cli_download.py", "core/pipelines/cli_mcx_download.py",
                 "core/pipelines/cli_upload.py"],
        "inputs": ["Allocation_Detail (FTP)", "CM/FO/CO/ClientNetEffectiveDeposits downloads",
                   "MCX RSK383 via TCS BANCS (same session as SA's MCX Deposit)", "RISK_REPORT + Kambala (NEW)"],
        "outputs": ["CLI_PRUDENTIAL_LIMIT.txt", "MCX_WebAllocationDeallocation.csv",
                    "BACKUPFILES AUTOMATED archive (CLI_FOLDER)"],
        "sends": "FTP upload of the limit file (no emails).",
        "failure_modes": ["FTP 530 login", "TCS BANCS session busy/stale lock", "missing RISK_REPORT/"
                          "Kambala (NEW variant pre-check stops the chain)", "wrong Run OLD password"],
        "recovery": "Stop releases the MCX session lock; standalone Download / MCX Web Allocation / Upload "
                    "buttons re-run individual steps without the full chain.",
    },
}

_FOLDER_PIPE = {"CLI_FOLDER": "cli", "SA_OUTPUT_DIR": "sa", "SA_COMMON_DIR": "shared",
                "MARGIN_SHORTAGE_FOLDER": "ms15", "MARGIN_TOP20_FOLDER": "ctcl",
                "FTP_BASE_LOCAL": "shared", "SA_LAS_FUNDING_DIR": "sa"}


# ── db ────────────────────────────────────────────────────────────────────────
def _db():
    con = sqlite3.connect(DB_PATH, timeout=15)
    con.execute("PRAGMA journal_mode=WAL")
    return con

def _init(con):
    con.executescript("""
    CREATE TABLE IF NOT EXISTS outputs(
      id INTEGER PRIMARY KEY, root TEXT, sub TEXT, name TEXT, ext TEXT,
      folder TEXT, pipe TEXT, mtime REAL, day TEXT, size INTEGER,
      deep INTEGER DEFAULT 0, sheets TEXT, n_rows INTEGER, columns TEXT,
      semantics TEXT, updated REAL, UNIQUE(root, sub));
    CREATE INDEX IF NOT EXISTS idx_out_day  ON outputs(day);
    CREATE INDEX IF NOT EXISTS idx_out_name ON outputs(name);
    CREATE TABLE IF NOT EXISTS clients(
      file_id INTEGER, value TEXT);
    CREATE INDEX IF NOT EXISTS idx_cli_val  ON clients(value);
    CREATE INDEX IF NOT EXISTS idx_cli_file ON clients(file_id);
    CREATE TABLE IF NOT EXISTS code(
      id INTEGER PRIMARY KEY, path TEXT UNIQUE, kind TEXT, doc TEXT,
      functions TEXT, classes TEXT, routes TEXT, imports TEXT, envs TEXT,
      lines INTEGER, mtime REAL);
    CREATE TABLE IF NOT EXISTS meta(k TEXT PRIMARY KEY, v TEXT);
    """)
    con.commit()


# ── Level 1: source code indexer ─────────────────────────────────────────────
_RE_DEF    = re.compile(r"^\s*def\s+([A-Za-z_]\w*)", re.M)
_RE_CLASS  = re.compile(r"^\s*class\s+([A-Za-z_]\w*)", re.M)
_RE_ROUTE  = re.compile(r"@app\.(get|post|put|delete)\(\s*[\"']([^\"']+)")
_RE_IMPORT = re.compile(r"^\s*(?:from\s+([\w.]+)\s+import|import\s+([\w.]+))", re.M)
_RE_ENV    = re.compile(r"getenv\(\s*[\"']([A-Z0-9_]+)")
_RE_JSFUNC = re.compile(r"function\s+([A-Za-z_]\w*)\s*\(")

def _index_code(con):
    n = 0
    for base, dirs, names in os.walk(REPO):
        dirs[:] = [d for d in dirs if d not in _SKIP_DIRS]
        for fn in names:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in (".py", ".html", ".js", ".bat", ".css"):
                continue
            fp = os.path.join(base, fn)
            rel = os.path.relpath(fp, REPO).replace("\\", "/")
            try:
                st = os.stat(fp)
            except OSError:
                continue
            row = con.execute("SELECT mtime FROM code WHERE path=?", (rel,)).fetchone()
            if row and abs(row[0] - st.st_mtime) < 1:
                continue
            try:
                text = open(fp, encoding="utf-8", errors="replace").read()
            except Exception:
                continue
            doc = ""
            m = re.search(r'"""(.{10,400}?)"""', text, re.S) or re.search(r"^(?:#|//)\s*(.{10,200})", text, re.M)
            if m:
                doc = " ".join(m.group(1).split())[:300]
            funcs   = sorted(set(_RE_DEF.findall(text) + (_RE_JSFUNC.findall(text) if ext in (".html", ".js") else [])))
            classes = sorted(set(_RE_CLASS.findall(text))) if ext == ".py" else []
            routes  = [f"{m1.upper()} {p}" for m1, p in _RE_ROUTE.findall(text)]
            imports = sorted(set(a or b for a, b in _RE_IMPORT.findall(text)))[:60] if ext == ".py" else []
            envs    = sorted(set(_RE_ENV.findall(text)))[:80]
            con.execute("""INSERT INTO code(path,kind,doc,functions,classes,routes,imports,envs,lines,mtime)
                           VALUES(?,?,?,?,?,?,?,?,?,?)
                           ON CONFLICT(path) DO UPDATE SET kind=excluded.kind,doc=excluded.doc,
                           functions=excluded.functions,classes=excluded.classes,routes=excluded.routes,
                           imports=excluded.imports,envs=excluded.envs,lines=excluded.lines,mtime=excluded.mtime""",
                        (rel, ext.lstrip("."), doc, json.dumps(funcs[:300]), json.dumps(classes),
                         json.dumps(routes), json.dumps(imports), json.dumps(envs),
                         text.count("\n") + 1, st.st_mtime))
            n += 1
    con.commit()
    return n


# ── Level 3: output indexer ──────────────────────────────────────────────────
def _sem_tags(headers):
    tags = {}
    for h in headers:
        if not h:
            continue
        for rx, tag, _label in _SEM_COMPILED:
            if rx.search(str(h)):
                tags.setdefault(tag, str(h)[:60])
                break
    return tags

def _client_cols(headers):
    out = []
    for i, h in enumerate(headers):
        if h and _SEM_COMPILED[0][0].search(str(h)):
            out.append(i)
    return out

def _deep_parse(fp, ext):
    """Return (sheets, n_rows, columns{sheet:[...]}, semantics{}, clients set)."""
    sheets, columns, sems, clients = [], {}, {}, set()
    n_rows = 0
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(fp, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets[:6]:
                sheets.append(ws.title)
                header, ccols = None, []
                for row in ws.iter_rows(min_row=1, max_row=MAX_ROWS_SCAN, values_only=True):
                    if header is None:
                        vals = ["" if c is None else str(c).strip() for c in row]
                        if sum(1 for v in vals if v) >= 2:
                            header = vals[:60]
                            columns[ws.title] = [v for v in header if v][:50]
                            sems.update(_sem_tags(header))
                            ccols = _client_cols(header)
                        continue
                    n_rows += 1
                    for i in ccols:
                        if i < len(row) and row[i] is not None and len(clients) < MAX_CLIENTS:
                            v = str(row[i]).strip()
                            if _CLIENT_VAL_RE.match(v) and not _CLIENT_NOT_RE.match(v):
                                clients.add(v.upper())
        finally:
            wb.close()
    elif ext in (".csv", ".tsv", ".txt"):
        delim = "\t" if ext == ".tsv" else ","
        with open(fp, encoding="utf-8", errors="replace", newline="") as f:
            rdr = csv.reader(f, delimiter=delim)
            header, ccols = None, []
            for i, row in enumerate(rdr):
                if i > MAX_ROWS_SCAN + 5000:
                    break
                if header is None:
                    vals = [c.strip() for c in row]
                    if sum(1 for v in vals if v) >= 2:
                        header = vals[:60]
                        sheets = ["csv"]
                        columns["csv"] = [v for v in header if v][:50]
                        sems.update(_sem_tags(header))
                        ccols = _client_cols(header)
                    continue
                n_rows += 1
                for ci in ccols:
                    if ci < len(row) and row[ci] and len(clients) < MAX_CLIENTS:
                        v = row[ci].strip()
                        if _CLIENT_VAL_RE.match(v) and not _CLIENT_NOT_RE.match(v):
                            clients.add(v.upper())
    return sheets, n_rows, columns, sems, clients


def _watch_roots():
    try:
        from core.ops_metrics import _WATCH_FOLDER_ENVS, _ROOT_LABELS
        out = []
        for k in _WATCH_FOLDER_ENVS:
            p = os.getenv(k, "").strip()
            if p and os.path.isdir(p):
                out.append((k, _ROOT_LABELS.get(k, k), p))
        return out
    except Exception:
        return []


def _index_outputs(con):
    deep_cutoff = time.time() - DEEP_DAYS * 86400
    n_new = deep = 0
    budget = 2000          # max new/changed rows per cycle — a giant first-time
                            # backlog spreads over several 15-min cycles instead
                            # of hammering the disk in one go
    for root_key, label, base in _watch_roots():
        pipe = _FOLDER_PIPE.get(root_key, "shared")
        for dirpath, dirs, names in os.walk(base):
            if n_new >= budget:
                dirs[:] = []
                break
            rel_dir = os.path.relpath(dirpath, base)
            if rel_dir.count(os.sep) >= 3:            # depth cap
                dirs[:] = []
            dirs[:] = [d for d in dirs if d not in _SKIP_DIRS]
            names = names[:4000]                       # per-directory safety cap
            for fn in names:
                if n_new >= budget:
                    break
                if fn.startswith("~$") or fn.startswith("."):
                    continue
                fp = os.path.join(dirpath, fn)
                sub = os.path.relpath(fp, base).replace("\\", "/")
                try:
                    st = os.stat(fp)
                except OSError:
                    continue
                row = con.execute("SELECT id, mtime, size, deep FROM outputs WHERE root=? AND sub=?",
                                  (root_key, sub)).fetchone()
                if row and abs(row[1] - st.st_mtime) < 1 and row[2] == st.st_size:
                    continue
                ext = os.path.splitext(fn)[1].lower()
                day = datetime.date.fromtimestamp(st.st_mtime).isoformat()
                want_deep = (st.st_mtime >= deep_cutoff and st.st_size < MAX_XLSX_MB * 2**20
                             and ext in (".xlsx", ".xlsm", ".csv", ".tsv"))
                sheets = columns = sems = None
                n_rows = None
                clients = set()
                if want_deep:
                    try:
                        s, nr, cols, sm, clients = _deep_parse(fp, ext)
                        sheets, n_rows, columns, sems = json.dumps(s), nr, json.dumps(cols), json.dumps(sm)
                        deep += 1
                        _STATE["deep_parsed"] += 1
                    except Exception:
                        _STATE["errors"] += 1
                        want_deep = False
                cur = con.execute("""INSERT INTO outputs(root,sub,name,ext,folder,pipe,mtime,day,size,deep,
                                     sheets,n_rows,columns,semantics,updated)
                                     VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                                     ON CONFLICT(root,sub) DO UPDATE SET mtime=excluded.mtime,day=excluded.day,
                                     size=excluded.size,deep=excluded.deep,sheets=excluded.sheets,
                                     n_rows=excluded.n_rows,columns=excluded.columns,
                                     semantics=excluded.semantics,updated=excluded.updated""",
                                  (root_key, sub, fn, ext.lstrip("."), label, pipe, st.st_mtime, day,
                                   st.st_size, 1 if want_deep else 0, sheets, n_rows, columns, sems,
                                   time.time()))
                # lastrowid is unreliable after an UPSERT that took the UPDATE
                # path — always resolve the id explicitly for existing rows so
                # client values can never be attached to the wrong file.
                if row:
                    fid = row[0]
                else:
                    fid = cur.lastrowid or con.execute(
                        "SELECT id FROM outputs WHERE root=? AND sub=?",
                        (root_key, sub)).fetchone()[0]
                if clients:
                    con.execute("DELETE FROM clients WHERE file_id=?", (fid,))
                    con.executemany("INSERT INTO clients(file_id,value) VALUES(?,?)",
                                    [(fid, v) for v in clients])
                n_new += 1
                if n_new % 25 == 0:
                    con.commit()
                    time.sleep(0.05)                   # stay polite to the disk
    con.commit()
    return n_new


# ── background indexer ────────────────────────────────────────────────────────
def _cycle():
    con = _db()
    try:
        _init(con)
        _STATE["phase"] = "code"
        _STATE["code_indexed"] += _index_code(con)
        _STATE["phase"] = "outputs"
        _STATE["files_indexed"] += _index_outputs(con)
        con.execute("INSERT INTO meta(k,v) VALUES('last_cycle',?) ON CONFLICT(k) DO UPDATE SET v=excluded.v",
                    (datetime.datetime.now().isoformat(timespec="seconds"),))
        con.commit()
        _STATE["last_cycle"] = datetime.datetime.now().strftime("%H:%M:%S")
        _STATE["phase"] = "idle"
    except Exception:
        _STATE["errors"] += 1
        _STATE["phase"] = "error"
    finally:
        con.close()

def _loop():
    while True:
        _cycle()
        time.sleep(RESCAN_SECS)

def ensure_indexer():
    global _THREAD
    with _LOCK:
        if _THREAD is None or not _THREAD.is_alive():
            _STATE["started"] = datetime.datetime.now().strftime("%H:%M:%S")
            _THREAD = threading.Thread(target=_loop, daemon=True, name="kb-indexer")
            _THREAD.start()

def status():
    ensure_indexer()
    con = _db()
    try:
        _init(con)
        out = dict(_STATE)
        out["outputs_total"] = con.execute("SELECT COUNT(*) FROM outputs").fetchone()[0]
        out["outputs_deep"]  = con.execute("SELECT COUNT(*) FROM outputs WHERE deep=1").fetchone()[0]
        out["code_files"]    = con.execute("SELECT COUNT(*) FROM code").fetchone()[0]
        out["client_values"] = con.execute("SELECT COUNT(DISTINCT value) FROM clients").fetchone()[0]
        return out
    finally:
        con.close()


# ── search / retrieval ────────────────────────────────────────────────────────
def _like(q):
    return "%" + q.replace("%", "").replace("_", " ").strip() + "%"

def search(q, limit=24):
    """Typed hits across all three levels; every hit is a real index row."""
    ensure_indexer()
    q = (q or "").strip()
    if len(q) < 2:
        return {"hits": []}
    con = _db()
    try:
        _init(con)
        hits = []
        for key, kb in PROCESS_KB.items():
            if q.lower() in (kb["name"] + " " + kb["purpose"] + " " + key).lower():
                hits.append({"type": "process", "key": key, "label": kb["name"],
                             "detail": kb["purpose"][:120]})
        for r in con.execute("""SELECT root,sub,name,folder,day,sheets,columns FROM outputs
                                WHERE name LIKE ? OR columns LIKE ? ORDER BY mtime DESC LIMIT ?""",
                             (_like(q), _like(q), limit)).fetchall():
            hits.append({"type": "output", "root": r[0], "sub": r[1], "label": r[2],
                         "detail": r[3] + " · " + (r[4] or "")})
        cv = con.execute("SELECT COUNT(DISTINCT file_id) FROM clients WHERE value=?",
                         (q.upper(),)).fetchone()[0]
        if cv:
            hits.insert(0, {"type": "client", "value": q.upper(), "label": "Client " + q.upper(),
                            "detail": f"appears in {cv} indexed reports"})
        for r in con.execute("""SELECT path,kind,doc,functions,routes FROM code
                                WHERE path LIKE ? OR functions LIKE ? OR routes LIKE ? OR envs LIKE ?
                                LIMIT 10""", (_like(q), _like(q), _like(q), _like(q))).fetchall():
            hits.append({"type": "code", "path": r[0], "label": r[0],
                         "detail": (r[2] or "")[:110]})
        return {"hits": hits[:limit]}
    finally:
        con.close()


def find_client(code, limit=15):
    ensure_indexer()
    con = _db()
    try:
        _init(con)
        rows = con.execute("""SELECT o.name,o.folder,o.day,o.root,o.sub FROM clients c
                              JOIN outputs o ON o.id=c.file_id WHERE c.value=?
                              ORDER BY o.mtime DESC LIMIT ?""", (code.upper(), limit)).fetchall()
        total = con.execute("SELECT COUNT(*) FROM clients WHERE value=?", (code.upper(),)).fetchone()[0]
        return {"client": code.upper(), "total": total,
                "files": [{"name": r[0], "folder": r[1], "day": r[2], "root": r[3], "sub": r[4]} for r in rows]}
    finally:
        con.close()


def process_kb(pipe=None):
    """Curated metadata + live stats (avg runtime / success rate / windows / recipients)."""
    out = {}
    try:
        from core import ops_metrics
        hist = ops_metrics.history(30)
    except Exception:
        hist = []
    for key, kb in PROCESS_KB.items():
        if pipe and key != pipe:
            continue
        runs = [e for e in hist if e["pipe"] == key]
        ok   = [e for e in runs if e["status"] == "ok"]
        durs = [e["duration_s"] for e in ok if e.get("duration_s")]
        wins = [os.getenv(kb["window_env"].format(i=i), "")
                for i in range(1, kb["runs_per_day"] + 1)]
        out[key] = {**kb, "windows": wins,
                    "stats_30d": {"executions": len(runs), "ok": len(ok),
                                  "failed": sum(1 for e in runs if e["status"] == "failed"),
                                  "success_rate": round(100 * len(ok) / (len(ok) + sum(1 for e in runs if e["status"] == "failed"))) if (ok or any(e["status"] == "failed" for e in runs)) else None,
                                  "avg_runtime_s": int(sum(durs) / len(durs)) if durs else None}}
    return out if not pipe else out.get(pipe, {})


_MONTHS = {m.lower(): i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July", "August",
     "September", "October", "November", "December"], 1)}


# ── Level 4: LIVE table analysis ─────────────────────────────────────────────
# The index knows what files exist and what columns they have; this layer goes
# one step further — when a question is ANALYTICAL ("top 5 clients with highest
# shortfall in the ctcl vs nmass 1st run"), it resolves the actual output file,
# opens it, applies the filter the question implies (e.g. REMARKS=SHORTFALL),
# ranks by the right numeric column, and hands the REAL ROWS to the model.
_ANALYTIC_RE = re.compile(r"\btop\b|\bhighest\b|\blowest\b|\bbottom\b|\blargest\b|"
                          r"\bsmallest\b|\bbiggest\b|\bmax\b|\bmin\b|how many|"
                          r"\bcount\b|\bsum\b|\btotal\b|\baverage\b|\bavg\b|which clients", re.I)
_ORDINALS = {"first": "01", "1st": "01", "second": "02", "2nd": "02", "third": "03",
             "3rd": "03", "fourth": "04", "4th": "04", "fifth": "05", "5th": "05"}
_STOP_TOKS = {"top", "with", "highest", "lowest", "clients", "client", "about",
              "tell", "output", "file", "run", "and", "the", "show", "list",
              "give", "what", "report", "data"}

def _fnum(v):
    """Parse a cell into a float, or None."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s or s in ("-", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None

def _read_rows(fp, ext, cap=40000):
    """(headers, rows) from the first sheet/CSV that has a header row."""
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(fp, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets[:4]:
                header, rows = None, []
                for row in ws.iter_rows(min_row=1, max_row=cap, values_only=True):
                    if header is None:
                        vals = ["" if c is None else str(c).strip() for c in row]
                        if sum(1 for v in vals if v) >= 2:
                            header = vals
                        continue
                    rows.append(row)
                if header and rows:
                    return header, rows
        finally:
            wb.close()
    elif ext in (".csv", ".tsv", ".txt"):
        delim = "\t" if ext == ".tsv" else ","
        with open(fp, encoding="utf-8", errors="replace", newline="") as f:
            header, rows = None, []
            for i, row in enumerate(csv.reader(f, delimiter=delim)):
                if i > cap:
                    break
                if header is None:
                    vals = [c.strip() for c in row]
                    if sum(1 for v in vals if v) >= 2:
                        header = vals
                    continue
                rows.append(row)
            if header and rows:
                return header, rows
    return None, None

def _analyze_struct(q, con):
    """Structured live computation over the best-matching output file, or None."""
    if not _ANALYTIC_RE.search(q):
        return None
    # -- resolve the file the question is about ------------------------------
    toks = [w for w in re.findall(r"[a-z0-9]{3,}", q) if w not in _STOP_TOKS]
    for word, num in _ORDINALS.items():
        if re.search(r"\b" + word + r"\b", q):
            toks.append(num)
    m = re.search(r"\brun\s*0?(\d)\b", q)
    if m:
        toks.append("0" + m.group(1))
    best, best_score = None, 0
    today = datetime.date.today().isoformat()
    for r in con.execute("""SELECT root,sub,name,ext,folder,day,mtime FROM outputs
                            WHERE ext IN ('.xlsx','.xlsm','.csv','.tsv','.txt')
                            ORDER BY mtime DESC LIMIT 500""").fetchall():
        nl = r[2].lower()
        score = sum(2 for t in toks if t in nl) + (1 if r[5] == today else 0)
        if score > best_score:
            best, best_score = r, score
    if not best or best_score < 4:      # need a real name match, not just recency
        return None
    root, sub, name, ext, folder, day, _mt = best
    base = os.getenv(root, "").strip()
    fp = os.path.join(base, sub) if base else None
    if not fp or not os.path.isfile(fp) or os.path.getsize(fp) > MAX_XLSX_MB * 2 * 1024 * 1024:
        return None
    headers, rows = _read_rows(fp, ext)
    if not headers or not rows:
        return None
    # -- classify columns -----------------------------------------------------
    ncols = len(headers)
    sample = rows[: min(len(rows), 400)]
    numeric_cols, text_cols = [], []
    for i in range(ncols):
        vals = [r[i] for r in sample if i < len(r) and r[i] not in (None, "")]
        if not vals:
            continue
        nums = sum(1 for v in vals if _fnum(v) is not None)
        (numeric_cols if nums >= len(vals) * 0.8 else text_cols).append(i)
    if not numeric_cols:
        return None
    # id column: client-code semantic header, else first mostly-unique text col
    id_col = next((i for i in text_cols + list(range(ncols))
                   if headers[i] and _SEM_COMPILED[0][0].search(headers[i])), None)
    if id_col is None:
        for i in text_cols:
            vals = [str(r[i]).strip() for r in sample if i < len(r) and r[i] not in (None, "")]
            if vals and len(set(vals)) > len(vals) * 0.8:
                id_col = i
                break
    if id_col is None:
        id_col = text_cols[0] if text_cols else 0
    # value column: prefer headers matching what the question asks about
    prefs = []
    if re.search(r"short(age|fall)", q):
        prefs += [r"short(age|fall)", r"differen", r"deficit"]
    qwords = re.findall(r"[a-z0-9]{2,}", q)                    # two-word header matches
    for wi in range(len(qwords) - 1):                          # first ("deposit ctcl" >
        prefs.append(re.escape(qwords[wi]) + r"[\s_]*" + re.escape(qwords[wi + 1]))
    prefs += [w for w in toks if len(w) >= 4]                  # then single words
    prefs += [r"differen", r"short(age|fall)", r"amount|value", r"margin", r"deposit"]
    val_col = None
    for p in prefs:
        try:
            rx = re.compile(p, re.I)
        except re.error:
            continue
        val_col = next((i for i in numeric_cols if headers[i] and rx.search(headers[i])), None)
        if val_col is not None:
            break
    if val_col is None:
        val_col = numeric_cols[-1]
    # -- filter: a question word that literally appears as a cell value -------
    filt_desc, data = "", rows
    for i in text_cols:
        if i == id_col:
            continue
        distinct = {str(r[i]).strip().lower() for r in sample if i < len(r) and r[i] not in (None, "")}
        if len(distinct) > 40:
            continue
        hitpair = next(((w, d) for w in re.findall(r"[a-z]{5,}", q)
                        for d in distinct if d and (w == d or w in d)), None)
        if hitpair:
            hit, dval = hitpair
            data = [r for r in rows if i < len(r) and r[i] is not None
                    and hit in str(r[i]).strip().lower()]
            filt_desc = f"{headers[i]} = {dval.upper()}"
            break
    scored = [(str(r[id_col]).strip() if id_col < len(r) and r[id_col] not in (None, "") else "?",
               _fnum(r[val_col]) if val_col < len(r) else None) for r in data]
    scored = [(c, v) for c, v in scored if v is not None]
    if not scored:
        return None
    n = 5
    m = re.search(r"\btop\s*(\d{1,2})\b|\b(\d{1,2})\s+clients\b", q)
    if m:
        n = min(20, int(m.group(1) or m.group(2)))
    reverse = not re.search(r"\blowest\b|\bsmallest\b|\bbottom\b|\bmin\b", q)
    scored.sort(key=lambda x: x[1], reverse=reverse)
    return {"name": name, "folder": folder, "day": day, "root": root, "sub": sub,
            "total_rows": len(rows), "matched": len(scored),
            "filter": filt_desc, "metric": headers[val_col],
            "direction": "top" if reverse else "bottom", "n": min(n, len(scored)),
            "items": [(c, v) for c, v in scored[:n]],
            "sum": sum(v for _, v in scored)}


def kb_fmt_num(v):
    """1,234,567.89 with trailing zeros trimmed — shared by context + answers."""
    return f"{v:,.2f}".rstrip("0").rstrip(".")


def analyze_question(question):
    """Public entry: structured live-table computation for an analytical
    question, or None if it isn't analytical / no output file matches."""
    con = _db()
    try:
        _init(con)
        return _analyze_struct((question or "").lower(), con)
    except Exception:
        return None
    finally:
        con.close()


def _analyze_table(question, q, con):
    """(lines, sources) context form of the same analysis — for retrieve()."""
    a = _analyze_struct(q, con)
    if not a:
        return [], []
    top = "; ".join(f"{c} = {kb_fmt_num(v)}" for c, v in a["items"])
    lines = [f"LIVE TABLE ANALYSIS of {a['name']} ({a['folder']}, {a['day']}): "
             f"{a['total_rows']} data rows" + (f" where {a['filter']}" if a["filter"] else "")
             + f" → {a['matched']} rows. {a['direction'].title()} {a['n']} by {a['metric']}: {top}. "
             f"Sum of {a['metric']} = {kb_fmt_num(a['sum'])}."]
    sources = [{"type": "output", "label": "Live analysis · " + a["name"],
                "detail": f"{a['matched']} rows · ranked by {a['metric']}",
                "root": a["root"], "sub": a["sub"]}]
    return lines, sources


def retrieve(question):
    """(context_text, sources[]) for the AI — evidence only, from the index."""
    ensure_indexer()
    q = (question or "").lower()
    lines, sources = [], []
    words = [w for w in re.findall(r"[a-z0-9_]{3,}", q) if w not in
             ("the", "and", "for", "with", "this", "that", "was", "are", "how",
              "why", "what", "which", "who", "show", "today", "yesterday")]
    con = _db()
    try:
        _init(con)
        # LIVE table analysis first — if the question is analytical (top N /
        # count / sum over a specific output file), compute the real answer
        # from the file's rows and lead the context with it.
        try:
            tl, ts = _analyze_table(question, q, con)
            lines += tl
            sources += ts
        except Exception:
            pass
        # processes mentioned
        for key, kb in PROCESS_KB.items():
            if key in q or kb["name"].lower() in q or any(w in kb["name"].lower() for w in words):
                p = process_kb(key)
                lines.append(f"PROCESS {p['name']}: {p['purpose']} Windows: "
                             + ", ".join(w for w in p["windows"] if w)
                             + f". 30-day: {p['stats_30d']['executions']} executions, "
                             + f"success {p['stats_30d']['success_rate']}%, avg {p['stats_30d']['avg_runtime_s']}s. "
                             + f"Sends: {p['sends']} Failure modes: {'; '.join(p['failure_modes'])}. "
                             + f"Recovery: {p['recovery']}")
                sources.append({"type": "process", "label": p["name"], "detail": "process knowledge base"})
        # explicit client codes in the question (uppercase-ish tokens)
        for tok in set(re.findall(r"\b[A-Za-z0-9]{4,15}\b", question or "")):
            fc = find_client(tok)
            if fc["total"]:
                lines.append(f"CLIENT {fc['client']}: found in {fc['total']} indexed reports. Recent: "
                             + "; ".join(f"{f['name']} ({f['folder']}, {f['day']})" for f in fc["files"][:6]))
                for f in fc["files"][:4]:
                    sources.append({"type": "output", "label": f["name"],
                                    "detail": f["folder"] + " · " + f["day"],
                                    "root": f["root"], "sub": f["sub"]})
                # ROW-LEVEL history from PostgreSQL — the pipelines archive
                # per-client snapshots there (ctcl_snapshots etc.), which is far
                # stronger evidence than file names when the question is about
                # values/trends. Guarded: skipped silently if PG is off.
                if any(k in q for k in ("shortfall", "shortage", "margin", "trend",
                                         "value", "chang", "history", "compare")):
                    try:
                        from core.db_postgres import shortfall_trend
                        tr = shortfall_trend(fc["client"], days=14) or []
                        if tr:
                            lines.append(f"POSTGRES shortfall_trend for {fc['client']} (last 14 days): "
                                         + "; ".join(str(r) for r in tr[:14]))
                            sources.append({"type": "stats", "label": "PostgreSQL · shortfall_trend",
                                            "detail": f"{len(tr)} data points for {fc['client']}"})
                    except Exception:
                        pass
        # output/code hits on the remaining words
        for w in words[:5]:
            for r in con.execute("""SELECT name,folder,day,columns,sheets,root,sub FROM outputs
                                    WHERE name LIKE ? ORDER BY mtime DESC LIMIT 3""", (_like(w),)).fetchall():
                cols = []
                try:
                    cd = json.loads(r[3] or "{}")
                    cols = next(iter(cd.values()), [])[:12]
                except Exception:
                    pass
                lines.append(f"OUTPUT FILE {r[0]} ({r[1]}, {r[2]})"
                             + (f", sheets {r[4]}" if r[4] else "")
                             + (f", columns: {', '.join(cols)}" if cols else ""))
                sources.append({"type": "output", "label": r[0], "detail": r[1] + " · " + (r[2] or ""),
                                "root": r[5], "sub": r[6]})
            for r in con.execute("""SELECT path,doc,routes FROM code
                                    WHERE path LIKE ? OR functions LIKE ? LIMIT 2""",
                                 (_like(w), _like(w))).fetchall():
                lines.append(f"CODE {r[0]}: {(r[1] or '')[:150]}"
                             + (f" Routes: {r[2]}" if r[2] and r[2] != "[]" else ""))
                sources.append({"type": "code", "label": r[0], "detail": (r[1] or "")[:90]})
        # month/count analytics
        if any(k in q for k in ("month", "week", "how many", "compare", "count", "most")) :
            rows = con.execute("""SELECT substr(day,1,7) m, pipe, COUNT(*) FROM outputs
                                  WHERE day >= date('now','-6 months')
                                  GROUP BY m, pipe ORDER BY m DESC""").fetchall()
            if rows:
                agg = {}
                for m, p, c in rows:
                    agg.setdefault(m, []).append(f"{p}:{c}")
                lines.append("OUTPUT FILES PER MONTH (by pipeline): "
                             + " | ".join(f"{m} → {', '.join(v)}" for m, v in list(agg.items())[:6]))
                sources.append({"type": "stats", "label": "Output index",
                                "detail": "files/month aggregation from knowledge.db"})
    finally:
        con.close()
    # dedupe sources
    seen, uniq = set(), []
    for s in sources:
        k = s["type"] + s["label"]
        if k not in seen:
            seen.add(k)
            uniq.append(s)
    return ("\n".join(lines[:24]), uniq[:10])


# ── insights (deterministic anomaly checks) ───────────────────────────────────
def insights():
    ensure_indexer()
    out = []
    try:
        from core import ops_metrics
        hist = ops_metrics.history(14)
        today = datetime.date.today().isoformat()
        todays = [e for e in hist if e["date"] == today]
        for e in todays:
            if e["status"] == "failed":
                out.append({"sev": "err", "text": f"{e['pipe_name']} / {e['label']} failed at {e['start'][11:16]}",
                            "entry": {"logdir": e["logdir"], "file": e["file"], "pipe": e["pipe"],
                                      "pipe_name": e["pipe_name"]}})
        tl = ops_metrics.timeline()
        for i in tl["items"]:
            if i["status"] == "overdue":
                out.append({"sev": "warn", "text": f"{i['id']} ({i['pipe_name']}) is overdue — window {i['window']} has passed"})
        # runtime anomaly: today's ok runs vs 14-day baseline per step label
        base = {}
        for e in hist:
            if e["status"] == "ok" and e.get("duration_s") and e["date"] != today:
                base.setdefault(e["pipe"] + e["label"], []).append(e["duration_s"])
        for e in todays:
            if e["status"] == "ok" and e.get("duration_s"):
                b = base.get(e["pipe"] + e["label"])
                if b and len(b) >= 3:
                    avg = sum(b) / len(b)
                    if e["duration_s"] > max(120, avg * 1.7):
                        out.append({"sev": "warn",
                                    "text": f"{e['label']} ({e['pipe_name']}) took {e['duration_s']}s today — "
                                            f"{round(e['duration_s']/avg,1)}× its 14-day average ({int(avg)}s)"})
        emf = sum(e.get("emails_failed", 0) for e in todays)
        if emf:
            out.append({"sev": "err", "text": f"{emf} email send failure(s) detected in today's logs"})
        # output volume drop vs last same weekday
        con = _db()
        try:
            _init(con)
            last_same = (datetime.date.today() - datetime.timedelta(days=7)).isoformat()
            for pipe in ("ctcl", "sa", "ms15", "cli"):
                t = con.execute("SELECT COUNT(*) FROM outputs WHERE day=? AND pipe=?", (today, pipe)).fetchone()[0]
                p = con.execute("SELECT COUNT(*) FROM outputs WHERE day=? AND pipe=?", (last_same, pipe)).fetchone()[0]
                if p >= 4 and t < p * 0.5 and datetime.datetime.now().hour >= 15:
                    out.append({"sev": "warn",
                                "text": f"{PROCESS_KB[pipe]['name']} produced {t} files today vs {p} last "
                                        f"{datetime.date.today().strftime('%A')} — possible missing outputs"})
        finally:
            con.close()
        if not out:
            out.append({"sev": "ok", "text": "No anomalies detected — executions, runtimes and output volumes look normal."})
    except Exception as e:
        out.append({"sev": "warn", "text": "Insight engine degraded: " + str(e)})
    return {"generated_at": datetime.datetime.now().strftime("%H:%M:%S"), "items": out[:12]}
