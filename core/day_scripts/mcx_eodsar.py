import polars as pl
from datetime import datetime, timedelta
import os
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# =====================================================
# CONFIG
# =====================================================
BASE_DIR = r"C:\PRIYANKA\MCX_EODSAR"
TM_CP_ID = "56720"
CC02_ID = "06436"
MARGIN_SHORTFALL_TRACK_DIR = r"D:\MARGIN SHORTFALL TRACK"


# =====================================================
# PRE-CHECK: ALL REQUIRED FILES (ONE GO)
# =====================================================
def precheck_files():
    missing = []
    files = set(os.listdir(BASE_DIR))

    def exists_pattern(starts):
        return any(
        f.upper().startswith(starts.upper()) and f.upper().endswith(".CSV")
        for f in files
    )


    # ---- MCX MARGIN
    if not exists_pattern(f"MCX_MARGIN_{TM_CP_ID}_"):
        missing.append(f"MCX_MARGIN_{TM_CP_ID}_<DATE>.csv")

    # ---- Allocation (pattern check)
    if not exists_pattern("Allocation_Detail_"):
        missing.append("Allocation_Detail_<ddmmyyyy>.csv")

    # ---- CC02 files (pattern check)
    for p in ["C_CC02", "F_CC02", "X_CC02", "O_CC02"]:
        if not exists_pattern(f"{p}_{CC02_ID}_"):
            missing.append(f"{p}_{CC02_ID}_<ddmmyyyy>.csv")

    # ---- NCDEX
    if not exists_pattern("ClientNetEffectiveDeposits_"):
        missing.append("ClientNetEffectiveDeposits_<ddmmyyyy>.csv")

    if missing:
        print("\n❌ FILE PRE-CHECK FAILED\n")
        for f in missing:
            print("  -", f)
        raise SystemExit(1)

    print("✅ All required files found. Proceeding...\n")

    # return actual MCX margin file (used later)
    return next(
        f for f in files
        if f.startswith(f"MCX_MARGIN_{TM_CP_ID}_") and f.endswith(".csv")
    )


# =====================================================
# HELPERS
# =====================================================
def normalize_client_id(df, col):
    return df.with_columns(
        pl.col(col)
        .cast(pl.Utf8)
        .str.strip_chars()
        .str.to_uppercase()
        .alias(col)
    )

def force_all_numeric_to_float(df: pl.DataFrame) -> pl.DataFrame:
    return df.with_columns([
        pl.when(pl.col(c).cast(pl.Float64, strict=False).is_not_null())
        .then(pl.col(c).cast(pl.Float64, strict=False))
        .otherwise(pl.col(c))
        .alias(c)
        for c in df.columns
    ])




# =====================================================
# FIND MCX MARGIN FILE (AFTER PRE-CHECK)
# =====================================================
margin_file = precheck_files()
margin_path = os.path.join(BASE_DIR, margin_file)

# =====================================================
# SAME-DATE FILE CHECK (all required files must be one trading day)
# =====================================================
import re as _sdc_re

def _same_date_parse(date_str, fmt):
    try:
        if fmt == "YYYYMMDD":
            return datetime.strptime(date_str, "%Y%m%d").date()
        elif fmt == "DDMMYYYY":
            return datetime.strptime(date_str, "%d%m%Y").date()
    except ValueError:
        return None
    return None

_same_date_groups = [
    ("MCX Margin file", rf"MCX_MARGIN_{TM_CP_ID}_(\d{{8}})\.csv", "YYYYMMDD"),
    ("Allocation Detail file", r"Allocation_Detail_(\d{8})\.csv", "DDMMYYYY"),
    ("C_CC02 file", rf"C_CC02_{CC02_ID}_(\d{{8}})\.csv", "DDMMYYYY"),
    ("F_CC02 file", rf"F_CC02_{CC02_ID}_(\d{{8}})\.csv", "DDMMYYYY"),
    ("X_CC02 file", rf"X_CC02_{CC02_ID}_(\d{{8}})\.csv", "DDMMYYYY"),
    ("O_CC02 file", rf"O_CC02_{CC02_ID}_(\d{{8}})\.csv", "DDMMYYYY"),
    ("ClientNetEffectiveDeposits file", r"ClientNetEffectiveDeposits_(\d{8})\.csv", "DDMMYYYY"),
]

_same_date_files_now = os.listdir(BASE_DIR)
_same_date_found = []
for _sdc_label, _sdc_pattern, _sdc_fmt in _same_date_groups:
    _sdc_rx = _sdc_re.compile(_sdc_pattern, _sdc_re.IGNORECASE)
    for _sdc_f in _same_date_files_now:
        _sdc_m = _sdc_rx.match(_sdc_f)
        if _sdc_m:
            _sdc_d = _same_date_parse(_sdc_m.group(1), _sdc_fmt)
            if _sdc_d:
                _same_date_found.append((_sdc_label, _sdc_f, _sdc_d))

if _same_date_found:
    _sdc_anchor_label, _sdc_anchor_file, _sdc_anchor_date = _same_date_found[0]
    _sdc_mismatched = [t for t in _same_date_found if t[2] != _sdc_anchor_date]
    if _sdc_mismatched:
        print("\n❌ Critical: Required files are not all from the same date. Stopping execution.")
        print(f"   Anchor date ({_sdc_anchor_label}): {_sdc_anchor_file} → {_sdc_anchor_date.strftime('%d-%m-%Y')}")
        for _sdc_lbl, _sdc_fn, _sdc_dt in _sdc_mismatched:
            print(f"   Mismatch  ({_sdc_lbl}): {_sdc_fn} → {_sdc_dt.strftime('%d-%m-%Y')}")
        raise SystemExit(1)
    else:
        print("✅ All required files confirmed to be from the same date.")

df_margin = pl.read_csv(margin_path, has_header=False)
df_margin = force_all_numeric_to_float(df_margin)




# =====================================================
# DATE PARSING
# =====================================================
raw_date = df_margin[0, 0]
if isinstance(raw_date, (int, float)):
    file_date = datetime(1899, 12, 30) + timedelta(days=int(round(raw_date)))

else:
    raw_date = str(raw_date).strip().upper()
    for fmt in ("%d%b%Y", "%d-%b-%y", "%d-%b-%Y", "%d/%m/%Y"):
        try:
            file_date = datetime.strptime(raw_date, fmt)
            break
        except ValueError:
            continue
    else:
        raise ValueError(f"Unrecognized date format: {raw_date}")

ddmmyyyy = file_date.strftime("%d%m%Y")
output_excel = f"MCX_EODSAR_{ddmmyyyy}.xlsx"



def format_working_sheet(ws):
    # Make header bold
    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    # Auto-adjust column width
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max_length + 2, 40)  # cap width to avoid huge columns
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width


# =====================================================
# SORT & FILTER
# =====================================================
df_margin = df_margin.sort(pl.col("column_4"), descending=True)
df_margin = df_margin.filter(pl.col("column_17") != 0)

# =====================================================
# WORKING HEADERS
# =====================================================
working_headers = [
    "Date","TM / CP ID","Client ID","Initial Margin","Other Margin","MTM",
    "Reserved_1","Reserved_2","MTM Collected","Initial Margin Collected",
    "Other Margin Collected","Peak Margin Threshold %",
    "Peak Margin","Peak Margin Collected","Peak Margin Shortfall",
    "Intra-day short Allocation","EOD Short Allocation",
    "Intra-day Short Allocation Reported",
    "Reason for short allocation reporting",
    ">>","DFF","category","T5 _remarks","REF","CHK",
    "mcx 2 _alocation file _ ledger",
    "c_cc02 excess","f_cc02 excess","X_cc02 excess","O_cc02 excess",
    "NCDEX_excess","REMARKS"
]

# =====================================================
# BUILD WORKING BASE
# =====================================================
df_working = df_margin.select(df_margin.columns[:17])
df_working.columns = working_headers[:17]
df_working = normalize_client_id(df_working, "Client ID")

# =====================================================
# ALLOCATION
# =====================================================
df_alloc = pl.read_csv(
    os.path.join(BASE_DIR, f"Allocation_Detail_{ddmmyyyy}.csv"),
    infer_schema_length=10000,
    schema_overrides={
        "MtfCashColl": pl.Float64
    }
)

df_alloc = force_all_numeric_to_float(df_alloc)

alloc_lookup = (
    df_alloc
    .select(["AccountID", "TotalAllocation"])
    .rename({
        "AccountID": "Client ID",
        "TotalAllocation": "mcx 2 _alocation file _ ledger"
    })
)
alloc_lookup = normalize_client_id(alloc_lookup, "Client ID")

# =====================================================
# CC02 FULL FILES + LOOKUPS
# =====================================================
def read_cc02(prefix):
    df = pl.read_csv(
        os.path.join(BASE_DIR, f"{prefix}_{CC02_ID}_{ddmmyyyy}.csv"),
        has_header=False,
        schema_overrides={
            "column_1": pl.Utf8
        },
        infer_schema_length=10000
    )

    df = force_all_numeric_to_float(df)

    df = df.with_columns(
        (pl.col("column_6") - pl.col("column_7")).alias("column_8")
    )

    lookup = normalize_client_id(
        df.select(
            pl.col("column_2").alias("Client ID"),
            pl.col("column_8").alias(prefix.lower() + " excess")
        ),
        "Client ID"
    )

    return df, lookup

c_df, c_lookup = read_cc02("C_CC02")
f_df, f_lookup = read_cc02("F_CC02")
x_df, x_lookup = read_cc02("X_CC02")
o_df, o_lookup = read_cc02("O_CC02")

# rename lookup columns properly
c_lookup = c_lookup.rename({"c_cc02 excess": "c_cc02 excess"})
f_lookup = f_lookup.rename({"f_cc02 excess": "f_cc02 excess"})
x_lookup = x_lookup.rename({"x_cc02 excess": "X_cc02 excess"})
o_lookup = o_lookup.rename({"o_cc02 excess": "O_cc02 excess"})

# =====================================================
# NCDEX
# =====================================================
df_ncdex = pl.read_csv(os.path.join(BASE_DIR, f"ClientNetEffectiveDeposits_{ddmmyyyy}.csv"))

df_ncdex = force_all_numeric_to_float(df_ncdex)
df_ncdex = df_ncdex.with_columns(
    (pl.col("Net Effective Deposits After Member Excess Cash ALC") - pl.col("Margin"))
    .alias("NCDEX_excess")
)

ncdex_lookup = normalize_client_id(
    df_ncdex.select(pl.col("Client/CP Code").alias("Client ID"), "NCDEX_excess"),
    "Client ID"
)

# =====================================================
# JOIN LOOKUPS
# =====================================================
df_working = (
    df_working
    .join(alloc_lookup, on="Client ID", how="left")
    .join(c_lookup, on="Client ID", how="left")
    .join(f_lookup, on="Client ID", how="left")
    .join(x_lookup, on="Client ID", how="left")
    .join(o_lookup, on="Client ID", how="left")
    .join(ncdex_lookup, on="Client ID", how="left")
)

# fill numeric nulls
df_working = df_working.with_columns([
    pl.col(c).fill_null(0)
    for c in [
        "mcx 2 _alocation file _ ledger",
        "c_cc02 excess","f_cc02 excess",
        "X_cc02 excess","O_cc02 excess",
        "NCDEX_excess"
    ]
])

# =====================================================
# CHK & T5
# =====================================================
df_working = df_working.with_columns(
    pl.when(pl.col("mcx 2 _alocation file _ ledger") >= 0)
    .then(
        pl.col("c_cc02 excess")
        + pl.col("f_cc02 excess")
        + pl.col("X_cc02 excess")
        + pl.col("O_cc02 excess")
        - pl.col("EOD Short Allocation")
    )
    .otherwise(
        pl.col("mcx 2 _alocation file _ ledger")
        + pl.col("c_cc02 excess")
        + pl.col("f_cc02 excess")
        + pl.col("X_cc02 excess")
        + pl.col("O_cc02 excess")
        - pl.col("EOD Short Allocation")
    )
    .alias("CHK")
)

df_working = df_working.with_columns(
    pl.when(pl.col("CHK") >= 0)
    .then(pl.lit("covered from CASH/ NFO/ CDS"))
    .otherwise(pl.lit("other"))
    .alias("T5 _remarks")
)

# add missing text columns
for col in working_headers:
    if col not in df_working.columns:
        df_working = df_working.with_columns(pl.lit("").alias(col))

df_working = df_working.select(working_headers)

# =====================================================
# WRITE EXCEL (ALL SHEETS)
# =====================================================
wb = Workbook()
ws = wb.active
ws.title = "WORKING"

for r in dataframe_to_rows(df_working.to_pandas(), index=False, header=True):
    ws.append(r)

# ✅ FORMAT WORKING SHEET
format_working_sheet(ws)

def add_sheet(df, name, header):
    ws = wb.create_sheet(name)
    for r in dataframe_to_rows(df.to_pandas(), index=False, header=header):
        ws.append(r)

add_sheet(df_alloc, "ALLOCATION", True)
add_sheet(c_df, "C_CC02", False)
add_sheet(f_df, "F_CC02", False)
add_sheet(x_df, "X_CC02", False)
add_sheet(o_df, "O_CC02", False)
add_sheet(df_ncdex, "NCDEX", True)

wb.save(os.path.join(BASE_DIR, output_excel))
print("✅ WORKING + ALL SHEETS GENERATED SUCCESSFULLY")

# -------- COPY FINAL OUTPUT TO SHARED MARGIN SHORTFALL TRACK FOLDER --------
import shutil as _mst_shutil
try:
    _mst_dir = os.path.join(MARGIN_SHORTFALL_TRACK_DIR, ddmmyyyy)
    os.makedirs(_mst_dir, exist_ok=True)
    _mst_shutil.copy2(os.path.join(BASE_DIR, output_excel), os.path.join(_mst_dir, output_excel))
    print(f"📤 Copied output to shared tracking folder: {_mst_dir}")
except Exception as e:
    print(f"⚠️ Could not copy output to shared tracking folder: {e}")



# =====================================================
# EMAIL → BACKUP → CLEANUP (FINAL & SAFE)
# =====================================================
import smtplib
import time
import shutil
from email.message import EmailMessage


def send_email_with_infinite_retry(output_excel_path, report_date):
    SMTP_SERVER = "smtp3.netcore.co.in"
    SMTP_PORT = 587
    SMTP_EMAIL = "kartavya.dhillon@mastertrust.co.in"
    SMTP_PASSWORD = ""

    TO_LIST = [
         "rms_del@mastertrust.co.in"

    ]

    CC_LIST = [
            ]

    SUBJECT = f"MCX EODSAR Report – {report_date}"

    BODY = f"""Dear Priyanka,

Please find attached the MCX End of Day SAR report for {report_date}.

Regards,
Best regards,
Kartavya Dhillon
Master Capital Services Ltd
📞 011-66331000 | Ext. 84

"""

    msg = EmailMessage()
    msg["From"] = SMTP_EMAIL
    msg["To"] = ", ".join(TO_LIST)
    msg["Cc"] = ", ".join(CC_LIST)
    msg["Subject"] = SUBJECT
    msg.set_content(BODY)

    with open(output_excel_path, "rb") as f:
        msg.add_attachment(
            f.read(),
            maintype="application",
            subtype="octet-stream",
            filename=os.path.basename(output_excel_path)
        )

    recipients = TO_LIST + CC_LIST

    while True:  # infinite retry
        try:
            with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=30) as server:
                server.starttls()
                server.login(SMTP_EMAIL, SMTP_PASSWORD)
                server.send_message(
                    msg,
                    from_addr=SMTP_EMAIL,
                    to_addrs=recipients
                )

            print("✅ Email sent successfully")
            break

        except Exception as e:
            print(f"❌ SMTP error: {e}")
            print("🔁 Retrying in 3 seconds...")
            time.sleep(3)


def backup_and_cleanup_files(base_dir, output_excel_name, files_used):
   

    # MCX_EODSAR_07012026.xlsx → 07012026
    date_part = output_excel_name.split("_")[-1].split(".")[0]

    # 🔑 FIXED ROOT BACKUP PATH
    BACKUP_ROOT = r"C:\PRIYANKA\BACKUP FILES AUTOMATED"

    backup_dir = os.path.join(
        BACKUP_ROOT,
        "MCX_EODSAR",
        date_part
    )
    os.makedirs(backup_dir, exist_ok=True)

    copied_files = []

    # -----------------------------
    # BACKUP
    # -----------------------------
    for file_path in files_used:
        if not os.path.isfile(file_path):
            continue

        filename = os.path.basename(file_path)

        # SAFETY: never touch scripts
        if filename.lower().endswith((".py", ".bat")):
            continue

        dest_path = os.path.join(backup_dir, filename)
        shutil.copy2(file_path, dest_path)
        copied_files.append(file_path)

    # -----------------------------
    # CLEANUP (DELETE)
    # -----------------------------
    for file_path in copied_files:
        try:
            os.remove(file_path)
        except Exception as e:
            print(f"⚠️ Unable to delete {file_path}: {e}")

    print(f"✅ Backup & cleanup completed → {backup_dir}")

# =====================================================
# EXECUTION ORDER (VERY IMPORTANT)
# =====================================================
output_excel_path = os.path.join(BASE_DIR, output_excel)

files_used = [
    margin_path,
    os.path.join(BASE_DIR, f"Allocation_Detail_{ddmmyyyy}.csv"),
    os.path.join(BASE_DIR, f"C_CC02_{CC02_ID}_{ddmmyyyy}.csv"),
    os.path.join(BASE_DIR, f"F_CC02_{CC02_ID}_{ddmmyyyy}.csv"),
    os.path.join(BASE_DIR, f"X_CC02_{CC02_ID}_{ddmmyyyy}.csv"),
    os.path.join(BASE_DIR, f"O_CC02_{CC02_ID}_{ddmmyyyy}.csv"),
    os.path.join(BASE_DIR, f"ClientNetEffectiveDeposits_{ddmmyyyy}.csv"),
    output_excel_path
]

# 1️⃣ SEND EMAIL FIRST
send_email_with_infinite_retry(
    output_excel_path=output_excel_path,
    report_date=ddmmyyyy
)

# 2️⃣ BACKUP + DELETE FILES
backup_and_cleanup_files(
    BASE_DIR,
    output_excel,
    files_used
)

