"""
LAS Table Image Generator — Standalone Test
=============================================
Replaces xlwings + Excel screenshot in LAS4.py / LASUPTO3.py.
Uses openpyxl (read data) + Pillow (draw table) — no Excel required.

HOW TO TEST:
  1. Set INPUT_XLSX to one of your real LAS xlsx files, e.g.:
       r"C:/LAS FUNDING/vertical_1.xlsx"
       r"C:/LAS FUNDING/VERTICAL_2.xlsx"
       r"C:/LAS FUNDING/LAS_FUNDING_OUTPUT_FILTERED.xlsx"
       r"C:/LAS FUNDING/ZONEWISE/LAS_FUNDING_CHND.xlsx"
  2. Run: python las_image_generator.py
  3. Open OUTPUT_IMAGE and compare with the WhatsApp table design.

Dependencies:
    pip install pillow openpyxl
"""

import os
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

# =============================================
# QUICK TEST CONFIG — change these two paths
# =============================================
INPUT_XLSX   = r"C:\LAS FUNDING\VERTICAL_2.xlsx"
OUTPUT_IMAGE = r"C:\LAS FUNDING\VERTICAL_2_SS_TEST.JPG"

# =============================================
# SIZING
# The Excel code sets font=17, row_height=62pt, col_width=20,
# then ScaleWidth(2) + ScaleHeight(2) → everything doubled.
# These pixel values match that scaled output (~2600px wide).
# =============================================

TITLE_ROW_H  = 90    # px
HEADER_ROW_H = 130   # px  (tall because some headers wrap to 3 lines)
DATA_ROW_H   = 90    # px

FONT_SIZE_TITLE  = 26
FONT_SIZE_HEADER = 20
FONT_SIZE_DATA   = 22

BORDER_W  = 2
BORDER_CLR = (0, 0, 0)

TITLE_BG = (173, 216, 230)   # #ADD8E6 — same as openpyxl fill in LAS4
TITLE_FG = (0, 0, 0)
HEADER_BG = (255, 255, 255)
HEADER_FG = (0, 0, 0)
DATA_BG   = (255, 255, 255)
DATA_FG   = (0, 0, 0)

PAD_X = 10   # horizontal text padding inside cell
PAD_Y = 8    # vertical text padding inside cell

# =============================================
# COLUMN CONFIG
# 14 columns (A:N), total ≈ 2570 px.
# col_display key includes \n where the header should wrap.
# =============================================
COL_CONFIG = {
    # display_key (with \n for wrap)           : (px_width, align)
    "ZONE":                                       (150,  "left"),
    "TradingCode":                                (195,  "left"),
    "client_name":                                (280,  "left"),
    "LAS FUNDING":                                (190,  "right"),
    "Deposit(Own\nClient Fund)":                  (205,  "right"),
    "Deposit (After\nadjustment\nShort NOV)":     (225,  "right"),
    "MTM.":                                       (160,  "right"),
    "Total Margin\nUtill":                        (205,  "right"),
    "UP_2%":                                      (160,  "right"),
    "DOWN_2%":                                    (160,  "right"),
    "UP_4%":                                      (160,  "right"),
    "DOWN_4%":                                    (160,  "right"),
    "UP_6%":                                      (160,  "right"),
    "DOWN_6%":                                    (160,  "right"),
}

# Text-align overrides for left-aligned columns
LEFT_COLS = {"ZONE", "TradingCode", "client_name"}

# Font paths (Windows Calibri — falls back to PIL default if missing)
FONT_REGULAR = r"C:\Windows\Fonts\calibri.ttf"
FONT_BOLD    = r"C:\Windows\Fonts\calibrib.ttf"


# =============================================
# INTERNAL HELPERS
# =============================================

def _load_font(path: str, size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(path, size)
    except Exception:
        # PIL built-in — very small but at least won't crash
        return ImageFont.load_default()


def _text_size(draw: ImageDraw, text: str, font) -> tuple[int, int]:
    bb = draw.textbbox((0, 0), text, font=font)
    return bb[2] - bb[0], bb[3] - bb[1]


def _wrap_text(text: str, font, max_width: int, draw: ImageDraw) -> list[str]:
    """Split at explicit \\n first, then word-wrap within max_width."""
    result = []
    for paragraph in str(text).split("\n"):
        words = paragraph.split()
        if not words:
            result.append("")
            continue
        line = ""
        for word in words:
            candidate = (line + " " + word).strip()
            w, _ = _text_size(draw, candidate, font)
            if w <= max_width - PAD_X * 2:
                line = candidate
            else:
                if line:
                    result.append(line)
                line = word
        if line:
            result.append(line)
    return result or [""]


def _draw_cell(
    draw: ImageDraw,
    x: int, y: int, w: int, h: int,
    text,
    font,
    align: str,
    bg: tuple,
    fg: tuple,
):
    # Fill background
    draw.rectangle([x, y, x + w - 1, y + h - 1], fill=bg)
    # Draw border
    draw.rectangle([x, y, x + w - 1, y + h - 1], outline=BORDER_CLR, width=BORDER_W)

    if text is None or str(text).strip() == "":
        return

    lines = _wrap_text(str(text), font, w, draw)

    line_heights = [_text_size(draw, ln, font)[1] for ln in lines]
    gap = 4
    total_h = sum(line_heights) + gap * (len(lines) - 1)

    cur_y = y + (h - total_h) // 2   # vertically centred

    for i, ln in enumerate(lines):
        lw, lh = _text_size(draw, ln, font)
        if align == "center":
            lx = x + (w - lw) // 2
        elif align == "right":
            lx = x + w - lw - PAD_X
        else:
            lx = x + PAD_X
        draw.text((lx, cur_y), ln, font=font, fill=fg)
        cur_y += line_heights[i] + gap


# =============================================
# MAIN PUBLIC FUNCTION
# =============================================

def xlsx_to_image(
    xlsx_path: str,
    output_path: str,
    quality: int = 95,
) -> bool:
    """
    Read an LAS xlsx file (title in A1, headers in row 2, data from row 3)
    and render it as a high-quality JPEG table image.

    Returns True on success, False if file missing / no data.
    """
    if not os.path.exists(xlsx_path):
        print(f"⚠️  File not found — skipping: {xlsx_path}")
        return False

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # Read all cells into a 2-D list (1-indexed → 0-indexed grid)
    grid = []
    for r in range(1, ws.max_row + 1):
        grid.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])

    # Strip trailing completely-empty rows
    while grid and all(v is None for v in grid[-1]):
        grid.pop()

    if len(grid) < 3:
        print(f"⚠️  No data in {os.path.basename(xlsx_path)} — skipping image")
        return False

    title_text  = str(grid[0][0]) if grid[0][0] else "Scenario Analysis LAS Account (in lac.)"
    header_vals = [str(v) if v is not None else "" for v in grid[1]]
    data_rows   = grid[2:]

    # --- Map xlsx headers to COL_CONFIG ---
    def _norm(s: str) -> str:
        return " ".join(s.replace("\n", " ").split()).lower()

    cfg_by_norm = {_norm(k): (k, w, a) for k, (w, a) in COL_CONFIG.items()}

    col_display = []
    col_widths  = []
    col_aligns  = []

    for hdr in header_vals:
        n = _norm(hdr)
        if n in cfg_by_norm:
            dk, w, a = cfg_by_norm[n]
            col_display.append(dk)   # display key has \n for wrapping
            col_widths.append(w)
            col_aligns.append(a)
        else:
            # Fallback for unknown columns
            col_display.append(hdr)
            col_widths.append(160)
            col_aligns.append("left" if hdr in LEFT_COLS else "right")

    total_w = sum(col_widths)
    total_h = TITLE_ROW_H + HEADER_ROW_H + DATA_ROW_H * len(data_rows)

    # --- Create canvas ---
    img  = Image.new("RGB", (total_w, total_h), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    font_title  = _load_font(FONT_BOLD,    FONT_SIZE_TITLE)
    font_header = _load_font(FONT_BOLD,    FONT_SIZE_HEADER)
    font_data   = _load_font(FONT_REGULAR, FONT_SIZE_DATA)

    # --- Title row (spans full width) ---
    draw.rectangle([0, 0, total_w - 1, TITLE_ROW_H - 1], fill=TITLE_BG)
    draw.rectangle([0, 0, total_w - 1, TITLE_ROW_H - 1], outline=BORDER_CLR, width=BORDER_W)
    tw, th = _text_size(draw, title_text, font_title)
    draw.text(((total_w - tw) // 2, (TITLE_ROW_H - th) // 2),
              title_text, font=font_title, fill=TITLE_FG)

    # --- Header row ---
    header_y = TITLE_ROW_H
    cx = 0
    for i, hdr in enumerate(col_display):
        _draw_cell(draw, cx, header_y, col_widths[i], HEADER_ROW_H,
                   hdr, font_header, "center", HEADER_BG, HEADER_FG)
        cx += col_widths[i]

    # --- Data rows ---
    for ri, drow in enumerate(data_rows):
        row_y = TITLE_ROW_H + HEADER_ROW_H + ri * DATA_ROW_H
        cx = 0
        for ci in range(len(col_widths)):
            val = drow[ci] if ci < len(drow) else None
            if isinstance(val, float):
                disp = f"{val:.2f}"
            elif val is None:
                disp = ""
            else:
                disp = str(val)
            _draw_cell(draw, cx, row_y, col_widths[ci], DATA_ROW_H,
                       disp, font_data, col_aligns[ci], DATA_BG, DATA_FG)
            cx += col_widths[ci]

    # --- Save ---
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    img.save(output_path, "JPEG", quality=quality)

    kb = os.path.getsize(output_path) / 1024
    print(f"✅ Saved: {output_path}  ({img.width}x{img.height}px, {kb:.0f} KB)")
    return True


# =============================================
# TEST RUNNER
# =============================================
if __name__ == "__main__":
    ok = xlsx_to_image(INPUT_XLSX, OUTPUT_IMAGE)

    if ok:
        print("✅ Test passed — compare the image with the WhatsApp design.")
    else:
        print("⚠️  Check that INPUT_XLSX is correct and the file has data.")

    # --- To test ALL files at once, uncomment below ---
    # BASE = r"C:\LAS FUNDING"
    # tests = [
    #     (rf"{BASE}\vertical_1.xlsx",                    rf"{BASE}\VERTICAL_1_SS_TEST.JPG"),
    #     (rf"{BASE}\VERTICAL_2.xlsx",                    rf"{BASE}\VERTICAL_2_SS_TEST.JPG"),
    #     (rf"{BASE}\LAS_FUNDING_OUTPUT_FILTERED.xlsx",   rf"{BASE}\FUNDING_SS_TEST.JPG"),
    #     (rf"{BASE}\ZONEWISE\LAS_FUNDING_CHND.xlsx",     rf"{BASE}\CHND_SS_TEST.JPG"),
    #     (rf"{BASE}\ZONEWISE\LAS_FUNDING_LUDH.xlsx",     rf"{BASE}\LUDH_SS_TEST.JPG"),
    #     (rf"{BASE}\ZONEWISE\LAS_FUNDING_MUMB.xlsx",     rf"{BASE}\MUMB_SS_TEST.JPG"),
    # ]
    # for src, dst in tests:
    #     xlsx_to_image(src, dst)
