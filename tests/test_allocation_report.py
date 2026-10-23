"""Helpers of the allocation-report engine (no mailbox, no network).
Run:  python -m unittest tests.test_allocation_report -v
"""
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from core.pipelines import allocation_report as ar


class AllocHelpers(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self._old = os.environ.get("ALLOC_REPORT_FOLDER")
        os.environ["ALLOC_REPORT_FOLDER"] = str(self.tmp / "base")

    def tearDown(self):
        if self._old is None:
            os.environ.pop("ALLOC_REPORT_FOLDER", None)
        else:
            os.environ["ALLOC_REPORT_FOLDER"] = self._old
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_clients_are_cleaned_and_deduplicated(self):
        self.assertEqual(ar._clients("ab123, sn120;AB123  x<y>"), ["AB123", "SN120"])
        self.assertEqual(ar._clients(""), [])

    def test_mails_must_be_valid(self):
        self.assertEqual(ar._clean_mails("a@x.com, bad, b@y.in;a@x.com"), ["a@x.com", "b@y.in"])

    def test_only_one_run_at_a_time(self):
        self.assertTrue(ar._acquire({"mode": "sofar", "by": "t"}))
        self.assertFalse(ar._acquire({"mode": "close"}))
        self.assertEqual(ar.lock_info()["mode"], "sofar")
        ar._release()
        self.assertIsNone(ar.lock_info())

    def test_client_filter_keeps_only_those_rows(self):
        from openpyxl import Workbook, load_workbook
        p = self.tmp / "UPWARDS_Logs_x.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["AllocationDate", "ClientCode", "AllocationAmount"])
        for c, a in (("AB123", 1), ("SN120", 2), ("ab123", 3), ("ZZ999", 4)):
            ws.append(["13-10-2026", c, a])
        wb.save(p)
        self.assertEqual(ar._filter_xlsx(str(p), {"AB123"}), 2)
        rows = list(load_workbook(p).active.iter_rows(values_only=True))
        self.assertEqual([r[1] for r in rows[1:]], ["AB123", "ab123"])
        self.assertEqual(ar._filter_xlsx(str(p), {"NONE"}), 0)
        self.assertFalse(p.exists())                      # nothing left -> no empty file

    def test_status_without_runs(self):
        st = ar.status()
        self.assertFalse(st["running"])
        self.assertEqual(st["last_manual"], {})

class BackOfficeSnapshots(unittest.TestCase):
    """BOD / AFTERNOON / EVENING allocation: named by order of the day's backup folders, optional EVENING, Back Office rows,
    and log rows chained on the latest snapshot."""

    @classmethod
    def setUpClass(cls):
        import importlib.util
        p = Path(__file__).resolve().parents[1] / "core" / "pipelines" / "allocation_report_src.py"
        spec = importlib.util.spec_from_file_location("alloc_src_for_test", str(p))
        cls.m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.m)

    def test_folders_are_named_by_position(self):
        tmp = tempfile.mkdtemp()
        try:
            for n in ("20-10-00", "08-42-42", "16-01-56", "not-a-time"):
                os.makedirs(os.path.join(tmp, n))
            c = self.m.pick_snapshot_folders(tmp)
            self.assertEqual(os.path.basename(c["BOD allocation"]), "08-42-42")
            self.assertEqual(os.path.basename(c["AFTERNOON allocation"]), "16-01-56")
            self.assertEqual(os.path.basename(c["EVENING allocation"]), "20-10-00")
            shutil.rmtree(os.path.join(tmp, "20-10-00"))
            self.assertEqual(set(self.m.pick_snapshot_folders(tmp)), {"BOD allocation", "AFTERNOON allocation"})
            shutil.rmtree(os.path.join(tmp, "16-01-56"))
            self.assertEqual(set(self.m.pick_snapshot_folders(tmp)), {"BOD allocation"})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_missing_evening_backup_is_skipped(self):
        tmp = tempfile.mkdtemp()
        try:
            for n in ("08-42-42", "16-01-56"):
                os.makedirs(os.path.join(tmp, n))
            chosen = self.m.pick_snapshot_folders(tmp)
            self.assertEqual(set(chosen), {"BOD allocation", "AFTERNOON allocation"})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_more_than_three_backup_folders_skips_the_run(self):
        import datetime as _dt
        root = tempfile.mkdtemp()
        old = self.m.BACKUP_ROOT
        try:
            self.m.BACKUP_ROOT = root
            day = os.path.join(root, "22-10-2026")
            for n in ("08-42-42", "12-00-00", "16-01-56"):
                os.makedirs(os.path.join(day, n))
            self.assertIsNone(self.m.check_too_many_snapshots(_dt.date(2026, 10, 22)))       # 3 is fine
            os.makedirs(os.path.join(day, "20-10-00"))
            warn = self.m.check_too_many_snapshots(_dt.date(2026, 10, 22))                    # 4 is not
            self.assertIn("WARNING", warn)
            self.assertIn("SKIPPED", warn)
            self.assertIsNone(self.m.check_too_many_snapshots(_dt.date(2026, 10, 21)))       # other day: nothing
        finally:
            self.m.BACKUP_ROOT = old
            shutil.rmtree(root, ignore_errors=True)

    def test_snapshot_rows_and_chaining(self):
        m = self.m
        rec = lambda t, amt, order: {"ClientCode": "C1", "Segment": "CM", "ActionType": "U", "Amount": amt,
                                     "RequestTime": t, "Minutes": m.parse_time_to_minutes(t), "_order": order,
                                     "ResponseStatus": "SUCCESS"}
        rows = [rec("10:00:00", 120.0, 0), rec("16:05:00", 160.0, 1)]
        tl, segs, summ, latest, _lt = m.build_timeline(
            rows, {"CM": {"C1": 100.0}}, {"CM": "08:40:00"},
            [("AFTERNOON allocation", {"CM": {"C1": 150.0}}, {"CM": "16:00:00"}),
             ("EVENING allocation", {"CM": {"C1": 170.0}}, {"CM": "20:00:00"})])
        labels = [(r[3], r[4]) for r in tl]
        self.assertEqual(labels[0], ("BOD allocation", "Back Office"))
        self.assertIn(("AFTERNOON allocation", "Back Office"), labels)
        self.assertIn(("EVENING allocation", "Back Office"), labels)
        after = next(r for r in tl if r[3] == "AFTERNOON allocation")
        self.assertEqual((after[5], after[6], after[7]), (150.0, 0.0, 150.0))      # previous = new = snapshot value
        late = next(r for r in tl if r[1] == "16:05:00")
        self.assertEqual((late[5], late[6], late[7], late[4]), (150.0, 10.0, 160.0, "Allocated"))   # follows the snapshot
        self.assertEqual(latest["C1"]["CM"], 170.0)
        self.assertEqual(tl[-1][3], "EVENING allocation")


if __name__ == "__main__":
    unittest.main()
